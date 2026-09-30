"""Phase 9/11: explainable scoring and analytics reports (offline, MemoryStore + API)."""

from __future__ import annotations

import io
import tempfile
import unittest
import uuid
from datetime import timedelta
from pathlib import Path

from cloud.intel.core.context import Ctx, ForbiddenError, ValidationError, utcnow
from cloud.intel.platform import Platform
from cloud.intel.scoring.service import BUYING_STAGES, COMPANY_KINDS, MODEL, run_scoring_task
from cloud.intel.store.memory import MemoryStore


def _platform():
    store = MemoryStore()
    user = str(uuid.uuid4())
    ws = store.create_workspace(user, "Scores", f"scores-{uuid.uuid4().hex[:8]}")
    ctx = Ctx(ws["id"], user, "owner")
    platform = Platform(store)
    platform.service("crm").ensure_defaults(ctx)
    return platform, ctx


def _opportunity(platform, ctx, company_id, **extra):
    pipeline = platform.store.first(ctx, "pipelines", {"is_default": True}) or platform.store.all(ctx, "pipelines")[0]
    stage = platform.store.all(ctx, "pipeline_stages", {"pipeline_id": pipeline["id"]})[0]
    values = {"company_id": company_id, "title": "Deal", "pipeline_id": pipeline["id"], "stage_id": stage["id"],
              "status": "open", **extra}
    return platform.store.insert(ctx, "opportunities", values)


class ScoringTests(unittest.TestCase):
    def setUp(self) -> None:
        self.platform, self.ctx = _platform()
        self.store = self.platform.store
        self.company = self.store.insert(self.ctx, "companies", {
            "name": "Acme Manufacturing", "domain": "acme.example", "website": "https://acme.example",
            "industry": "Manufacturing", "country": "United States", "employee_count": 800,
            "technologies": ["SAP S/4HANA"]})
        self.store.insert(self.ctx, "company_technologies", {
            "company_id": self.company["id"], "technology": "SAP S/4HANA", "category": "SAP", "source": "jobs",
            "observed_at": utcnow() - timedelta(days=10), "confidence": 0.9, "evidence_text": "SAP S/4HANA migration"})
        self.store.insert(self.ctx, "hiring_signals", {
            "company_id": self.company["id"], "signal_type": "PROJECT_IMPLEMENTATION", "detected_at": utcnow(),
            "confidence": 0.8, "strength": 70, "summary": "S/4HANA implementation roles", "fingerprint": "fp1",
            "status": "active"})

    def test_company_scores_are_sums_of_explained_factors(self) -> None:
        out = self.platform.service("scoring").score_company(self.ctx, self.company["id"])
        self.assertEqual(set(out["scores"]), set(COMPANY_KINDS))
        for kind, part in out["scores"].items():
            self.assertEqual(part["model"], MODEL)
            self.assertTrue(part["factors"], kind)
            for factor in part["factors"]:
                self.assertEqual(set(factor), {"name", "weight", "value", "points", "reason"})
                self.assertLessEqual(factor["points"], factor["weight"])
                self.assertTrue(factor["reason"])
            self.assertAlmostEqual(part["score"], min(100.0, round(sum(f["points"] for f in part["factors"]), 1)),
                                   places=1)
        tech = {f["name"]: f for f in out["scores"]["technology"]["factors"]}
        self.assertEqual(tech["core_platform"]["points"], 25)
        self.assertEqual(tech["recency"]["points"], 20)
        self.assertTrue(any(e["type"] == "company_technology" for e in out["scores"]["technology"]["evidence"]))
        self.assertTrue(any(e["type"] == "hiring_signal" and e["id"] for e in out["scores"]["hiring"]["evidence"]))
        self.assertEqual(out["scores"]["buying_stage"]["label"], "researching")

        saved = self.store.get(self.ctx, "companies", self.company["id"])
        self.assertEqual(saved["technology_score"], out["scores"]["technology"]["score"])
        self.assertEqual(saved["buying_stage"], "researching")
        self.assertIsNotNone(saved["scored_at"])
        # The legacy breakdown shape the UI and research agent read is kept.
        self.assertIn("account", saved["score_breakdown"])
        self.assertIn("component", saved["score_breakdown"]["account"][0])

    def test_snapshots_append_only_on_change(self) -> None:
        scoring = self.platform.service("scoring")
        scoring.score_company(self.ctx, self.company["id"])
        first = self.store.count(self.ctx, "score_snapshots", {"entity_id": self.company["id"]})
        self.assertEqual(first, len(COMPANY_KINDS))
        scoring.score_company(self.ctx, self.company["id"])
        self.assertEqual(self.store.count(self.ctx, "score_snapshots", {"entity_id": self.company["id"]}), first)
        _opportunity(self.platform, self.ctx, self.company["id"])
        scoring.score_company(self.ctx, self.company["id"])
        history = scoring.history(self.ctx, "company", self.company["id"], kind="buying_stage")
        self.assertEqual([h["label"] for h in history], ["evaluating", "researching"])

    def test_buying_stage_follows_evidence(self) -> None:
        scoring = self.platform.service("scoring")
        bare = self.store.insert(self.ctx, "companies", {"name": "Nobody Inc"})
        self.assertEqual(scoring.score_company(self.ctx, bare["id"])["scores"]["buying_stage"]["label"], "unaware")
        contact = self.store.insert(self.ctx, "contacts", {"company_id": bare["id"], "full_name": "Pat Doe",
                                                           "email": "pat@nobody.example"})
        self.store.insert(self.ctx, "message_events", {"contact_id": contact["id"], "event": "replied",
                                                       "occurred_at": utcnow()})
        part = scoring.score_company(self.ctx, bare["id"])["scores"]["buying_stage"]
        self.assertEqual(part["label"], "engaged")
        self.assertEqual({f["name"]: f for f in part["factors"]}["engagement"]["points"], 25)
        _opportunity(self.platform, self.ctx, bare["id"], status="won")
        self.assertEqual(scoring.score_company(self.ctx, bare["id"])["scores"]["buying_stage"]["label"], "customer")
        self.assertEqual(BUYING_STAGES[-1], "customer")

    def test_no_technology_scores_zero_with_reason(self) -> None:
        bare = self.store.insert(self.ctx, "companies", {"name": "Plain Co"})
        part = self.platform.service("scoring").score_company(self.ctx, bare["id"])["scores"]["technology"]
        self.assertEqual(part["score"], 0)
        self.assertIn("no technologies", part["factors"][0]["reason"])

    def test_contact_score_and_evidence(self) -> None:
        contact = self.store.insert(self.ctx, "contacts", {
            "company_id": self.company["id"], "full_name": "Jo CIO", "seniority": "c_level", "function": "it",
            "email": "jo@acme.example", "email_status": "VALID", "confidence": 0.9})
        self.store.insert(self.ctx, "email_validations", {
            "email": "jo@acme.example", "status": "VALID", "score": 95, "provider": "emaillistverify",
            "validated_at": utcnow(), "expires_at": utcnow() + timedelta(days=30)})
        out = self.platform.service("scoring").score_contact(self.ctx, contact["id"])
        part = out["scores"]["contact"]
        self.assertEqual(part["score"], 98.0)
        self.assertEqual({f["name"] for f in part["factors"]}, {"seniority", "function", "email", "confidence"})
        self.assertTrue(any(e["type"] == "email_validation" for e in part["evidence"]))
        self.assertEqual(self.store.get(self.ctx, "contacts", contact["id"])["contact_score"], 98.0)

    def test_explain_is_read_only_for_viewers(self) -> None:
        viewer = str(uuid.uuid4())
        self.store.add_member(self.ctx, viewer, "viewer")
        vctx = Ctx(self.ctx.workspace_id, viewer, "viewer")
        out = self.platform.service("scoring").explain(vctx, "company", self.company["id"])
        self.assertFalse(out["persisted"])
        self.assertIn("technology", out["how"])
        self.assertEqual(self.store.count(self.ctx, "score_snapshots"), 0)
        with self.assertRaises(ForbiddenError):
            self.platform.service("scoring").rescore(vctx, "company", self.company["id"])
        with self.assertRaises(ValidationError):
            self.platform.service("scoring").explain(vctx, "deal", self.company["id"])

    def test_scoring_task_rescores_all_companies(self) -> None:
        self.store.insert(self.ctx, "companies", {"name": "Second Co"})
        result = run_scoring_task(self.platform, self.ctx.as_system(), {"params": {"all": True}}, None)
        self.assertEqual(result, {"targets": 2, "scored": 2, "errors": 0})


class ReportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.platform, self.ctx = _platform()
        self.store = self.platform.store
        self.reports = self.platform.service("reports")

    def test_every_report_runs_empty_and_says_so(self) -> None:
        for info in self.reports.catalog():
            out = self.reports.run(self.ctx, info["key"])
            self.assertEqual(out["report"], info["key"])
            self.assertTrue(out["columns"], info["key"])
            self.assertIn("start", out["range"])
            if info["key"] not in ("prospect_conversion", "funnel"):
                self.assertTrue(out["empty"], info["key"])

    def test_campaign_attribution_never_invents_rates_or_money(self) -> None:
        company = self.store.insert(self.ctx, "companies", {"name": "Acme"})
        campaign = self.store.insert(self.ctx, "campaigns", {"key": "c1", "name": "SAP push"})
        contact = self.store.insert(self.ctx, "contacts", {"full_name": "A B", "company_id": company["id"]})
        for event in ("sent", "sent", "replied"):
            self.store.insert(self.ctx, "message_events", {"campaign_id": campaign["id"], "contact_id": contact["id"],
                                                           "event": event, "occurred_at": utcnow()})
        _opportunity(self.platform, self.ctx, company["id"], campaign_id=campaign["id"], status="won")
        row = self.reports.run(self.ctx, "campaign_attribution")["rows"][0]
        self.assertEqual(row["sent"], 2)
        self.assertEqual(row["reply_rate"], 0.5)
        self.assertEqual(row["won"], 1)
        self.assertIsNone(row["won_value"])  # no amount was recorded
        _opportunity(self.platform, self.ctx, company["id"], campaign_id=campaign["id"], status="won", amount=1200)
        self.assertEqual(self.reports.run(self.ctx, "campaign_attribution")["rows"][0]["won_value"], 1200.0)

    def test_funnel_and_conversion_rates(self) -> None:
        for i in range(4):
            self.store.insert(self.ctx, "companies", {"name": f"Co {i}"})
        rows = {r["step"]: r for r in self.reports.run(self.ctx, "funnel")["rows"]}
        self.assertEqual(rows["Companies added"]["count"], 4)
        self.assertIsNone(rows["Companies added"]["conversion"])
        self.assertEqual(rows["Contacts added"]["conversion"], 0.0)
        self.assertIsNone(rows["Emails sent"]["conversion"])  # previous step was 0

    def test_date_range_and_filters_are_validated(self) -> None:
        with self.assertRaises(ValidationError):
            self.reports.run(self.ctx, "funnel", start="2026-02-01", end="2026-01-01")
        with self.assertRaises(ValidationError):
            self.reports.run(self.ctx, "funnel", filters={"owner_id": "x"})
        with self.assertRaises(ValidationError):
            self.reports.run(self.ctx, "nope")
        old = self.reports.run(self.ctx, "activity_performance", start="2020-01-01", end="2020-01-03")
        self.assertEqual(len(old["series"][0]["points"]), 3)

    def test_hiring_trends_and_ai_usage(self) -> None:
        company = self.store.insert(self.ctx, "companies", {"name": "Hiring Co"})
        self.store.insert(self.ctx, "hiring_signals", {"company_id": company["id"], "signal_type": "HIRING_SPIKE",
                                                       "detected_at": utcnow(), "fingerprint": "f", "status": "active"})
        out = self.reports.run(self.ctx, "hiring_trends")
        self.assertEqual(out["rows"][0]["signal_type"], "HIRING_SPIKE")
        self.assertEqual(sum(p["y"] for p in out["series"][0]["points"]), 1)
        self.store.insert(self.ctx.as_system(), "ai_usage", {"provider": "gemini", "model": "m", "purpose": "scraper",
                                                             "success": True, "total_tokens": 100})
        row = self.reports.run(self.ctx, "ai_usage")["rows"][0]
        self.assertEqual((row["calls"], row["tokens"]), (1, 100))
        self.assertIsNone(row["estimated_cost_usd"])

    def test_export_csv_and_xlsx(self) -> None:
        self.store.insert(self.ctx, "companies", {"name": "=HYPERLINK(1)", "lifecycle": "prospect"})
        data, name, content_type = self.reports.export(self.ctx, "prospect_conversion", "csv")
        text = data.decode("utf-8-sig")
        self.assertTrue(name.endswith(".csv"))
        self.assertIn("Lifecycle,Companies,Share", text)
        self.assertIn("100.0%", text)
        data, name, content_type = self.reports.export(self.ctx, "prospect_conversion", "xlsx")
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(data))
        self.assertEqual(book.sheetnames, ["prospect_conversion", "about"])
        with self.assertRaises(ValidationError):
            self.reports.export(self.ctx, "funnel", "pdf")
        self.assertTrue(self.store.count(self.ctx, "audit_log", {"action": "report.export"}) >= 2)

    def test_saved_views(self) -> None:
        row = self.reports.save_view(self.ctx, {"name": "Q3 campaigns", "report": "campaign_attribution",
                                                "date_range": {"start": "2026-07-01", "end": "2026-09-30"}})
        self.assertEqual(row["report"], "campaign_attribution")
        with self.assertRaises(ValidationError):
            self.reports.save_view(self.ctx, {"name": "x", "report": "bogus"})


class ScoringApiTests(unittest.TestCase):
    def setUp(self) -> None:
        from fastapi.testclient import TestClient

        from cloud.api.auth import DevTokenIssuer
        from cloud.api.main import create_app
        from cloud.api.settings import Settings
        from cloud.intel.platform import PlatformConfig
        from cloud.shared.storage import LocalFileStorage
        from cloud.worker.dispatcher import NullDispatcher

        scratch = tempfile.TemporaryDirectory()
        self.addCleanup(scratch.cleanup)
        root = Path(scratch.name)
        self.platform = Platform(MemoryStore(), storage=LocalFileStorage(root / "platform"),
                                 config=PlatformConfig(files_dir=root / "platform"))
        issuer = DevTokenIssuer("scoring-api-tests-secret-0123456789abcdef")
        app = create_app(Settings(auth_mode="dev", results_dir=root / "results"),
                         storage=LocalFileStorage(root / "results"), token_verifier=issuer,
                         dispatcher=NullDispatcher(), platform=self.platform)
        self.client = TestClient(app)
        self.client.headers["Authorization"] = f"Bearer {issuer.issue('alice@example.com')['access_token']}"
        self.client.__enter__()
        self.addCleanup(self.client.__exit__, None, None, None)
        self.ws = self.client.post("/api/v1/workspaces", json={"name": "Scores API", "seed": False}).json()["id"]
        self.base = f"/api/v1/w/{self.ws}"

    def test_score_and_report_endpoints(self) -> None:
        company = self.client.post(self.base + "/companies", json={"name": "Acme", "industry": "Manufacturing"}).json()
        r = self.client.get(f"{self.base}/scores/company/{company['id']}")
        self.assertEqual(r.status_code, 200, r.text)
        self.assertFalse(r.json()["persisted"])
        r = self.client.post(f"{self.base}/scores/company/{company['id']}")
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(len(self.client.get(f"{self.base}/scores/company/{company['id']}/history").json()["items"]), 5)
        self.assertEqual(self.client.get(self.base + "/scores/model").json()["model"], MODEL)
        self.assertEqual(self.client.get(self.base + "/scores/deal/x").status_code, 422)

        catalog = self.client.get(self.base + "/analytics/reports").json()["items"]
        self.assertGreaterEqual(len(catalog), 13)
        r = self.client.get(self.base + "/analytics/reports/prospect_conversion", params={"start": "2025-06-01"})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["totals"]["companies"], 1)
        self.assertEqual(self.client.get(self.base + "/analytics/reports/funnel",
                                         params={"bogus": "1"}).status_code, 422)
        r = self.client.get(self.base + "/analytics/reports/funnel/export", params={"format": "xlsx"})
        self.assertEqual(r.status_code, 200)
        self.assertIn("attachment", r.headers["content-disposition"])
        r = self.client.post(self.base + "/analytics/saved-reports", json={"name": "Mine", "report": "funnel"})
        self.assertEqual(r.status_code, 201, r.text)
        self.assertEqual(self.client.get(self.base + "/analytics/saved-reports").json()["total"], 1)
        # The existing dashboard route still works.
        self.assertEqual(self.client.get(self.base + "/analytics/dashboard").status_code, 200)
        r = self.client.post(self.base + "/scores/rescore", json={"all": True})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["kind"], "scoring")


if __name__ == "__main__":
    unittest.main()
