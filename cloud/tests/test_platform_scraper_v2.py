"""AI scraper V2-V4, offline.

V2: pagination (next / numbered / cursor / load-more link / load-more button), loop
protection, job-detail enrichment and precedence, browser fallback and its limits,
discovery, multi-URL concurrency, per-domain limits, retries and backoff, blocked
pages, checkpointing.
V3: natural-language schemas, custom fields and types, schema editing, validation,
normalisation, field statuses, conflicts, confidence, templates.
V4: queue and worker execution, pause/resume/cancel/retry/restart, crash recovery,
duplicate prevention, persistence, exports, CRM matching and PROPOSE → REVIEW → APPLY,
the research agent and the AI Control Room tools, observability, the HTTP API.

Pages come from a fake HTTP session, the browser from a fake renderer, AI from fakes.
"""

from __future__ import annotations

import csv
import io
import json
import logging
import tempfile
import threading
import time
import unittest
import uuid
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Dict, List

from cloud.intel.core.context import Ctx, ValidationError
from cloud.intel.core.http import SafeFetcher
from cloud.intel.platform import Platform, PlatformConfig
from cloud.intel.scraper.crawler import VisitedSet
from cloud.intel.scraper.fetcher import PageFetcher, RenderResult
from cloud.intel.scraper.limits import DomainLimiter, RunBudget
from cloud.intel.scraper.models import CrawlOptions, FieldValue, Outcome, merge_value
from cloud.intel.scraper.normalizer import normalize_phone, normalize_value, normalize_website
from cloud.intel.scraper.planner import instruction_to_schema
from cloud.intel.scraper.runner import AIBudget, execute_run, flatten, scrape_one
from cloud.intel.scraper.service import clean_schema
from cloud.intel.scraper.validator import apply_filters, check_value, validate_fields
from cloud.intel.store.memory import MemoryStore
from cloud.intel.tasks.worker import TaskPaused, run_task_inline
from cloud.shared.storage import LocalFileStorage
from cloud.tests.test_platform_ai_fakes import FakeResponse, FakeSession, fake_resolver
from cloud.tests.test_platform_scraper import ORG_PAGE, PAGES, FakeAI, values

TODAY = date.today().isoformat()


def listing(jobs, *, next_href=None, numbered=(), extra=""):
    items = "".join(f'<li><a href="/careers/{slug}"><h3>{title}</h3><span class="location">{loc}</span></a></li>'
                    for slug, title, loc in jobs)
    nav = ""
    if next_href:
        nav += f'<a rel="next" href="{next_href}">Next</a>'
    if numbered:
        nav = '<nav class="pagination">' + "".join(f'<a href="{h}">{n}</a>' for n, h in numbered) + nav + "</nav>"
    return f"<html><head><title>Careers at Pager Inc</title></head><body><ul>{items}</ul>{nav}{extra}</body></html>"


PAGINATED = {
    "https://pg.example/careers": (200, listing([("a1", "Analyst One", "Austin, TX"), ("a2", "Analyst Two", "Remote")],
                                                next_href="/careers?page=2",
                                                numbered=[(2, "/careers?page=2"), (3, "/careers?page=3")])),
    "https://pg.example/careers?page=2": (200, listing([("a3", "Analyst Three", "Tulsa, OK"),
                                                        ("a4", "Analyst Four", "Tulsa, OK")],
                                                       next_href="/careers?page=3")),
    "https://pg.example/careers?page=3": (200, listing([("a5", "Analyst Five", "Dallas, TX")])),
}

LOOP = {
    "https://loop.example/jobs": (200, listing([("l1", "Loop One", "X")], next_href="/jobs?p=2")),
    "https://loop.example/jobs?p=2": (200, listing([("l2", "Loop Two", "X")], next_href="/jobs")),
}

CURSOR = {
    "https://cur.example/jobs": (200, listing([("c1", "Cursor One", "X")], next_href="/jobs?cursor=abc123")),
    "https://cur.example/jobs?cursor=abc123": (200, listing([("c2", "Cursor Two", "X")])),
}

LOAD_MORE_LINK = {
    "https://lm.example/jobs": (200, listing([("m1", "More One", "X")],
                                             extra='<button data-url="/jobs?offset=20">Load more jobs</button>')),
    "https://lm.example/jobs?offset=20": (200, listing([("m2", "More Two", "X")])),
}

LOAD_MORE_BUTTON = {
    "https://btn.example/jobs": (200, listing([("b1", "Button One", "X")],
                                              extra='<button class="load-more">Show more</button>')),
}
BUTTON_RENDERED = listing([("b1", "Button One", "X"), ("b2", "Button Two", "X"), ("b3", "Button Three", "X")])

SHELL = ("<html><head><title>Careers</title><script src='/a.js'></script><script src='/b.js'></script></head>"
         "<body><div id='root'></div></body></html>")
SHELL_RENDERED = listing([("s1", "Rendered Engineer", "Denver, CO")])

DETAIL_LISTING = ("<html><head><title>Careers at Detail Co</title></head><body><ul>"
                  '<li><a href="/careers/erp-lead"><h3>ERP Lead</h3><span class="location">Austin, TX</span></a></li>'
                  '<li><a href="/careers/sap-analyst"><h3>SAP Analyst</h3></a></li></ul></body></html>')


def detail_page(title, location, salary_min, *, extra=""):
    ld = {"@type": "JobPosting", "title": title, "datePosted": TODAY, "employmentType": "FULL_TIME",
          "hiringOrganization": {"@type": "Organization", "name": "Detail Co"},
          "jobLocation": {"@type": "Place", "address": {"addressLocality": location[0], "addressRegion": location[1],
                                                        "addressCountry": "US"}},
          "baseSalary": {"@type": "MonetaryAmount", "currency": "USD",
                         "value": {"@type": "QuantitativeValue", "minValue": salary_min, "maxValue": salary_min + 20000,
                                   "unitText": "YEAR"}},
          "description": "<p>We need 5+ years of experience with SAP S/4HANA and Python. PMP preferred.</p>" + extra}
    return f"""<html><head><script type="application/ld+json">{json.dumps(ld)}</script></head>
<body><main><h1>{title}</h1><p>Hiring Manager: Dana Reyes</p></main></body></html>"""


DETAILS = {
    "https://det.example/careers": (200, DETAIL_LISTING),
    "https://det.example/careers/erp-lead": (200, detail_page("ERP Lead", ("Dallas", "TX"), 120000)),
    "https://det.example/careers/sap-analyst": (200, detail_page("SAP Analyst", ("Tulsa", "OK"), 90000)),
}


class SeqSession(FakeSession):
    """A URL may map to a list of responses, served in order (the last one repeats);
    an Exception in the list is raised instead of answering."""

    def request(self, method, url, **kwargs):
        entry = self.pages.get(url)
        if isinstance(entry, list):
            self.calls.append(url)
            item = entry.pop(0) if len(entry) > 1 else entry[0]
            if isinstance(item, BaseException):
                raise item
            status, body, *rest = item
            return FakeResponse(status, body if isinstance(body, str) else json.dumps(body), rest[0] if rest else None)
        return super().request(method, url, **kwargs)


def http_for(pages, session_cls=SeqSession):
    return SafeFetcher(session=session_cls(pages), resolver=fake_resolver(), per_host_delay=0)


class FakeRenderer:
    def __init__(self, pages: Dict[str, str]) -> None:
        self.pages, self.calls = pages, []

    def render(self, url, *, interact=False, max_clicks=10):
        self.calls.append((url, interact))
        html = self.pages.get(url)
        return RenderResult(html=html or "", final_url=url, status=200 if html else 0, duration_ms=12.5,
                            clicks=1 if interact else 0, error=None if html else "nothing rendered")


def fetcher(pages, *, renderer=None, sleep=None, **options):
    opts = CrawlOptions.from_mapping(options)
    budget = RunBudget(max_runtime_s=opts.max_runtime_s, max_records=opts.max_records,
                       max_browser_pages=opts.max_browser_pages)
    f = PageFetcher(http_for(pages), renderer, options=opts, budget=budget, sleep=sleep or (lambda s: None),
                    limiter=DomainLimiter(concurrency=opts.domain_concurrency,
                                          max_requests=opts.max_requests_per_domain, sleep=lambda s: None))
    return f, opts, budget


def crawl(url, instruction, pages, *, renderer=None, ai=None, **options):
    f, opts, budget = fetcher(pages, renderer=renderer, **options)
    page = scrape_one(url, instruction_to_schema(instruction), f, ai, options=opts, budget=budget,
                      visited=VisitedSet())
    return page, f


# =====================================================================================================
# V2
# =====================================================================================================


class PaginationTests(unittest.TestCase):
    def test_next_and_numbered_pagination_collect_every_page_once(self) -> None:
        page, f = crawl("https://pg.example/careers", "Find all job titles and job URLs", PAGINATED)
        self.assertEqual([r["job_title"] for r in values(page)],
                         ["Analyst One", "Analyst Two", "Analyst Three", "Analyst Four", "Analyst Five"])
        fetched = [p for p in page.pages if p["outcome"] == Outcome.OK]
        self.assertEqual([p["url"] for p in fetched], ["https://pg.example/careers", "https://pg.example/careers?page=2",
                                                       "https://pg.example/careers?page=3"])
        self.assertEqual([p["kind"] for p in fetched], ["input", "listing", "listing"])
        self.assertEqual(f.http._session.calls.count("https://pg.example/careers?page=3"), 1, "never twice")
        skipped = [p for p in page.pages if p["outcome"] == Outcome.SKIPPED]
        self.assertTrue(all(p["error"] == "already visited in this run" for p in skipped))

    def test_pagination_can_be_turned_off_and_is_bounded(self) -> None:
        page, _ = crawl("https://pg.example/careers", "Find all job titles", PAGINATED, pagination=False)
        self.assertEqual(len(page.records), 2)
        page, _ = crawl("https://pg.example/careers", "Find all job titles", PAGINATED, max_pages=2)
        self.assertEqual(len(page.records), 4)
        self.assertTrue(any("page limit" in p for p in page.problems))
        self.assertIn(Outcome.LIMIT, {p["outcome"] for p in page.pages})

    def test_an_official_api_board_is_not_paged_again(self) -> None:
        board = listing([("x1", "HTML Job", "X")], next_href="/acmemfg?page=2")
        pages = {**PAGES, "https://boards.greenhouse.io/acmemfg": (200, board)}
        page, f = crawl("https://boards.greenhouse.io/acmemfg", "Find all job titles", pages)
        self.assertEqual([r["job_title"] for r in values(page)], ["ERP Analyst"])
        self.assertNotIn("https://boards.greenhouse.io/acmemfg?page=2", f.http._session.calls)

    def test_listing_page_limit_leaves_room_for_details(self) -> None:
        page, _ = crawl("https://pg.example/careers", "Find all job titles", PAGINATED, max_listing_pages=1)
        self.assertEqual(len(page.records), 4)
        self.assertTrue(any("listing page limit (1)" in p for p in page.problems))

    def test_a_pagination_loop_ends(self) -> None:
        page, f = crawl("https://loop.example/jobs", "Find all job titles", LOOP)
        self.assertEqual([r["job_title"] for r in values(page)], ["Loop One", "Loop Two"])
        self.assertEqual(len(f.http._session.calls), 3)   # robots.txt + two pages

    def test_cursor_and_load_more_link(self) -> None:
        page, _ = crawl("https://cur.example/jobs", "Find all job titles", CURSOR)
        self.assertEqual([r["job_title"] for r in values(page)], ["Cursor One", "Cursor Two"])
        page, _ = crawl("https://lm.example/jobs", "Find all job titles", LOAD_MORE_LINK)
        self.assertEqual([r["job_title"] for r in values(page)], ["More One", "More Two"])

    def test_load_more_button_needs_the_browser(self) -> None:
        page, _ = crawl("https://btn.example/jobs", "Find all job titles", LOAD_MORE_BUTTON)
        self.assertEqual(len(page.records), 1)
        self.assertTrue(any("enable browser rendering" in p for p in page.problems))
        renderer = FakeRenderer({"https://btn.example/jobs": BUTTON_RENDERED})
        page, _ = crawl("https://btn.example/jobs", "Find all job titles", LOAD_MORE_BUTTON, renderer=renderer,
                        browser=True)
        self.assertEqual([r["job_title"] for r in values(page)], ["Button One", "Button Two", "Button Three"])
        self.assertEqual(renderer.calls, [("https://btn.example/jobs", True)])
        rendered = [p for p in page.pages if p["browser_used"]]
        self.assertEqual(len(rendered), 1)
        self.assertIn("load-more button", rendered[0]["browser_reason"])
        self.assertEqual(rendered[0]["browser_outcome"], Outcome.OK)
        self.assertEqual(rendered[0]["browser_duration_ms"], 12.5)
        self.assertTrue(page.records[1]["job_title"].browser, "browser values are marked")


class DetailTests(unittest.TestCase):
    def test_detail_pages_enrich_and_merge_by_precedence(self) -> None:
        instruction = ("Find job titles, job URLs, location, salary, skills, years of experience, certifications, "
                       "technology, seniority, employment type, posted date, hiring manager and description")
        page, _ = crawl("https://det.example/careers", instruction, DETAILS, follow_details=True)
        rows = values(page)
        self.assertEqual([r["job_title"] for r in rows], ["ERP Lead", "SAP Analyst"])
        lead = page.records[0]
        self.assertEqual(lead["location"].value, "Dallas, TX, US", "structured data beats the listing card")
        self.assertEqual(lead["location"].method, "json-ld")
        self.assertEqual([a.value for a in lead["location"].alternatives], ["Austin, TX"], "the conflict is kept")
        self.assertEqual(lead["location"].source_url, "https://det.example/careers/erp-lead")
        self.assertEqual(lead["job_url"].source_url, "https://det.example/careers", "each value keeps its source")
        self.assertIn("120000", str(lead["salary"].value))
        self.assertIn("SAP S/4HANA", lead["technology"].value)
        self.assertEqual(lead["years_experience"].value, 5)
        self.assertIn("PMP", lead["certifications"].value)
        self.assertEqual(lead["hiring_manager"].value, "Dana Reyes")
        self.assertEqual([p["kind"] for p in page.pages if p["outcome"] == "OK"], ["input", "detail", "detail"])
        row = flatten(lead, instruction_to_schema(instruction), source_url="x", extracted_at=datetime.now(timezone.utc))
        self.assertEqual(row["_field_status"]["location"], "conflict")
        self.assertEqual(row["_conflicts"]["location"]["alternatives"][0]["value"], "Austin, TX")
        self.assertEqual(row["seniority"], "Lead")
        self.assertEqual(row["_field_status"]["seniority"], "inferred")
        self.assertEqual(row["employment_type"], "Full-time")

    def test_detail_limit_and_weaker_evidence_never_wins(self) -> None:
        page, _ = crawl("https://det.example/careers", "Find job titles and salary", DETAILS, follow_details=True,
                        max_detail_pages=1)
        self.assertTrue(any("detail page limit (1)" in p for p in page.problems))
        strong = FieldValue("ERP Analyst", "ats-api", 0.95, "api", "https://api")
        weak = FieldValue("ERP Analyst II", "ai", 0.6, "quote", "https://page")
        merged = merge_value(strong, weak)
        self.assertEqual((merged.value, merged.method), ("ERP Analyst", "ats-api"))
        self.assertEqual(merged.alternatives[0].value, "ERP Analyst II")
        merged = merge_value(FieldValue("x", "heading", 0.7, browser=True), FieldValue("y", "regex", 0.5))
        self.assertEqual(merged.value, "y", "HTTP-deterministic beats browser-rendered")


class BrowserFallbackTests(unittest.TestCase):
    def test_a_javascript_shell_is_rendered_only_when_enabled(self) -> None:
        pages = {"https://js.example/careers": (200, SHELL)}
        page, _ = crawl("https://js.example/careers", "Find all job titles", pages)
        self.assertEqual(page.records, [{}] if page.records else [])
        self.assertIn("no job postings found on the page", page.problems)
        renderer = FakeRenderer({"https://js.example/careers": SHELL_RENDERED})
        page, _ = crawl("https://js.example/careers", "Find all job titles", pages, renderer=renderer, browser=True)
        self.assertEqual([r["job_title"] for r in values(page)], ["Rendered Engineer"])
        visit = page.pages[0]
        self.assertTrue(visit["browser_used"])
        self.assertEqual(visit["browser_reason"], "the page is an empty JavaScript shell")

    def test_browser_limits_and_challenges(self) -> None:
        pages = {"https://js.example/careers": (200, SHELL), "https://js2.example/careers": (200, SHELL)}
        renderer = FakeRenderer({"https://js.example/careers": SHELL_RENDERED,
                                 "https://js2.example/careers": SHELL_RENDERED})
        f, opts, budget = fetcher(pages, renderer=renderer, browser=True, max_browser_pages=1)
        schema = instruction_to_schema("Find all job titles")
        first = scrape_one("https://js.example/careers", schema, f, options=opts, budget=budget)
        second = scrape_one("https://js2.example/careers", schema, f, options=opts, budget=budget)
        self.assertEqual(len(first.records), 1)
        self.assertEqual(len(renderer.calls), 1, "the run's browser budget is 1 page")
        self.assertFalse(second.pages[0]["browser_used"])
        challenge = FakeRenderer({"https://js.example/careers": "<html><title>Just a moment...</title>"
                                                               "<script src='/cdn-cgi/challenge-platform/x'></script></html>"})
        page, _ = crawl("https://js.example/careers", "Find all job titles", pages, renderer=challenge, browser=True)
        self.assertEqual(page.outcome, Outcome.CAPTCHA)
        self.assertIn("not bypassed", page.problems[0])


class DiscoveryTests(unittest.TestCase):
    def test_homepage_without_links_tries_bounded_common_paths(self) -> None:
        home = "<html><head><title>Omega Tools</title></head><body><p>Welcome</p></body></html>"
        pages = {"https://omega.example/": (200, home),
                 "https://omega.example/careers": (200, listing([("o1", "Omega Machinist", "Ohio")]))}
        page, f = crawl("https://omega.example/", "Get company name and job titles", pages)
        self.assertEqual([r["job_title"] for r in values(page)], ["Omega Machinist"])
        self.assertEqual(page.pages[1]["kind"], "discovery")
        self.assertLessEqual(len([u for u in f.http._session.calls if "omega" in u]), 5)

    def test_candidates_are_ranked(self) -> None:
        from cloud.intel.scraper.discovery import rank_candidates

        ranked = rank_candidates("https://acme.example/", [("About", "https://acme.example/about"),
                                                           ("Jobs", "https://acme.example/join"),
                                                           ("x", "https://boards.greenhouse.io/acme")])
        self.assertEqual(ranked[0][0], "https://job-boards.greenhouse.io/acme")
        self.assertEqual(ranked[1][0], "https://acme.example/join")


class LimiterRetryTests(unittest.TestCase):
    def test_transient_failures_are_retried_with_backoff(self) -> None:
        slept: List[float] = []
        pages = {"https://flaky.example/": [(503, "busy"), (503, "busy"), (200, ORG_PAGE)]}
        f, *_ = fetcher(pages, sleep=slept.append, max_retries=2)
        page = f.fetch("https://flaky.example/")
        self.assertEqual((page.outcome, page.attempts), (Outcome.OK, 3))
        self.assertEqual(slept, [1.0, 2.0])

    def test_retry_after_is_honoured_and_bounded(self) -> None:
        slept: List[float] = []
        pages = {"https://slow.example/": [(429, "slow", {"Retry-After": "5"}), (200, ORG_PAGE)],
                 "https://slower.example/": [(429, "slow", {"Retry-After": "3600"}), (200, ORG_PAGE)]}
        f, *_ = fetcher(pages, sleep=slept.append, max_retries=2, max_backoff_s=30)
        self.assertEqual(f.fetch("https://slow.example/").outcome, Outcome.OK)
        self.assertEqual(slept, [5.0])
        page = f.fetch("https://slower.example/")
        self.assertEqual(page.outcome, Outcome.RATE_LIMITED, "an hour's wait is not waited out")
        self.assertEqual(page.attempts, 1)

    def test_refusals_and_timeouts(self) -> None:
        import requests

        slept: List[float] = []
        pages = {"https://no.example/": [(403, "no"), (200, ORG_PAGE)], "https://gone.example/": [(404, "x")],
                 "https://to.example/": [requests.exceptions.ReadTimeout("read timed out"), (200, ORG_PAGE)]}
        f, *_ = fetcher(pages, sleep=slept.append, max_retries=2)
        self.assertEqual(f.fetch("https://no.example/").attempts, 1, "403 is never retried")
        self.assertEqual(f.fetch("https://gone.example/").outcome, Outcome.NOT_FOUND)
        timed = f.fetch("https://to.example/")
        self.assertEqual((timed.outcome, timed.attempts), (Outcome.OK, 2))

    def test_per_domain_request_budget(self) -> None:
        page, _ = crawl("https://pg.example/careers", "Find all job titles", PAGINATED, max_requests_per_domain=2)
        # pages 1 and 2 use the domain's budget; page 3 is reported, not fetched
        limited = [p for p in page.pages if p["outcome"] == Outcome.LIMIT]
        self.assertEqual([p["url"] for p in limited], ["https://pg.example/careers?page=3"])
        self.assertIn("per-domain request limit", limited[0]["error"])
        self.assertEqual(len(page.records), 4)

    def test_domain_concurrency_is_bounded(self) -> None:
        limiter = DomainLimiter(concurrency=2, max_requests=100)
        active, peak, lock = [0], [0], threading.Lock()

        def hit():
            limiter.acquire("https://a.example/x")
            with lock:
                active[0] += 1
                peak[0] = max(peak[0], active[0])
            time.sleep(0.02)
            with lock:
                active[0] -= 1
            limiter.release("https://a.example/x")

        threads = [threading.Thread(target=hit) for _ in range(8)]
        for t in threads:
            t.start()
        for t in threads:
            t.join()
        self.assertEqual(peak[0], 2)
        self.assertEqual(limiter.requests("a.example"), 8)


# =====================================================================================================
# V3
# =====================================================================================================


class SchemaV3Tests(unittest.TestCase):
    def test_the_natural_language_example(self) -> None:
        schema = instruction_to_schema("Find US manufacturing companies using SAP, include company name, website, "
                                       "ERP, industry, employee count, hiring manager, and all open SAP jobs posted "
                                       "in the last 14 days.")
        by = {f["name"]: f for f in schema["fields"]}
        self.assertEqual(schema["entities"], ["company", "job"])
        self.assertEqual(set(by), {"company_name", "website", "industry", "employee_count", "erp", "hiring_manager",
                                   "job_title", "posted_date"})
        self.assertEqual((by["employee_count"]["type"], by["posted_date"]["type"], by["website"]["type"]),
                         ("integer", "date", "url"))
        self.assertEqual((by["website"]["normalize"], by["company_name"]["level"], by["job_title"]["level"]),
                         ("website", "company", "job"))
        self.assertTrue(by["job_title"]["required"])
        self.assertIn({"field": "posted_date", "op": "within_days", "value": 14, "mode": "hard"}, schema["filters"])
        self.assertIn({"field": "job_title", "op": "contains_any", "value": ["SAP"], "mode": "hard"}, schema["filters"])
        self.assertEqual(schema["criteria"], {"country": "US", "industries": ["manufacturing"], "technologies": ["SAP"]})

    def test_custom_fields_and_their_types(self) -> None:
        schema = instruction_to_schema("Get company name, Uses SAP?, implementation partner, cloud provider, "
                                       "decision maker, job family, number of plants, next audit date, "
                                       "investor relations page")
        by = {f["name"]: f for f in schema["fields"]}
        self.assertEqual(by["uses_sap"]["type"], "boolean")
        self.assertEqual(by["uses_sap"]["hint"], "sap")
        self.assertEqual(by["implementation_partner"]["source"], "custom")
        self.assertEqual(by["number_plants"]["type"], "integer")
        self.assertEqual(by["next_audit_date"]["type"], "date")
        self.assertEqual(by["investor_relations"]["type"], "url", "'page' in the phrase makes it a URL field")
        for standard in ("cloud_provider", "decision_maker", "job_family"):
            self.assertEqual(by[standard]["source"], "rules")

    def test_schema_editing(self) -> None:
        schema = instruction_to_schema("Get company name, website and industry")
        edited = {**schema, "fields": [
            {**schema["fields"][2], "name": "Sector", "label": "Sector"},          # renamed, moved first
            {**schema["fields"][0], "required": True},
            {"name": "tier", "type": "enum", "enum": ["A", "B", "C"], "level": "company"},   # added
            {"name": "ticker", "type": "string", "pattern": "[A-Z]{1,5}", "max_length": 5},
        ]}                                                                          # website removed
        clean = clean_schema(edited)
        self.assertEqual([f["name"] for f in clean["fields"]], ["sector", "company_name", "tier", "ticker"])
        self.assertEqual(clean["fields"][2]["enum"], ["A", "B", "C"])
        with self.assertRaises(ValidationError):
            clean_schema({"fields": [{"name": "x", "type": "enum"}]})
        with self.assertRaises(ValidationError):
            clean_schema({"fields": [{"name": "x", "pattern": "(["}]})
        with self.assertRaises(ValidationError):
            clean_schema({"fields": []})

    def test_ai_can_only_describe_phrases_the_user_wrote(self) -> None:
        ai = FakeAI({"fields": [{"phrase": "implementation partner", "name": "impl_partner", "type": "string",
                                 "level": "company", "description": "SI partner", "hint": "look for 'partner'"},
                                {"phrase": "ceo salary", "name": "ceo_salary", "type": "decimal", "level": "company",
                                 "description": "invented"}]})
        schema = instruction_to_schema("Get company name and implementation partner", ai=ai)
        names = [f["name"] for f in schema["fields"]]
        self.assertEqual(names, ["company_name", "impl_partner"])
        self.assertEqual(schema["fields"][1]["hint"], "look for 'partner'")


class ValidationNormalizationTests(unittest.TestCase):
    def test_types(self) -> None:
        cases = [({"name": "p", "type": "phone"}, "+19185550142", "+19185550142"),
                 ({"name": "p", "type": "phone"}, "call us", None),
                 ({"name": "n", "type": "integer"}, "1,200", 1200), ({"name": "n", "type": "integer"}, "1.5", None),
                 ({"name": "d", "type": "decimal"}, "$1,234.50", 1234.5),
                 ({"name": "e", "type": "enum", "enum": ["Remote", "Hybrid"]}, "Remote", "Remote"),
                 ({"name": "e", "type": "enum", "enum": ["Remote", "Hybrid"]}, "Moon", None),
                 ({"name": "s", "type": "string", "max_length": 3}, "abcd", None),
                 ({"name": "s", "type": "string", "pattern": "[A-Z]+"}, "ACME", "ACME"),
                 ({"name": "s", "type": "string", "pattern": "[A-Z]+"}, "acme", None),
                 ({"name": "t", "type": "datetime"}, "2026-09-03T10:00:00+00:00", "2026-09-03T10:00:00+00:00"),
                 ({"name": "b", "type": "boolean"}, "yes", None), ({"name": "a", "type": "array"}, "x", ["x"]),
                 ({"name": "u", "type": "url"}, "ftp://x", None), ({"name": "m", "type": "email"}, "A@B.com", "a@b.com")]
        for spec, value, expected in cases:
            clean, error = check_value(spec, value)
            self.assertEqual(clean, expected, (spec, value))
            self.assertEqual(error is None, expected is not None, (spec, value, error))

    def test_statuses(self) -> None:
        fields = [{"name": "a", "type": "string"}, {"name": "b", "type": "url"}, {"name": "c", "type": "string",
                                                                                  "required": True},
                  {"name": "d", "type": "string"}, {"name": "e", "type": "string"}]
        _clean, statuses, errors = validate_fields({"a": "x", "b": "not a url", "d": "y", "e": "z"}, fields,
                                                   methods={"d": "ai"}, conflicts=["e"])
        self.assertEqual(statuses, {"a": "valid", "b": "invalid", "c": "missing", "d": "inferred", "e": "conflict"})
        self.assertEqual(set(errors), {"b", "c"})

    def test_invalid_ai_values_are_rejected_with_evidence(self) -> None:
        schema = {"entity": "company", "fields": [{"name": "employee_count", "type": "integer"},
                                                  {"name": "company_name", "type": "string"}]}
        record = {"employee_count": FieldValue("about a thousand", "ai", 0.45, "about a thousand people", "https://x"),
                  "company_name": FieldValue("Acme", "json-ld", 0.95, "Organization.name", "https://x")}
        row = flatten(record, schema, source_url="https://x", extracted_at=datetime.now(timezone.utc))
        self.assertIsNone(row["employee_count"], "never replaced by a guess")
        self.assertEqual(row["_field_status"]["employee_count"], "invalid")
        self.assertEqual(row["_rejected"]["employee_count"]["value"], "about a thousand")
        self.assertEqual(row["_rejected"]["employee_count"]["evidence"], "about a thousand people")
        self.assertEqual(row["confidence"], 0.95)

    def test_normalization(self) -> None:
        self.assertEqual(normalize_website("https://www.example.com/"), "https://example.com")
        self.assertEqual(normalize_website("www.example.com/about?x=1"), "https://example.com")
        self.assertEqual(normalize_phone("(918) 555-0142"), "+19185550142")
        self.assertEqual(normalize_phone("+44 20 7946 0958"), "+442079460958")
        self.assertEqual(normalize_value("location", "string", "Tulsa, OK, United States", rule="location"), "Tulsa, OK, US")
        self.assertEqual(normalize_value("employee_count", "integer", "1.2k employees"), 1200)
        self.assertEqual(normalize_value("employee_count", "integer", "501-1,000"), 501)
        self.assertEqual(normalize_value("revenue", "decimal", "$1.2B"), 1.2e9)
        self.assertEqual(normalize_value("job_title", "string", "  Senior   Engineer ", rule="job_title"),
                         "Senior Engineer")
        self.assertEqual(normalize_value("remote_mode", "enum", "onsite", enum=["Remote", "Hybrid", "On-site"]),
                         "On-site")

    def test_keyword_boolean_and_soft_filters(self) -> None:
        pages = {"https://kw.example/": (200, "<html><head><title>KW Corp</title></head><body><p>We run SAP S/4HANA "
                                              "across all plants.</p></body></html>"),
                 "https://nokw.example/": (200, "<html><head><title>No KW</title></head><body>hello</body></html>")}
        page, _ = crawl("https://kw.example/", "Get company name and Uses SAP?", pages)
        self.assertIs(values(page)[0]["uses_sap"], True)
        self.assertIn("SAP", page.records[0]["uses_sap"].evidence)
        page, _ = crawl("https://nokw.example/", "Get company name and Uses SAP?", pages)
        self.assertNotIn("uses_sap", values(page)[0], "absence is not evidence of 'no'")
        rows = [{"country": "US"}, {"country": "CA"}, {"country": None}]
        kept, dropped = apply_filters(rows, [{"field": "country", "op": "equals", "value": "US", "mode": "soft"}])
        self.assertEqual((len(kept), dropped), (3, 0))
        self.assertEqual([r["_filters"]["country equals"] for r in kept], ["pass", "fail", "unknown"])


# =====================================================================================================
# V4
# =====================================================================================================


class Base(unittest.TestCase):
    def setUp(self) -> None:
        from cloud.shared.queue import InMemoryJobQueue

        scratch = tempfile.TemporaryDirectory()
        self.addCleanup(scratch.cleanup)
        self.root = Path(scratch.name)
        self.store = MemoryStore()
        self.user = str(uuid.uuid4())
        ws = self.store.create_workspace(self.user, "W", f"w-{uuid.uuid4().hex[:8]}")
        self.ctx = Ctx(ws["id"], self.user, "owner")
        self.pages: Dict[str, Any] = {**PAGES, **PAGINATED, **DETAILS}
        self.sessions: List[FakeSession] = []
        self.renderer = FakeRenderer({"https://btn.example/jobs": BUTTON_RENDERED})
        self.queue = InMemoryJobQueue()
        self.platform = self.make_platform()
        self.service = self.platform.service("scraper")

    def make_platform(self) -> Platform:
        def factory():
            http = http_for(self.pages)
            self.sessions.append(http._session)
            return http

        return Platform(self.store, storage=LocalFileStorage(self.root), queue=self.queue,
                        config=PlatformConfig(extra={"fetcher_factory": factory,
                                                     "renderer_factory": lambda: self.renderer}))

    def calls(self) -> List[str]:
        return [c for s in self.sessions for c in s.calls]

    def run_task(self, run):
        return run_task_inline(self.platform, self.ctx.workspace_id, run["task_id"])

    def read(self, run, key) -> bytes:
        run = self.store.get(self.ctx, "scrape_runs", run["id"])
        with self.platform.storage.open(run["stats"]["files"][key]["storage_key"]) as handle:
            return handle.read()


class _Reporter:
    def __init__(self, pause_after: int = 10 ** 6, cancel_after: int = 10 ** 6) -> None:
        self.checks, self.pause_after, self.cancel_after = 0, pause_after, cancel_after
        self.checkpoint: Dict[str, Any] = {}

    def progress(self, *a, **k):
        pass

    def is_cancelled(self):
        self.checks += 1
        return self.checks > self.cancel_after

    def should_pause(self):
        return self.checks > self.pause_after


class QueueWorkerTests(Base):
    def test_a_run_is_a_queued_task_the_worker_executes(self) -> None:
        from cloud.intel.tasks.worker import PlatformWorker

        run = self.service.start(self.ctx, ["https://pg.example/careers"], "Find all job titles and job URLs")
        task = self.platform.tasks.get(self.ctx, run["task_id"])
        self.assertEqual((task["kind"], task["status"]), ("scraper", "queued"))
        self.assertEqual(run["status"], "queued")
        self.assertTrue(PlatformWorker(self.platform, worker_id="w1").process_next())
        run = self.service.get(self.ctx, run["id"])
        self.assertEqual(run["status"], "completed")
        self.assertEqual(run["stats"]["records"], 5)
        self.assertEqual(self.platform.tasks.get(self.ctx, run["task_id"])["status"], "completed")

    def test_one_worker_runs_several_runs_at_once(self) -> None:
        from cloud.intel.tasks.worker import PlatformWorker

        runs = [self.service.start(self.ctx, [url], "Find all job titles")
                for url in ("https://pg.example/careers", "https://det.example/careers")]
        worker = PlatformWorker(self.platform, worker_id="w2", poll_seconds=0.02, maintenance_seconds=3600)
        thread = threading.Thread(target=worker.run, kwargs={"concurrency": 2}, daemon=True)
        thread.start()
        deadline = time.monotonic() + 30
        while time.monotonic() < deadline:
            if all(self.service.get(self.ctx, r["id"])["status"] == "completed" for r in runs):
                break
            time.sleep(0.05)
        worker.stop()
        thread.join(timeout=10)
        self.assertEqual([self.service.get(self.ctx, r["id"])["status"] for r in runs], ["completed", "completed"])

    def test_high_volume_runs_need_confirmation(self) -> None:
        urls = [f"https://site{n}.example/" for n in range(60)]
        plan = self.service.plan(self.ctx, urls, "Get company name")
        self.assertTrue(plan["requires_confirmation"])
        with self.assertRaises(ValidationError):
            self.service.start(self.ctx, urls, "Get company name")
        run = self.service.start(self.ctx, urls, "Get company name", confirm=True)
        self.assertEqual(run["stats"]["url_count"], 60)

    def test_plan_shows_sources_limits_browser_and_cost(self) -> None:
        plan = self.service.plan(self.ctx, "https://boards.greenhouse.io/acmemfg\nhttps://beta.example/\n"
                                            "https://det.example/careers", "Extract company name, website, job title "
                                                                           "and job URL.",
                                 options={"max_pages": 10, "follow_details": True, "browser": True})
        kinds = [s["kind"] for s in plan["sources"]]
        self.assertEqual(kinds, ["ats_board", "homepage", "careers_page"])
        self.assertTrue(plan["sources"][0]["official_api"])
        self.assertEqual(plan["limits"]["max_pages"], 10)
        self.assertTrue(plan["browser"]["requested"])
        self.assertEqual(plan["estimate"]["estimated_cost_usd"], 0.0)
        self.assertEqual(plan["estimate"]["ai_calls_max"], 0, "no AI provider in this workspace")
        self.assertEqual(self.store.count(self.ctx, "scrape_runs"), 0, "planning saves nothing")


class LifecycleTests(Base):
    def test_pause_and_resume_a_queued_run(self) -> None:
        run = self.service.start(self.ctx, ["https://pg.example/careers"], "Find all job titles")
        run = self.service.pause(self.ctx, run["id"])
        self.assertEqual(run["status"], "paused")
        self.assertEqual(self.platform.tasks.get(self.ctx, run["task_id"])["status"], "paused")
        run = self.service.resume(self.ctx, run["id"])
        self.assertEqual(run["status"], "queued")
        self.assertEqual(self.run_task(run)["status"], "completed")
        self.assertEqual(self.service.get(self.ctx, run["id"])["stats"]["records"], 5)

    def test_pause_mid_run_then_resume_without_refetching(self) -> None:
        urls = ["https://pg.example/careers", "https://det.example/careers", "https://beta.example/"]
        run = self.service.start(self.ctx, urls, "Find all job titles", options={"concurrency": 1})
        with self.assertRaises(TaskPaused):
            execute_run(self.platform, self.ctx, run["id"], _Reporter(pause_after=1),
                        task={"params": {"run_id": run["id"]}, "attempts": 1})
        run = self.service.get(self.ctx, run["id"])
        self.assertEqual(run["status"], "paused")
        saved = self.store.all(self.ctx, "scrape_results", {"run_id": run["id"]})
        self.assertEqual(len(saved), 1, "the first URL was saved before the pause")
        before = self.calls().count("https://pg.example/careers")
        execute_run(self.platform, self.ctx, run["id"], _Reporter(), task={"params": {"run_id": run["id"]},
                                                                            "attempts": 1})
        run = self.service.get(self.ctx, run["id"])
        self.assertEqual(run["status"], "completed")
        self.assertEqual(self.calls().count("https://pg.example/careers"), before, "saved pages are not refetched")
        self.assertEqual(len(self.store.all(self.ctx, "scrape_results", {"run_id": run["id"]})), 3)
        self.assertTrue(run["stats"]["recoveries"])

    def test_cancel_retry_restart(self) -> None:
        self.pages["https://down.example/"] = [(500, "x")]
        run = self.service.start(self.ctx, ["https://down.example/", "https://beta.example/"], "Get company name",
                                 options={"max_retries": 0})
        self.run_task(run)
        run = self.service.get(self.ctx, run["id"])
        self.assertEqual(run["stats"]["outcomes"], {"FAILED": 1, "OK": 1})
        self.pages["https://down.example/"] = [(200, "<html><head><title>Down Inc</title></head></html>")]
        run = self.service.retry(self.ctx, run["id"])
        self.run_task(run)
        self.assertEqual(self.service.get(self.ctx, run["id"])["stats"]["outcomes"], {"OK": 2})
        again = self.service.restart(self.ctx, run["id"])
        self.assertNotEqual(again["id"], run["id"])
        self.assertEqual(again["stats"]["restarted_from"], run["id"])
        self.assertEqual([i["url"] for i in again["stats"]["inputs"]], ["https://down.example/", "https://beta.example/"])
        self.service.cancel(self.ctx, again["id"])
        self.assertEqual(self.service.get(self.ctx, again["id"])["status"], "cancelled")
        self.assertEqual(self.service.get(self.ctx, run["id"])["status"], "completed", "the old run is kept")


class CrashRecoveryTests(Base):
    def test_a_crashed_run_resumes_from_its_checkpoint_without_duplicates(self) -> None:
        class Crash(BaseException):
            pass

        self.pages["https://det.example/careers"] = [Crash("worker killed"), (200, DETAIL_LISTING)]
        run = self.service.start(self.ctx, ["https://pg.example/careers", "https://det.example/careers"],
                                 "Find all job titles", options={"concurrency": 1})
        with self.assertRaises(Crash):
            execute_run(self.platform, self.ctx, run["id"], _Reporter(), task={"params": {"run_id": run["id"]},
                                                                                "attempts": 1})
        crashed = self.service.get(self.ctx, run["id"])
        self.assertIn(crashed["status"], ("fetching", "saving", "extracting", "paginating"))
        self.assertEqual(len(self.store.all(self.ctx, "scrape_results", {"run_id": run["id"]})), 1)
        pg_calls = self.calls().count("https://pg.example/careers")
        # The queue re-delivers the task (attempt 2).
        task = self.run_task(run)
        self.assertEqual(task["status"], "completed")
        run = self.service.get(self.ctx, run["id"])
        self.assertEqual(run["stats"]["recoveries"][-1]["inputs_already_saved"], 1)
        self.assertEqual(self.calls().count("https://pg.example/careers"), pg_calls)
        self.assertEqual(run["stats"]["records"], 7)
        pages = self.store.all(self.ctx, "scrape_pages", {"run_id": run["id"]})
        self.assertEqual(len(pages), len({p["url_key"] for p in pages}), "no page row twice")

    def test_maintenance_marks_runs_whose_task_died(self) -> None:
        run = self.service.start(self.ctx, ["https://beta.example/"], "Get company name")
        self.store.update(self.ctx, "scrape_runs", run["id"], {"status": "fetching"})
        system = self.ctx.as_system()
        task = self.store.get(system, "platform_tasks", run["task_id"])
        self.store.update(system, "platform_tasks", task["id"], {"status": "failed", "error": "lease expired"})
        self.assertEqual(self.service.recover(self.ctx), 1)
        run = self.service.get(self.ctx, run["id"])
        self.assertEqual(run["status"], "failed")
        self.assertIn("retry to resume", run["error"])

    def test_results_survive_a_restart_of_the_api(self) -> None:
        run = self.service.start(self.ctx, ["https://pg.example/careers"], "Find all job titles")
        self.run_task(run)
        fresh = self.make_platform().service("scraper")   # a new process on the same database and storage
        self.assertEqual(fresh.records(self.ctx, run["id"], "jobs")["total"], 5)
        self.assertEqual(fresh.pages(self.ctx, run["id"])["total"], 3)


class ExportDedupeTests(Base):
    def test_exports_and_observability(self) -> None:
        run = self.service.start(self.ctx, ["https://pg.example/careers", "https://walled.example/",
                                            "https://det.example/careers"], "Find all job titles and job URLs",
                                 options={"follow_details": True})
        self.run_task(run)
        run = self.service.get(self.ctx, run["id"])
        obs = run["stats"]["observability"]
        for key in ("total_urls", "completed_urls", "failed_urls", "blocked_urls", "pages_visited", "records_found",
                    "companies_found", "jobs_found", "browser_pages", "ai_calls", "ai_failures", "validation_errors",
                    "duplicates_removed", "duration_seconds"):
            self.assertIn(key, obs)
        self.assertEqual((obs["total_urls"], obs["completed_urls"], obs["blocked_urls"]), (3, 2, 1))
        self.assertEqual(obs["jobs_found"], 7)
        self.assertEqual(obs["pages_visited"], 12)   # 3 + 5 (404) details | 1 blocked | 1 + 2 details
        ndjson = self.read(run, "ndjson").decode().splitlines()
        self.assertEqual(len(ndjson), 7)
        self.assertIn("_evidence", json.loads(ndjson[0]))
        payload = json.loads(self.read(run, "json"))
        for key in ("schema", "summary", "records", "companies", "jobs", "inputs", "pages", "errors"):
            self.assertIn(key, payload)
        self.assertTrue(any(e["outcome"] == "BLOCKED" for e in payload["errors"]))
        pages = list(csv.DictReader(io.StringIO(self.read(run, "pages.csv").decode("utf-8-sig"))))
        self.assertEqual({p["kind"] for p in pages}, {"input", "listing", "detail"})
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(self.read(run, "xlsx")))
        self.assertEqual(book.sheetnames, ["All Fields", "Companies", "Jobs", "Pages", "Errors", "Run Summary", "Inputs"])
        evidence = self.service.evidence(self.ctx, run["id"])
        self.assertEqual(evidence["total"], 7)
        self.assertIn("job_title", evidence["items"][0]["fields"])

    def test_duplicates_across_inputs_keep_every_source(self) -> None:
        self.pages["https://pg.example/jobs-mirror"] = (200, listing([("a1", "Analyst One", "Austin, TX")]))
        run = self.service.start(self.ctx, ["https://pg.example/careers", "https://pg.example/jobs-mirror"],
                                 "Find all job titles and job URLs")
        self.run_task(run)
        run = self.service.get(self.ctx, run["id"])
        self.assertEqual(run["stats"]["duplicates_removed"], 1)
        rows = self.service.records(self.ctx, run["id"], "jobs")["items"]
        one = next(r for r in rows if r["job_title"] == "Analyst One")
        self.assertEqual(one["source_urls"], ["https://pg.example/careers", "https://pg.example/jobs-mirror"])
        self.assertEqual(one["_merged_from"][0]["input_url"], "https://pg.example/jobs-mirror")


class CrmTests(Base):
    def setUp(self) -> None:
        super().setUp()
        crm = self.platform.service("crm")
        crm.ensure_defaults(self.ctx)
        crm.create_company(self.ctx, {"name": "Acme Manufacturing", "website": "https://acme-mfg.com",
                                      "careers_url": "https://acme-mfg.com/old-careers"})
        run = self.service.start(self.ctx, ["https://www.acme-mfg.com/", "https://beta.example/"],
                                 "Get company name, website and careers URL")
        self.run_task(run)
        self.run = self.service.get(self.ctx, run["id"])

    def test_matching_is_read_only(self) -> None:
        before = self.store.count(self.ctx, "companies")
        match = self.service.match_crm(self.ctx, self.run["id"])
        by_name = {m["company_name"]: m for m in match["items"]}
        self.assertEqual(by_name["Acme Manufacturing"]["match"], "conflict")
        self.assertIn("careers_url", by_name["Acme Manufacturing"]["conflicts"])
        self.assertEqual(by_name["Beta Corp"]["match"], "new")
        self.assertEqual(self.store.count(self.ctx, "companies"), before)

    def test_propose_review_apply(self) -> None:
        before = self.store.count(self.ctx, "companies")
        made = self.service.propose(self.ctx, self.run["id"], ["company", "task"])
        self.assertEqual(made["created"], 4)
        self.assertEqual(self.service.propose(self.ctx, self.run["id"], ["company", "task"])["created"], 0,
                         "idempotent")
        self.assertEqual(self.store.count(self.ctx, "companies"), before, "proposing changes nothing")
        proposals = self.service.proposals(self.ctx, self.run["id"])["items"]
        beta = next(p for p in proposals if p["action"] == "company" and p["payload"]["values"]["name"] == "Beta Corp")
        with self.assertRaises(ValidationError):
            self.service.apply_proposals(self.ctx, [beta["id"]])
        self.service.review(self.ctx, [beta["id"]], "approved")
        result = self.service.apply_proposals(self.ctx, [beta["id"]])
        self.assertEqual(result, {"applied": 1, "failed": 0})
        self.assertEqual(self.store.count(self.ctx, "companies"), before + 1)
        applied = self.store.get(self.ctx, "scrape_proposals", beta["id"])
        self.assertEqual((applied["status"], applied["applied_entity_type"]), ("applied", "companies"))
        task = next(p for p in proposals if p["action"] == "task" and p["record_key"].startswith("beta.example"))
        self.service.review(self.ctx, [task["id"]], "approved")
        self.service.apply_proposals(self.ctx, [task["id"]])
        self.assertEqual(self.store.count(self.ctx, "crm_tasks"), 1)
        rejected = next(p for p in proposals if p["action"] == "task" and p["id"] != task["id"])
        self.service.review(self.ctx, [rejected["id"]], "rejected")
        with self.assertRaises(ValidationError):
            self.service.apply_proposals(self.ctx, [rejected["id"]])

    def test_a_possible_duplicate_is_never_merged_automatically(self) -> None:
        proposal = self.store.insert(self.ctx, "scrape_proposals", {
            "run_id": self.run["id"], "record_key": "dup", "action": "company", "status": "approved",
            "match": "possible_duplicate", "payload": {"values": {"name": "Acme Mfg"}}})
        result = self.service.apply_proposals(self.ctx, [proposal["id"]])
        self.assertEqual(result, {"applied": 0, "failed": 1})
        self.assertIn("needs review", self.store.get(self.ctx, "scrape_proposals", proposal["id"])["error"])


class TemplateTests(Base):
    def test_builtin_and_saved_templates(self) -> None:
        templates = self.service.templates
        names = [t["name"] for t in templates.list(self.ctx)]
        for expected in ("Company Research", "Company + Website", "Company + Jobs", "Company + ERP Technology",
                         "Hiring Intelligence", "Contact Discovery", "Job Extraction", "Custom Research"):
            self.assertIn(expected, names)
        with self.assertRaises(ValidationError):
            templates.update(self.ctx, "builtin:company-jobs", {"name": "x"})
        copy = templates.duplicate(self.ctx, "builtin:hiring-intelligence")
        self.assertEqual(copy["name"], "Hiring Intelligence (copy)")
        self.assertTrue(copy["options"]["follow_details"])
        saved = templates.create(self.ctx, {"name": "SAP jobs", "instruction": "Find all job titles and job URLs",
                                            "options": {"max_pages": 2}})
        edited = templates.update(self.ctx, saved["id"], {"description": "SAP hiring"})
        self.assertEqual(edited["description"], "SAP hiring")
        from cloud.intel.core.context import ConflictError

        with self.assertRaises(ConflictError):
            templates.create(self.ctx, {"name": "SAP jobs", "instruction": "x"})
        run = self.service.start(self.ctx, ["https://pg.example/careers"], template_id=saved["id"])
        self.assertEqual(run["stats"]["options"]["max_pages"], 2)
        self.assertEqual(run["instruction"], "Find all job titles and job URLs")
        self.run_task(run)
        self.assertEqual(self.service.get(self.ctx, run["id"])["stats"]["records"], 4, "2 pages of 2 jobs")
        templates.delete(self.ctx, saved["id"])
        self.assertNotIn("SAP jobs", [t["name"] for t in templates.list(self.ctx)])


class AgentIntegrationTests(Base):
    def test_the_research_agent_invokes_the_scraper(self) -> None:
        before = self.store.count(self.ctx, "companies")
        research = self.platform.service("research")
        plan = research.plan(self.ctx, "Scrape https://pg.example/careers and https://det.example/careers and "
                                       "collect their open jobs.")
        self.assertIn("scrape_jobs", [s["tool"] for s in plan["plan"]])
        approved = research.approve(self.ctx, plan["id"])
        task = run_task_inline(self.platform, self.ctx.workspace_id, approved["task_id"])
        self.assertEqual(task["status"], "completed", task.get("error"))
        run = self.store.get(self.ctx, "research_runs", plan["id"])
        step = next(s for s in run["plan"] if s["tool"] == "scrape_jobs")
        self.assertEqual(step["status"], "done", step.get("result"))
        scrape = self.service.get(self.ctx, step["result"]["scrape_run_id"])
        self.assertEqual(scrape["status"], "completed")
        self.assertEqual(scrape["stats"]["records"], 7)
        self.assertEqual(self.store.count(self.ctx, "companies"), before, "research never changes the CRM itself")

    def test_control_room_tools(self) -> None:
        from cloud.intel.agent.tools import TOOLS

        self.assertEqual(TOOLS["plan_scrape"].risk, "read")
        self.assertFalse(TOOLS["plan_scrape"].needs_approval({}))
        self.assertTrue(TOOLS["apply_scrape_proposals"].needs_approval({}), "CRM changes always need approval")
        self.assertTrue(TOOLS["run_ai_scraper"].needs_approval({"affected": 51}))
        self.assertFalse(TOOLS["run_ai_scraper"].needs_approval({"affected": 3}))

    def test_toolkit(self) -> None:
        from cloud.intel.scraper import toolkit

        run = toolkit.create_scrape_run(self.platform, self.ctx, ["https://pg.example/careers"], "Find all job titles")
        self.assertIsNone(run.get("task_id"), "created, not started")
        run = toolkit.start_scrape(self.platform, self.ctx, run["id"])
        self.run_task(run)
        info = toolkit.get_scrape_run(self.platform, self.ctx, run["id"])
        self.assertEqual((info["status"], info["records"]), ("completed", 5))
        self.assertEqual(toolkit.get_scrape_results(self.platform, self.ctx, run["id"], "jobs")["total"], 5)
        self.assertEqual(toolkit.get_scrape_pages(self.platform, self.ctx, run["id"])["total"], 3)
        self.assertEqual(toolkit.export_scrape_results(self.platform, self.ctx, run["id"], "ndjson")["content_type"],
                         "application/x-ndjson")


class LoggingTests(Base):
    def test_structured_logs_carry_no_secrets(self) -> None:
        import os
        from unittest import mock

        with mock.patch.dict(os.environ, {"GEMINI_API_KEY": "g-secret-key-123456"}), \
                self.assertLogs("cloud.intel.scraper", level="INFO") as logs:
            run = self.service.start(self.ctx, ["https://pg.example/careers"], "Find all job titles")
            self.run_task(run)
        text = "\n".join(logs.output)
        self.assertIn("scrape.input", text)
        self.assertIn("scrape.finish", text)
        self.assertNotIn("g-secret-key-123456", text)


class ScraperApiV4Tests(unittest.TestCase):
    def setUp(self) -> None:
        from fastapi.testclient import TestClient

        from cloud.api.auth import DevTokenIssuer
        from cloud.api.main import create_app
        from cloud.api.settings import Settings
        from cloud.worker.dispatcher import NullDispatcher

        scratch = tempfile.TemporaryDirectory()
        self.addCleanup(scratch.cleanup)
        root = Path(scratch.name)
        pages = {**PAGES, **PAGINATED}
        self.platform = Platform(MemoryStore(), storage=LocalFileStorage(root / "platform"),
                                 config=PlatformConfig(files_dir=root / "platform",
                                                       extra={"fetcher_factory": lambda: http_for(pages)}))
        issuer = DevTokenIssuer("scraper-v4-api-tests-secret-0123456789abcdef")
        app = create_app(Settings(auth_mode="dev", results_dir=root / "results"),
                         storage=LocalFileStorage(root / "results"), token_verifier=issuer,
                         dispatcher=NullDispatcher(), platform=self.platform)
        self.client = TestClient(app)
        self.client.headers["Authorization"] = f"Bearer {issuer.issue('alice@example.com')['access_token']}"
        self.client.__enter__()
        self.addCleanup(self.client.__exit__, None, None, None)
        self.ws = self.client.post("/api/v1/workspaces", json={"name": "V4", "seed": False}).json()["id"]
        self.base = f"/api/v1/w/{self.ws}/scraper"
        self.platform.service("crm").ensure_defaults(Ctx(self.ws, issuer.user_id_for("alice@example.com"), "owner"))

    def test_the_whole_flow_over_http(self) -> None:
        plan = self.client.post(self.base + "/plan", json={"urls": "https://pg.example/careers",
                                                           "instruction": "Find all job titles and job URLs",
                                                           "options": {"max_pages": 5}})
        self.assertEqual(plan.status_code, 200, plan.text)
        self.assertEqual(plan.json()["estimate"]["estimated_cost_usd"], 0.0)
        schema = plan.json()["schema"]
        schema["fields"].append({"name": "location", "type": "string", "level": "job"})
        run = self.client.post(self.base + "/runs", json={"urls": "https://pg.example/careers", "schema": schema,
                                                          "options": {"max_pages": 5}}).json()
        self.assertEqual([f["name"] for f in run["schema"]["fields"]], ["job_title", "job_url", "location"])
        paused = self.client.post(f"{self.base}/runs/{run['id']}/pause")
        self.assertEqual(paused.json()["status"], "paused", paused.text)
        self.assertEqual(self.client.post(f"{self.base}/runs/{run['id']}/resume").json()["status"], "queued")
        run_task_inline(self.platform, self.ws, self.client.get(f"{self.base}/runs/{run['id']}").json()["task_id"])
        done = self.client.get(f"{self.base}/runs/{run['id']}").json()
        self.assertEqual(done["status"], "completed")
        self.assertEqual(self.client.get(f"{self.base}/runs/{run['id']}/pages").json()["total"], 3)
        self.assertEqual(self.client.get(f"{self.base}/runs/{run['id']}/errors").status_code, 200)
        self.assertEqual(self.client.get(f"{self.base}/runs/{run['id']}/evidence").json()["total"], 5)
        for fmt, view in (("ndjson", "all"), ("csv", "pages"), ("csv", "errors"), ("xlsx", "all")):
            response = self.client.get(f"{self.base}/runs/{run['id']}/files/{fmt}", params={"view": view})
            self.assertEqual(response.status_code, 200, (fmt, view))
        self.assertEqual(self.client.get(f"{self.base}/runs/{run['id']}/files/pdf").status_code, 404)
        match = self.client.get(f"{self.base}/runs/{run['id']}/crm/match")
        self.assertEqual(match.status_code, 200, match.text)
        made = self.client.post(f"{self.base}/runs/{run['id']}/crm/propose", json={"actions": ["job"]})
        self.assertEqual(made.json()["created"], 5)
        proposals = self.client.get(f"{self.base}/runs/{run['id']}/proposals").json()["items"]
        refused = self.client.post(self.base + "/proposals/apply", json={"ids": [proposals[0]["id"]]})
        self.assertEqual(refused.status_code, 422)
        self.client.post(self.base + "/proposals/review", json={"ids": [proposals[0]["id"]], "decision": "approved"})
        applied = self.client.post(self.base + "/proposals/apply", json={"ids": [proposals[0]["id"]]})
        self.assertEqual(applied.json(), {"applied": 1, "failed": 0}, applied.text)
        restart = self.client.post(f"{self.base}/runs/{run['id']}/restart")
        self.assertEqual(restart.status_code, 201)

    def test_templates_over_http(self) -> None:
        listed = self.client.get(self.base + "/templates").json()["items"]
        self.assertEqual(len([t for t in listed if t["builtin"]]), 8)
        made = self.client.post(self.base + "/templates", json={"name": "Mine", "instruction": "Get company name"})
        self.assertEqual(made.status_code, 201, made.text)
        tid = made.json()["id"]
        self.assertEqual(self.client.patch(f"{self.base}/templates/{tid}", json={"category": "custom"}).json()["category"],
                         "custom")
        dup = self.client.post(f"{self.base}/templates/{tid}/duplicate", json={})
        self.assertEqual(dup.json()["name"], "Mine (copy)")
        self.assertEqual(self.client.delete(f"{self.base}/templates/{tid}").status_code, 204)
        self.assertEqual(self.client.get(f"{self.base}/templates/{tid}").status_code, 404)


if __name__ == "__main__":
    unittest.main()


try:
    import playwright  # noqa: F401
    _HAVE_PLAYWRIGHT = True
except ImportError:  # pragma: no cover
    _HAVE_PLAYWRIGHT = False


@unittest.skipUnless(_HAVE_PLAYWRIGHT, "Playwright is not installed")
class RealBrowserTests(unittest.TestCase):
    """Real headless Chromium, no network: pages are served from memory after the renderer's
    own public-address check has passed."""

    PAGE = """<html><head><title>Careers at Clicky</title>
<script src="http://127.0.0.1/evil.js"></script></head><body>
<ul id="jobs"><li><a href="/careers/one"><h3>Job One</h3></a></li></ul>
<button id="more" onclick="document.getElementById('jobs').insertAdjacentHTML('beforeend',
 '<li><a href=&quot;/careers/two&quot;><h3>Job Two</h3></a></li><li><a href=&quot;/careers/three&quot;><h3>Job Three</h3></a></li>');
 this.remove();">Load more</button></body></html>"""

    def renderer(self):
        from cloud.intel.scraper.fetcher import PlaywrightRenderer

        pages = {"https://clicky.example/jobs": self.PAGE}

        class OfflineRenderer(PlaywrightRenderer):
            def _fulfill(self, route):
                body = pages.get(route.request.url)
                if body is None:
                    route.abort()
                else:
                    route.fulfill(status=200, content_type="text/html", body=body)

        return OfflineRenderer(resolver=fake_resolver(), timeout_ms=20000)

    def test_chromium_clicks_load_more_and_blocks_private_addresses(self) -> None:
        renderer = self.renderer()
        result = renderer.render("https://clicky.example/jobs", interact=True, max_clicks=3)
        self.assertIsNone(result.error)
        self.assertEqual(result.clicks, 1)
        self.assertIn("Job Three", result.html)
        self.assertIn("http://127.0.0.1/evil.js", renderer.blocked, "a private address is never loaded")
        pages = {"https://clicky.example/jobs": (200, self.PAGE)}
        page, _ = crawl("https://clicky.example/jobs", "Find all job titles", pages, renderer=renderer, browser=True)
        self.assertEqual([r["job_title"] for r in values(page)], ["Job One", "Job Two", "Job Three"])
        self.assertTrue(any(p["browser_used"] and p["browser_outcome"] == "OK" for p in page.pages))


BOARD = """<html><head><title>Python Job Board | Python.org</title>
<meta property="og:site_name" content="Python.org"></head><body>
<nav id="mainnav"><a href="/jobs/">Jobs</a><a href="/jobs/create/">Submit a job</a></nav>
<aside class="sidebar"><h3>Job types</h3><ul>
  <li><a href="/jobs/type/back-end/">Back end</a></li><li><a href="/jobs/type/big-data/">Big Data</a></li></ul>
  <h3>Locations</h3><ul><li><a href="/jobs/location/remote-remote/">Remote – Remote</a></li></ul></aside>
<ol class="list-recent-jobs">
{items}
</ol>
<a href="/community/jobs/howto/">job submission how-to</a> <a href="/jobs/feed/rss/">Subscribe via RSS</a>
<ul class="pagination"><li><a href="?page=2">Next</a></li></ul></body></html>"""
BOARD_ITEM = """<li><h2 class="listing-company"><span class="listing-company-name">
  <a href="/jobs/{id}/">{title}</a><br/>{company}</span>
  <span class="listing-location"><a href="/jobs/location/{loc_slug}/">{loc}</a></span></h2>
  <span class="listing-job-type"><a href="/jobs/type/back-end/">Back end</a></span>
  <span class="listing-posted">Posted: <time datetime="2026-09-2{day}T08:00:00+00:00">2{day} September 2026</time></span></li>"""


def board(jobs):
    return BOARD.replace("{items}", "".join(BOARD_ITEM.format(**j) for j in jobs))


class JobBoardTests(unittest.TestCase):
    """A public job board listing other companies' jobs, with filter/taxonomy links everywhere
    (modelled on the page the V2 live test ran against)."""

    def test_only_postings_are_jobs_and_each_keeps_its_company(self) -> None:
        jobs1 = [{"id": 8139, "title": "Senior Staff Engineer", "company": "Kraken", "loc": "Remote (UK / EU)",
                  "loc_slug": "remote-uk-eu", "day": 8},
                 {"id": 8137, "title": "Django Developer", "company": "Widget Ltd", "loc": "Birmingham, UK",
                  "loc_slug": "birmingham-uk", "day": 7},
                 {"id": 8136, "title": "ML Engineer", "company": "Deep Co", "loc": "Worldwide",
                  "loc_slug": "worldwide", "day": 6}]
        jobs2 = [{"id": 8101, "title": "Python Developer", "company": "Remote First", "loc": "Anywhere",
                  "loc_slug": "anywhere", "day": 1}]
        pages = {"https://board.example/jobs/": (200, board(jobs1)),
                 "https://board.example/jobs/?page=2": (200, board(jobs2).replace('<a href="?page=2">Next</a>', "")),
                 "https://board.example/jobs/8139/": (200, detail_page("Senior Staff Engineer", ("London", "UK"),
                                                                       90000))}
        page, f = crawl("https://board.example/jobs/", "Find all job titles, job URLs, company name, location and "
                                                       "posted date.", pages, follow_details=True, max_detail_pages=1)
        rows = values(page)
        self.assertEqual([r["job_title"] for r in rows], ["Senior Staff Engineer", "Django Developer", "ML Engineer",
                                                          "Python Developer"])
        self.assertEqual([r["company_name"] for r in rows][1:], ["Widget Ltd", "Deep Co", "Remote First"])
        self.assertEqual(rows[1]["posted_date"], "2026-09-27T08:00:00+00:00")
        self.assertEqual(rows[1]["location"], "Birmingham, UK")
        detail_fetches = [p["url"] for p in page.pages if p["kind"] == "detail" and p["outcome"] == "OK"]
        self.assertEqual(detail_fetches, ["https://board.example/jobs/8139/"], "details are postings, not filters")
        for bad in ("type/back-end", "location/", "create", "feed/rss", "howto"):
            self.assertFalse(any(bad in str(r["job_url"]) for r in rows), bad)


PYORG = """<html><head><title>Python Job Board | Python.org</title>
<meta property="og:site_name" content="Python.org"></head><body>
<div id="touchnav-wrapper">
 <header class="main-header"><nav id="mainnav" class="python-navigation main-navigation">
   <a href="/jobs/">Jobs</a> <a href="/jobs/create/">Submit a job</a></nav></header>
 <div id="content" class="content-wrapper"><div class="container">
  <aside class="left-sidebar" role="complementary"><h3>Job types</h3><ul>
   <li><a href="/jobs/types/">All types</a></li><li><a href="/jobs/type/back-end/">Back end</a></li>
   <li><a href="/jobs/type/big-data/">Big Data</a></li><li><a href="/jobs/category/developer-engineer/">Developer</a></li>
   <li><a href="/jobs/location/remote-remote/">Remote - Remote</a></li></ul></aside>
  <section class="main-content with-left-sidebar" role="main">
   <div class="row"><div class="list-widget"><ol class="list-recent-jobs list-row-container menu">
{items}
   </ol></div></div>
   <p><a href="/community/jobs/howto/">job submission how-to</a> <a href="/jobs/feed/rss/">Subscribe via RSS</a></p>
   {next}
  </section>
 </div></div>
 <footer class="main-footer"><a href="/jobs/create/">Submit a job</a> <a href="/jobs/">All jobs</a></footer>
</div></body></html>"""


def pyorg(jobs, next_page=None):
    nav = (f'<ul class="pagination menu"><li><a href="?page={next_page}">Next &raquo;</a></li></ul>'
           if next_page else "")
    return PYORG.replace("{items}", "".join(BOARD_ITEM.format(**j) for j in jobs)).replace("{next}", nav)


PYORG_JOBS_1 = [{"id": 8139, "title": "Senior Staff Engineer", "company": "Kraken", "loc": "Remote (UK / EU)",
                 "loc_slug": "remote-uk-eu", "day": 8},
                {"id": 8137, "title": "Django Developer", "company": "Widget Ltd", "loc": "Birmingham, UK",
                 "loc_slug": "birmingham-uk", "day": 7},
                {"id": 8136, "title": "ML Engineer", "company": "Deep Co", "loc": "Worldwide",
                 "loc_slug": "worldwide", "day": 6}]
PYORG_JOBS_2 = [{"id": 8101, "title": "Python Developer", "company": "Remote First", "loc": "Anywhere",
                 "loc_slug": "anywhere", "day": 1}]
PYORG_PAGES = {
    "https://www.python.org/jobs/": (200, pyorg(PYORG_JOBS_1, next_page=2)),
    "https://www.python.org/jobs/?page=2": (200, pyorg(PYORG_JOBS_2)),
    "https://www.python.org/jobs/8139/": (200, detail_page("Senior Staff Engineer", ("London", "UK"), 90000)
                                          .replace('"name": "Detail Co"', '"name": "Kraken"')),
    "https://www.python.org/jobs/8137/": (200, detail_page("Django Developer", ("Birmingham", "UK"), 60000)
                                          .replace('"name": "Detail Co"', '"name": "Widget Ltd"')),
}


class PythonOrgRegressionTests(unittest.TestCase):
    """The live-test regression: python.org wraps its job list in
    <section class="main-content with-left-sidebar" role="main">. A substring match on
    "sidebar" rejected every real job; only whole class tokens and explicit roles count now."""

    INSTRUCTION = ("Find all job titles, job URLs, company name, location, posted date, description and "
                   "employment type.")

    def run_board(self, pages=None, **options):
        opts = {"max_listing_pages": 2, "follow_details": True, "max_detail_pages": 2, "max_pages": 8, **options}
        return crawl("https://www.python.org/jobs/", self.INSTRUCTION, pages or PYORG_PAGES, **opts)

    def test_real_jobs_are_found_and_navigation_is_not(self) -> None:
        page, _ = self.run_board()
        rows = values(page)
        self.assertEqual([r["job_url"] for r in rows],
                         ["https://www.python.org/jobs/8139/", "https://www.python.org/jobs/8137/",
                          "https://www.python.org/jobs/8136/", "https://www.python.org/jobs/8101/"])
        for bad in ("/type", "/category/", "/location/", "/create", "/feed/", "/howto"):
            self.assertFalse(any(bad in r["job_url"] for r in rows), bad)
        fetched = [(p["kind"], p["url"]) for p in page.pages if p["outcome"] == Outcome.OK]
        self.assertEqual(fetched, [("input", "https://www.python.org/jobs/"),
                                   ("listing", "https://www.python.org/jobs/?page=2"),
                                   ("detail", "https://www.python.org/jobs/8139/"),
                                   ("detail", "https://www.python.org/jobs/8137/")],
                         "pagination proceeds after jobs are found; details are postings")
        self.assertTrue(any("detail page limit (2)" in p for p in page.problems))
        self.assertFalse(any("candidate careers pages" in p for p in page.problems), "no discovery needed")
        urls = [p["url"] for p in page.pages]
        self.assertEqual(len(urls), len(set(urls)), "no page twice")
        self.assertEqual(len({r["job_url"] for r in rows}), len(rows), "no record twice")

    def test_listing_and_detail_fields_merge(self) -> None:
        page, _ = self.run_board()
        first = page.records[0]
        self.assertEqual(first["company_name"].value, "Kraken")
        self.assertEqual(first["location"].value, "London, UK, US", "structured detail data beats the card")
        self.assertEqual([a.value for a in first["location"].alternatives], ["Remote (UK / EU)"])
        self.assertEqual(first["employment_type"].value, "FULL_TIME")
        self.assertIn("SAP S/4HANA", first["description"].value)
        self.assertEqual(first["job_url"].source_url, "https://www.python.org/jobs/")
        self.assertEqual(first["description"].source_url, "https://www.python.org/jobs/8139/")
        third = page.records[2]
        self.assertEqual((third["company_name"].value, third["location"].value), ("Deep Co", "Worldwide"))
        self.assertNotIn("description", third, "not enriched beyond the detail limit")

    def test_container_tokens_and_roles(self) -> None:
        from bs4 import BeautifulSoup

        from cloud.intel.scraper.extractor import _in_chrome

        html = """<div class="main-content with-left-sidebar"><a id="a" href="/jobs/1/">x</a></div>
        <div class="content-area"><a id="b" href="/jobs/2/">x</a></div>
        <div class="sidebar"><a id="c" href="/jobs/3/">x</a></div>
        <div role="navigation"><a id="d" href="/jobs/4/">x</a></div>
        <nav><a id="e" href="/jobs/5/">x</a></nav>
        <div class="navbar-wrapper-thing"><a id="f" href="/jobs/6/">x</a></div>
        <div id="sidebar"><main><a id="g" href="/jobs/7/">x</a></main></div>"""
        soup = BeautifulSoup(html, "lxml")
        verdict = {a["id"]: _in_chrome(a) for a in soup.find_all("a")}
        self.assertEqual(verdict, {"a": False, "b": False, "c": True, "d": True, "e": True, "f": False, "g": False})

    def test_filter_and_howto_pages_are_never_careers_candidates(self) -> None:
        from cloud.intel.scraper.discovery import rank_candidates

        links = [("All types", "https://www.python.org/jobs/types/"),
                 ("Back end", "https://www.python.org/jobs/type/back-end/"),
                 ("job submission how-to", "https://www.python.org/community/jobs/howto/"),
                 ("Remote", "https://www.python.org/jobs/location/remote/"),
                 ("Jobs", "https://www.python.org/jobs/")]
        ranked = rank_candidates("https://www.python.org/", links)
        self.assertEqual([u for u, _s, _w in ranked], ["https://www.python.org/jobs/"])
        # A listing whose jobs cannot be read never wanders into its own filter pages.
        empty = {"https://www.python.org/jobs/": (200, pyorg([]))}
        page, f = self.run_board(empty)
        calls = f.http._session.calls
        self.assertFalse(any(("/type" in u or "howto" in u or "/location/" in u) for u in calls), calls)


CATEGORY_ITEM = """<li><h2 class="listing-company"><span class="listing-company-name">
  <a href="/jobs/{id}/">{title}</a><br/>{company}</span>
  <span class="listing-location"><a href="/jobs/location/{loc_slug}/">{loc}</a></span></h2>
  <span class="listing-job-type"><a href="/jobs/type/back-end/">Back end</a></span>
  <span class="listing-posted">Posted: <time datetime="2026-09-2{day}T08:00:00+00:00">2{day} September 2026</time></span>
  <span class="listing-company-category"><a href="/jobs/category/developer-engineer/">Developer / Engineer</a></span></li>"""
DECORATED_DETAIL = """<html><head><title>Job: {title} at {company} | Python.org</title>
<meta property="og:title" content="Job: {title} at {company}"></head><body>
<nav><a href="/jobs/">Jobs</a></nav><main><h1 class="listing-company"><span class="company-name">{title}</span></h1>
<p>We need 5+ years of experience with Python and Django. {company} is hiring.</p></main></body></html>"""


class LiveFindingsRegressionTests(unittest.TestCase):
    """Found by the final live python.org run: a card's category element was read as the
    company, and a detail page's decorated og:title replaced the clean listing title."""

    def test_company_and_title_survive_detail_merge(self) -> None:
        jobs = [{"id": 8139, "title": "Senior Staff Engineer", "company": "Kraken", "loc": "Remote (UK / EU)",
                 "loc_slug": "remote-uk-eu", "day": 8},
                {"id": 8137, "title": "Django Developer", "company": "Widget Ltd", "loc": "Birmingham, UK",
                 "loc_slug": "birmingham-uk", "day": 7}]
        listing_html = PYORG.replace("{items}", "".join(CATEGORY_ITEM.format(**j) for j in jobs)).replace("{next}", "")
        pages = {"https://www.python.org/jobs/": (200, listing_html),
                 "https://www.python.org/jobs/8139/": (200, DECORATED_DETAIL.format(title="Senior Staff Engineer",
                                                                                   company="Kraken"))}
        page, _ = crawl("https://www.python.org/jobs/", "Find all job titles, job URLs, company name and description.",
                        pages, follow_details=True, max_detail_pages=1)
        rows = values(page)
        self.assertEqual([r["company_name"] for r in rows], ["Kraken", "Widget Ltd"], "never the category")
        self.assertEqual(rows[0]["job_title"], "Senior Staff Engineer", "the listing title is kept")
        first = page.records[0]
        self.assertEqual(first["job_title"].source_url, "https://www.python.org/jobs/")
        self.assertIn("Job: Senior Staff Engineer at Kraken", [a.value for a in first["job_title"].alternatives])
        self.assertIn("5+ years", first["description"].value, "detail data still merges")
        self.assertEqual(first["description"].source_url, "https://www.python.org/jobs/8139/")

    def test_a_structured_detail_title_may_still_win(self) -> None:
        from cloud.intel.scraper.detail import merge_job

        listing = {"job_title": FieldValue("ERP Lead", "link", 0.7, "a", "https://l")}
        merged = merge_job(listing, {"job_title": FieldValue("ERP Lead (SAP)", "json-ld", 0.95, "ld", "https://d")})
        self.assertEqual(merged["job_title"].value, "ERP Lead (SAP)")
        self.assertEqual([a.value for a in merged["job_title"].alternatives], ["ERP Lead"])
