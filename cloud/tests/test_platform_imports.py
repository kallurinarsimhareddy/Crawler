"""Track A: the internal data engine — multi-file import batches, explicit
mapping, merge with provenance — and the import/CRM HTTP API."""

from __future__ import annotations

import io
import json
import tempfile
import unittest
import uuid
from pathlib import Path

from openpyxl import Workbook

from cloud.intel.core.context import ConflictError, Ctx, ValidationError
from cloud.intel.imports.parse import iter_rows, parse_file
from cloud.intel.platform import Platform, PlatformConfig
from cloud.intel.store.memory import MemoryStore
from cloud.intel.tasks.worker import run_task_inline
from cloud.shared.storage import LocalFileStorage


def csv_bytes(rows, header=("Company Name", "Website", "Industry", "City", "Notes"), encoding="utf-8"):
    lines = [",".join(header)] + [",".join(r) for r in rows]
    return ("\n".join(lines) + "\n").encode(encoding)


def xlsx_bytes(rows, header=("Account Name", "URL", "Industry", "City")):
    wb = Workbook()
    ws = wb.active
    ws.append(list(header))
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ImportTestBase(unittest.TestCase):
    def setUp(self) -> None:
        scratch = tempfile.TemporaryDirectory()
        self.addCleanup(scratch.cleanup)
        store = MemoryStore()
        self.user = str(uuid.uuid4())
        ws = store.create_workspace(self.user, "Imports", f"imp-{uuid.uuid4().hex[:8]}")
        self.ctx = Ctx(ws["id"], self.user, "owner")
        self.platform = Platform(store, storage=LocalFileStorage(Path(scratch.name)),
                                 config=PlatformConfig(files_dir=Path(scratch.name)))
        self.imports = self.platform.service("imports")
        self.store = store


class TestParse(unittest.TestCase):
    def test_csv_cp1252_semicolon_and_problems(self) -> None:
        data = "Name;Ville\nSociété Générale;Paris\n\n".encode("cp1252")
        parsed = parse_file("f.csv", data)
        self.assertEqual(parsed.columns, ["Name", "Ville"])
        self.assertEqual(parsed.row_count, 1)
        self.assertEqual(list(iter_rows("csv", data))[0][1]["Name"], "Société Générale")
        dup = parse_file("d.csv", b"Name,Name\nA,B\n")
        self.assertTrue(any("duplicate" in p for p in dup.problems))
        self.assertIn("the file is empty", parse_file("e.csv", b"").problems)

    def test_xlsx_and_json(self) -> None:
        parsed = parse_file("f.xlsx", xlsx_bytes([("Acme", "acme.com", "Mfg", "Tulsa"), ("Beta", 12.0, None, None)]))
        self.assertEqual(parsed.row_count, 2)
        rows = list(iter_rows("xlsx", xlsx_bytes([("Beta", 12.0, None, None)])))
        self.assertEqual(rows[0][1]["URL"], "12")  # numbers stay as written, no ".0"
        js = json.dumps([{"company": "A", "domain": "a.com"}, {"company": "B", "extra": 1}]).encode()
        self.assertEqual(parse_file("f.json", js).columns, ["company", "domain", "extra"])
        from cloud.intel.imports.parse import ParseError

        with self.assertRaises(ParseError):
            parse_file("bad.json", b"{\"x\": 1}")


class TestBatchLifecycle(ImportTestBase):
    def _upload_thirteen(self, batch_id):
        # 12 compatible files across three formats, deliberately overlapping companies...
        for i in range(5):
            self.imports.add_file(self.ctx, batch_id, f"csv_{i}.csv", csv_bytes([
                (f"Company {i}", f"https://www.company{i}.com", "Manufacturing", "Tulsa", "=cmd()"),
                ("Shared Widgets Inc", "sharedwidgets.com", "Manufacturing", "", f"from csv {i}"),
            ]))
        for i in range(4):
            self.imports.add_file(self.ctx, batch_id, f"xlsx_{i}.xlsx", xlsx_bytes([
                (f"Xl Company {i}", f"xl{i}.com", "Retail", "Dallas"),
                ("Shared Widgets", "http://sharedwidgets.com/about", "Aerospace", "Austin"),
            ]))
        for i in range(3):
            self.imports.add_file(self.ctx, batch_id, f"json_{i}.json", json.dumps([
                {"Company Name": f"Json Co {i}", "Website": f"jsonco{i}.io", "Industry": "Software"},
                {"Company Name": "No Domain Co", "Website": "", "Industry": "Software"},
            ]).encode())
        # ...and one file that cannot be imported: no company column at all.
        self.imports.add_file(self.ctx, batch_id, "contacts_only.csv",
                              b"First Name,Last Name,Email\nJane,Doe,jane@x.com\n")

    def test_thirteen_files_validate_map_and_merge(self) -> None:
        batch = self.imports.create_batch(self.ctx, "Q3 lists", "companies")
        self._upload_thirteen(batch["id"])
        report = self.imports.validate(self.ctx, batch["id"])
        self.assertEqual((report["compatible"], report["incompatible"]), (12, 1))
        bad = [f for f in report["files"] if f["status"] == "incompatible"]
        self.assertEqual(bad[0]["filename"], "contacts_only.csv")
        self.assertIn("required field: name", bad[0]["problems"][0])

        suggestions = self.imports.suggest_mapping(self.ctx, batch["id"])
        self.assertFalse(suggestions["applied"])
        by_column = {s["column"]: s["target"] for s in suggestions["suggestions"]}
        self.assertEqual(by_column["Account Name"], "company.name")
        self.assertEqual(by_column["URL"], "company.website")
        # nothing was stored by asking for suggestions
        self.assertEqual(self.store.get(self.ctx, "import_batches", batch["id"])["mapping"], {})
        with self.assertRaises(ConflictError):
            self.imports.merge(self.ctx, batch["id"])

        with self.assertRaises(ValidationError):  # required field not mapped
            self.imports.set_mapping(self.ctx, batch["id"], {"Website": "company.website"})
        with self.assertRaises(ValidationError):
            self.imports.set_mapping(self.ctx, batch["id"], {"Company Name": "company.password"})

        # A mapping that leaves the XLSX files' name column unmapped makes them incompatible, by name.
        partial = {"Company Name": "company.name", "Website": "company.website", "Notes": None}
        report = self.imports.set_mapping(self.ctx, batch["id"], partial)
        self.assertEqual(report["compatible"], 8)
        xlsx = [f for f in report["files"] if f["filename"].startswith("xlsx")]
        self.assertTrue(all(f["status"] == "incompatible" for f in xlsx))

        # Files name the company column differently: both map to company.name, explicitly.
        mapping = {"Company Name": "company.name", "Account Name": "company.name", "Website": "company.website",
                   "URL": "company.website", "Industry": "company.industry", "City": "company.city", "Notes": None}
        report = self.imports.set_mapping(self.ctx, batch["id"], mapping)
        self.assertEqual((report["compatible"], report["incompatible"]), (12, 1))

        task = self.imports.merge(self.ctx, batch["id"])
        run_task_inline(self.platform, self.ctx.workspace_id, task["id"])
        done = self.platform.tasks.get(self.ctx, task["id"])
        self.assertEqual(done["status"], "completed", done["error"])
        stats = done["result"]
        self.assertEqual(stats["rows"], 24)
        # new: Company 0-4, Xl Company 0-3, Json Co 0-2, Shared Widgets Inc, No Domain Co
        self.assertEqual(stats["created"], 5 + 4 + 3 + 1 + 1)
        self.assertEqual(stats["duplicates"], 8)       # Shared Widgets again, same domain, in 8 more files
        self.assertEqual(stats["needs_review"], 2)     # "No Domain Co" by name alone, twice
        self.assertEqual(stats["rows"], stats["created"] + stats["duplicates"] + stats["needs_review"])
        batch_row = self.store.get(self.ctx, "import_batches", batch["id"])
        self.assertEqual(batch_row["status"], "merged")

        shared = self.store.first(self.ctx, "companies", {"domain": "sharedwidgets.com"})
        self.assertEqual(shared["source_count"], 9)
        self.assertEqual(shared["industry"], "Manufacturing")   # first seen value kept...
        self.assertEqual(shared["city"], "Austin")               # ...empty value filled
        records = self.store.all(self.ctx, "source_records", {"entity_id": shared["id"]})
        self.assertEqual(len(records), 9)
        csv0 = self.store.first(self.ctx, "import_files", {"batch_id": batch["id"], "filename": "csv_0.csv"})
        first = [r for r in records if r["import_file_id"] == csv0["id"]][0]
        self.assertEqual((first["import_batch_id"], first["row_number"]), (batch["id"], 2))
        self.assertEqual(first["original"]["Notes"], "from csv 0")  # unmapped column kept, exactly as written
        conflicts = [r["normalized"].get("conflicts") or {} for r in records]
        self.assertTrue(any(c.get("industry", {}).get("incoming") == "Aerospace" for c in conflicts))

        pending = self.imports.rows(self.ctx, batch["id"], status="pending").rows
        self.assertEqual(len(pending), 2)
        resolved = self.imports.resolve_row(self.ctx, batch["id"], pending[0]["id"], "merge_into",
                                            pending[0]["problems"][0].split("possible duplicate of ")[1].split(":")[0])
        self.assertEqual(resolved["status"], "merged")

    def test_duplicate_file_and_limits(self) -> None:
        batch = self.imports.create_batch(self.ctx, "dups", "companies")
        data = csv_bytes([("A", "a.com", "x", "y", "")])
        self.imports.add_file(self.ctx, batch["id"], "a.csv", data)
        with self.assertRaises(ConflictError):
            self.imports.add_file(self.ctx, batch["id"], "a-copy.csv", data)
        with self.assertRaises(ValidationError):
            self.imports.add_file(self.ctx, batch["id"], "a.pdf", b"%PDF")
        with self.assertRaises(ValidationError):
            self.imports.create_batch(self.ctx, "x", "planets")

    def test_contacts_batch_and_pause_resume(self) -> None:
        batch = self.imports.create_batch(self.ctx, "people", "companies_and_contacts")
        rows = [(f"Co {i}", f"co{i}.com", f"First{i}", "Last", f"p{i}@co{i}.com", "IT Director") for i in range(450)]
        header = ("Company", "Domain", "First Name", "Last Name", "Email", "Title")
        self.imports.add_file(self.ctx, batch["id"], "people.csv", csv_bytes(rows, header=header))
        self.imports.set_mapping(self.ctx, batch["id"], {
            "Company": "company.name", "Domain": "company.domain", "First Name": "contact.first_name",
            "Last Name": "contact.last_name", "Email": "contact.email", "Title": "contact.title"})
        task = self.imports.merge(self.ctx, batch["id"])
        self.platform.tasks.pause(self.ctx, task["id"])          # paused before it starts
        self.platform.tasks.resume(self.ctx, task["id"])
        # simulate a pause request arriving mid-run
        from unittest import mock

        from cloud.intel.tasks.service import TaskReporter

        with mock.patch.object(TaskReporter, "should_pause", side_effect=[False, True]):
            run_task_inline(self.platform, self.ctx.workspace_id, task["id"])
        paused = self.platform.tasks.get(self.ctx, task["id"])
        self.assertEqual(paused["status"], "paused")
        self.assertEqual(paused["progress"]["checkpoint"]["row"], 399)
        self.platform.tasks.resume(self.ctx, task["id"])
        run_task_inline(self.platform, self.ctx.workspace_id, task["id"])
        done = self.platform.tasks.get(self.ctx, task["id"])
        self.assertEqual(done["status"], "completed", done["error"])
        self.assertEqual(self.store.count(self.ctx, "contacts"), 450)
        self.assertEqual(self.store.count(self.ctx, "companies"), 450)
        self.assertEqual(self.store.count(self.ctx, "import_rows", {"batch_id": batch["id"]}), 450)
        contact = self.store.first(self.ctx, "contacts", {"email": "p7@co7.com"})
        self.assertEqual(self.store.get(self.ctx, "companies", contact["company_id"])["domain"], "co7.com")


if __name__ == "__main__":
    unittest.main()
