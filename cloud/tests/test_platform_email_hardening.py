"""Email validation production hardening: the one final-status resolver, a deterministic
input matrix, de-duplication, explicit credit-exact paid checks, failure isolation and
recovery, background ingest of large pastes and damaged files, tenant isolation, exports.

Everything here is offline and synthetic: DNS, SMTP and EmailListVerify are fakes. Nothing in
this file is (or pretends to be) a genuine mailbox verification.
"""

from __future__ import annotations

import csv
import io
import tempfile
import time
import unittest
import uuid
from pathlib import Path
from unittest import mock

from cloud.intel.core.context import ConflictError, Ctx, ValidationError
from cloud.intel.email import engine
from cloud.intel.email import jobs as jobs_module
from cloud.intel.email.engine import DnsAnswer, final_filters, final_status
from cloud.intel.email.providers import EmailListVerifyProvider, LocalValidator, ValidationResult
from cloud.intel.email.service import EmailValidationService, max_age
from cloud.intel.tasks.worker import run_task_inline
from cloud.tests.test_platform_email_jobs import FakeDns, make_platform
from cloud.tests.test_platform_sources_support import FakeResponse, FakeSession

# --- the deterministic DNS world the matrix runs in ----------------------------------------------

#: domain -> how it answers. Anything unlisted has MX + SPF + DMARC.
WORLD = {
    "gone-domain.com": "nxdomain",     # does not exist
    "nomail-domain.com": "no_mail",    # exists, no MX and no A/AAAA
    "nullmx-domain.com": "null_mx",    # RFC 7505 null MX
    "flaky-dns.com": "servfail",       # temporary DNS failure
}


class WorldDns:
    def __init__(self):
        self.calls = []

    def lookup(self, name, rtype):
        self.calls.append((name, rtype))
        domain = name.removeprefix("_dmarc.")
        kind = WORLD.get(domain)
        if kind == "nxdomain":
            return DnsAnswer(error="nxdomain")
        if kind == "servfail":
            return DnsAnswer(error="servfail")
        if kind == "no_mail":
            return DnsAnswer(error="no_answer")
        if kind == "null_mx" and rtype == "MX":
            return DnsAnswer(["0 ."])
        if rtype == "MX":
            return DnsAnswer([f"10 mx.{domain}."])
        if rtype == "TXT":
            return DnsAnswer(["v=DMARC1; p=none"] if name.startswith("_dmarc.") else ["v=spf1 -all"])
        return DnsAnswer(["192.0.2.1"])


def world_resolver(domain):
    kind = WORLD.get(domain)
    if kind == "servfail":
        return None
    return kind not in ("nxdomain", "no_mail", "null_mx")


class ScriptedElv:
    """EmailListVerify stand-in answering by local part; records every address it is asked."""

    name, paid, cost_per_check = "emaillistverify", True, 1.0

    def __init__(self, answers=None, default="ok"):
        self.answers = answers or {}
        self.default = default
        self.calls = []

    def check(self, email):
        from cloud.intel.email.providers import ELV_STATUS

        self.calls.append(email)
        code = self.answers.get(email.split("@")[0], self.default)
        if isinstance(code, Exception):
            raise code
        status, score = ELV_STATUS[code]
        return ValidationResult(email, status, score, self.name, {"result_code": code, "mapping_verified": True})


def world_platform():
    platform, ctx, _ = make_platform()
    email = EmailValidationService(platform, local=LocalValidator(resolver=world_resolver))
    platform.override("email", email)
    jobs = platform.service("email_jobs")
    jobs.dns_client = WorldDns()
    return platform, ctx, email, jobs


def run_rows(platform, ctx, jobs, emails, **settings):
    job = jobs.create_from_rows(ctx, name="t", rows=[{"email": e} for e in emails], email_field="email",
                                source_type="manual", start=True, settings=settings or None)
    done = run_task_inline(platform, ctx.workspace_id, job["task_id"])
    assert done["status"] == "completed", done.get("error")
    return jobs.get(ctx, job["id"]), jobs.items(ctx, job["id"], limit=500).rows


def enable_elv(platform, ctx, provider, balance=1000):
    registry = platform.service("providers")
    registry.set_credentials(ctx, "emaillistverify", {"api_key": "elv-test-key"})
    row = registry.connection(ctx, "emaillistverify")
    platform.store.update(ctx, "provider_connections", row["id"], {"status": "verified"})
    platform.service("credits").sync(ctx, "emaillistverify", balance, source="test")
    platform.service("email").paid_factory = lambda c: provider


# --- 1. the resolver ------------------------------------------------------------------------------

class ResolverTests(unittest.TestCase):
    STATUSES = ("VALID", "INVALID", "RISKY", "UNKNOWN", "DISPOSABLE", "ROLE", "FREE_PROVIDER")
    PROVIDERS = ("local", "emaillistverify", "other", None)

    def test_only_three_outcomes_and_valid_needs_a_mailbox_verifier(self) -> None:
        for status in self.STATUSES:
            for provider in self.PROVIDERS:
                final = final_status(status, provider)
                self.assertIn(final, ("VALID", "INVALID", "NOT_VERIFIED"))
                if final == "VALID":
                    self.assertEqual((status, provider), ("VALID", "emaillistverify"))

    def test_store_filters_select_exactly_what_the_resolver_says(self) -> None:
        from cloud.intel.store.base import parse_filters  # noqa: F401 - the filters are plain store filters

        def matches(filters, status, provider):
            for key, value in filters.items():
                field, _, op = key.partition("__")
                actual = {"status": status, "provider": provider}[field]
                if (op == "in" and actual not in value) or (not op and actual != value):
                    return False
            return True

        for status in self.STATUSES:
            for provider in ("local", "emaillistverify"):  # the only providers that write rows
                if (status, provider) == ("VALID", "local"):
                    continue  # never written: built-in checks have no VALID and strict mode demotes one
                wanted = final_status(status, provider)
                hits = [f for f in engine.FINAL_STATUSES if matches(final_filters(f), status, provider)]
                self.assertEqual(hits, [wanted], (status, provider))

    def test_the_built_in_layers_never_produce_valid(self) -> None:
        validator = LocalValidator(resolver=world_resolver)
        for email in ("ceo@acme.com", "john.smith@bigcorp.io", "x@gmail.com", "info@acme.com", "a@mailinator.com",
                      "noreply@acme.com", "test@acme.com", "user@example.com", "bad..dots@acme.com",
                      "a@gone-domain.com", "a@flaky-dns.com", "", "nonsense"):
            self.assertNotEqual(validator.check(email).status, "VALID", email)

    def test_max_age_zero_is_honoured(self) -> None:
        self.assertEqual((max_age(0), max_age(None), max_age("7"), max_age("x"), max_age(999)), (0, 30, 7, 30, 365))


# --- 2. the deterministic input matrix -----------------------------------------------------------

class MatrixTests(unittest.TestCase):
    """Synthetic addresses covering every evidence class. Mailbox verification in the second
    half comes from a scripted fake, clearly NOT a real provider answer."""

    MATRIX = {
        # address: (expected final without a verifier, why)
        "jane.doe@acme-corp.com": ("NOT_VERIFIED", "valid format + MX: not mailbox evidence"),
        "bad..dots@acme-corp.com": ("INVALID", "malformed"),
        "no-at-sign": ("INVALID", "malformed"),
        "john@gone-domain.com": ("INVALID", "NXDOMAIN"),
        "john@nomail-domain.com": ("INVALID", "no MX and no address"),
        "john@nullmx-domain.com": ("INVALID", "null MX"),
        "someone@gmail.com": ("NOT_VERIFIED", "free provider"),
        "temp@mailinator.com": ("NOT_VERIFIED", "disposable: a risk, not proof of invalidity"),
        "info@acme-corp.com": ("NOT_VERIFIED", "role account"),
        "no-reply@acme-corp.com": ("NOT_VERIFIED", "no-reply"),
        "test@acme-corp.com": ("NOT_VERIFIED", "placeholder-looking, but a real domain"),
        "user@example.com": ("INVALID", "RFC 2606 reserved domain"),
        "ann@flaky-dns.com": ("NOT_VERIFIED", "temporary DNS failure is never INVALID"),
    }

    def test_built_in_only(self) -> None:
        platform, ctx, _, jobs = world_platform()
        job, items = run_rows(platform, ctx, jobs, list(self.MATRIX))
        got = {i["row"]["email"]: i["final_status"] for i in items}
        for email, (final, why) in self.MATRIX.items():
            self.assertEqual(got[email], final, f"{email}: {why}")
        self.assertEqual(job["counts"]["final_valid"], 0)
        flaky = next(i for i in items if i["row"]["email"] == "ann@flaky-dns.com")
        dom = next(s for s in flaky["stages"] if s["stage"] == "domain")
        self.assertEqual((dom["status"], dom["error_code"], dom["confidence"]), ("unknown", "servfail", "low"))
        nullmx = next(i for i in items if i["row"]["email"] == "john@nullmx-domain.com")
        self.assertIn("null MX", nullmx["checks"]["domain_invalid"])

    def test_every_result_carries_inspectable_stage_evidence(self) -> None:
        platform, ctx, _, jobs = world_platform()
        _, items = run_rows(platform, ctx, jobs, ["jane.doe@acme-corp.com"])
        stages = {s["stage"]: s for s in items[0]["stages"]}
        for name in ("normalization", "syntax", "domain", "mx", "spf", "dmarc", "role", "disposable",
                     "free_provider", "no_reply", "mailbox", "final"):
            self.assertIn(name, stages)
            self.assertEqual(set(stages[name]), {"stage", "label", "status", "result", "confidence", "source",
                                                 "checked_at", "latency_ms", "error_code", "details"})
        self.assertEqual(stages["mx"]["status"], "pass")
        self.assertIn("says nothing about this mailbox", stages["mx"]["details"]["note"])
        self.assertEqual((stages["mailbox"]["status"], stages["final"]["result"]), ("skipped", "NOT_VERIFIED"))
        self.assertIsNotNone(stages["domain"]["latency_ms"])

    def test_with_a_scripted_verifier_only_its_ok_is_valid(self) -> None:
        platform, ctx, _, jobs = world_platform()
        elv = ScriptedElv({"jane.doe": "ok", "catchall": "ok_for_all", "dead": "email_disabled", "unk": "unknown",
                           "info": "ok", "someone": "ok"})
        enable_elv(platform, ctx, elv)
        emails = ["jane.doe@acme-corp.com", "catchall@acme-corp.com", "dead@acme-corp.com", "unk@acme-corp.com",
                  "info@acme-corp.com", "someone@gmail.com", "temp@mailinator.com", "john@gone-domain.com"]
        job, _ = run_rows(platform, ctx, jobs, emails)
        estimate = jobs.unknowns_estimate(ctx, job["id"])
        queued = jobs.verify_unknowns(ctx, job["id"], confirm=True, expected_credits=estimate["credits"])
        self.assertEqual(run_task_inline(platform, ctx.workspace_id, queued["task_id"])["status"], "completed")
        got = {i["row"]["email"]: i["final_status"] for i in jobs.items(ctx, job["id"], limit=50).rows}
        self.assertEqual(got, {"jane.doe@acme-corp.com": "VALID", "catchall@acme-corp.com": "NOT_VERIFIED",
                               "dead@acme-corp.com": "INVALID", "unk@acme-corp.com": "NOT_VERIFIED",
                               "info@acme-corp.com": "VALID", "someone@gmail.com": "VALID",
                               "temp@mailinator.com": "NOT_VERIFIED", "john@gone-domain.com": "INVALID"})
        # Disposable and confirmed-invalid addresses are never sent; nothing beyond the estimate is.
        self.assertNotIn("temp@mailinator.com", elv.calls)
        self.assertNotIn("john@gone-domain.com", elv.calls)
        self.assertEqual(len(elv.calls), estimate["to_check"])


# --- 3. de-duplication ---------------------------------------------------------------------------

class CountingValidator(LocalValidator):
    def __init__(self):
        super().__init__(resolver=world_resolver)
        self.checked = []

    def _check(self, email):
        self.checked.append(email)
        return super()._check(email)


class DedupTests(unittest.TestCase):
    def setUp(self) -> None:
        self.platform, self.ctx, _, self.jobs = world_platform()
        self.local = CountingValidator()
        self.email = EmailValidationService(self.platform, local=self.local)
        self.platform.override("email", self.email)

    def test_the_same_address_twice_and_a_hundred_times_is_checked_once(self) -> None:
        for copies in (2, 100):
            self.local.checked.clear()
            address = f"dup{copies}@acme-corp.com"
            job, items = run_rows(self.platform, self.ctx, self.jobs, [address] * copies)
            self.assertEqual(self.local.checked, [address])
            self.assertEqual(job["counts"]["processed"], copies)  # every row still gets its result
            self.assertEqual({i["final_status"] for i in items}, {"NOT_VERIFIED"})

    def test_case_and_whitespace_variants_are_one_address(self) -> None:
        variants = ["Jane.Doe@Acme-Corp.com", " jane.doe@acme-corp.com ", "JANE.DOE@ACME-CORP.COM",
                    "jane.doe@acme-corp.com\t", "mailto:jane.doe@acme-corp.com", "Jane <jane.doe@acme-corp.com>"]
        job, items = run_rows(self.platform, self.ctx, self.jobs, variants)
        self.assertEqual(self.local.checked, ["jane.doe@acme-corp.com"])
        self.assertEqual({i["email"] for i in items}, {"jane.doe@acme-corp.com"})

    def test_separate_jobs_reuse_the_cache_and_never_pay_twice(self) -> None:
        elv = ScriptedElv()
        enable_elv(self.platform, self.ctx, elv)
        job, _ = run_rows(self.platform, self.ctx, self.jobs, ["pay.once@acme-corp.com"])
        estimate = self.jobs.unknowns_estimate(self.ctx, job["id"])
        queued = self.jobs.verify_unknowns(self.ctx, job["id"], confirm=True, expected_credits=estimate["credits"])
        run_task_inline(self.platform, self.ctx.workspace_id, queued["task_id"])
        self.assertEqual(elv.calls, ["pay.once@acme-corp.com"])
        self.local.checked.clear()
        job2, items = run_rows(self.platform, self.ctx, self.jobs, ["PAY.ONCE@acme-corp.com"])
        self.assertEqual((items[0]["final_status"], items[0]["cached"]), ("VALID", True))
        self.assertEqual(self.local.checked, [])  # served from the cache
        self.assertEqual(self.jobs.unknowns_estimate(self.ctx, job2["id"])["to_check"], 0)
        self.assertEqual(elv.calls, ["pay.once@acme-corp.com"])
        self.assertEqual(self.platform.service("credits").balance(self.ctx, "emaillistverify")["consumed"], 1)


# --- 4. paid checks: explicit, exact, capped ---------------------------------------------------------

class PaidSafetyTests(unittest.TestCase):
    def setUp(self) -> None:
        self.platform, self.ctx, self.email, self.jobs = world_platform()
        self.elv = ScriptedElv()
        enable_elv(self.platform, self.ctx, self.elv)

    def test_a_job_can_never_start_paid(self) -> None:
        job = self.jobs.create_from_rows(self.ctx, name="t", rows=[{"email": "a@acme-corp.com"}], email_field="email")
        with self.assertRaises(ValidationError):
            self.jobs.start(self.ctx, job["id"], settings={"allow_paid": True})
        before = self.platform.store.count(self.ctx, "email_validation_jobs", {})
        with self.assertRaises(ValidationError):  # refused before anything is created
            self.jobs.create_from_rows(self.ctx, name="t", rows=[{"email": "a@acme-corp.com"}], email_field="email",
                                       start=True, settings={"allow_paid": True})
        self.assertEqual(self.platform.store.count(self.ctx, "email_validation_jobs", {}), before)
        self.assertEqual(self.elv.calls, [])

    def test_the_confirmed_budget_is_a_hard_cap(self) -> None:
        emails = [f"person{i}@acme-corp.com" for i in range(10)]
        job, _ = run_rows(self.platform, self.ctx, self.jobs, emails)
        estimate = self.jobs.unknowns_estimate(self.ctx, job["id"])
        self.assertEqual(estimate["credits"], 10)
        queued = self.jobs.verify_unknowns(self.ctx, job["id"], confirm=True, expected_credits=10)
        self.assertEqual(queued["settings"]["paid_budget"], 10)
        # Simulate an estimate that turned out low: only 4 of the confirmed credits are left.
        settings = {**queued["settings"], "paid_used": 6}
        self.platform.store.update(self.ctx, "email_validation_jobs", job["id"], {"settings": settings})
        run_task_inline(self.platform, self.ctx.workspace_id, queued["task_id"])
        self.assertEqual(len(self.elv.calls), 4)
        items = self.jobs.items(self.ctx, job["id"], limit=50).rows
        skipped = [i for i in items if "budget" in str((i["checks"] or {}).get("paid_skipped"))]
        self.assertEqual(len(skipped), 6)
        self.assertTrue(all(i["final_status"] == "NOT_VERIFIED" for i in skipped))
        final = self.jobs.get(self.ctx, job["id"])
        self.assertEqual((final["settings"]["allow_paid"], final["settings"]["paid_budget"]), (False, None))

    def test_a_refused_credit_reservation_does_not_fail_the_job(self) -> None:
        job, _ = run_rows(self.platform, self.ctx, self.jobs, ["a@acme-corp.com", "b@acme-corp.com"])
        estimate = self.jobs.unknowns_estimate(self.ctx, job["id"])
        queued = self.jobs.verify_unknowns(self.ctx, job["id"], confirm=True, expected_credits=estimate["credits"])
        self.platform.service("credits").set_hard_limit(self.ctx, "emaillistverify", 0)  # changed meanwhile
        done = run_task_inline(self.platform, self.ctx.workspace_id, queued["task_id"])
        self.assertEqual(done["status"], "completed", done.get("error"))
        items = self.jobs.items(self.ctx, job["id"], limit=50).rows
        self.assertEqual({i["checks"].get("paid_error_code") for i in items}, {"no_credits"})
        self.assertEqual(self.elv.calls, [])
        # Still unresolved, so a later Verify Unknowns can settle them.
        self.assertEqual(self.jobs.unknowns_estimate(self.ctx, job["id"])["unresolved"], 2)

    def test_repeated_provider_failures_trip_the_breaker(self) -> None:
        from cloud.intel.email import service as service_module

        session = FakeSession({"verifyEmail": FakeResponse(503, text="")})
        provider = EmailListVerifyProvider("k", session=session, sleep=lambda s: None, max_retries=0)
        self.email.paid_factory = lambda c: provider
        emails = [f"p{i}@acme-corp.com" for i in range(12)]
        results = self.email.validate(self.ctx, emails, allow_paid=True, on_credit_error="record")
        self.assertEqual(len(session.calls), service_module.BREAKER_FAILURES)
        self.assertTrue(all(r["final_status"] == "NOT_VERIFIED" for r in results))
        # The workspace cools down: the next call does not hit the failing provider at all.
        again = self.email.validate(self.ctx, ["q@acme-corp.com"], allow_paid=True)
        self.assertEqual(len(session.calls), service_module.BREAKER_FAILURES)
        self.assertIn("paused", again[0]["checks"]["paid_skipped"])

    def test_possibly_charged_failures_are_counted(self) -> None:
        session = FakeSession({"verifyEmail": FakeResponse(200, text="<html>maintenance</html>")})
        self.email.paid_factory = lambda c: EmailListVerifyProvider("k", session=session, sleep=lambda s: None)
        [result] = self.email.validate(self.ctx, ["m@acme-corp.com"], allow_paid=True)
        self.assertEqual((result["status"], result["provider"]), ("UNKNOWN", "local"))  # still re-checkable
        self.assertTrue(result["checks"]["paid_charge_uncertain"])
        self.assertEqual(self.platform.service("credits").balance(self.ctx, "emaillistverify")["consumed"], 1)


class QuickValidateApiTests(unittest.TestCase):
    def setUp(self) -> None:
        from fastapi.testclient import TestClient

        from cloud.api.auth import DevTokenIssuer
        from cloud.api.main import create_app
        from cloud.api.settings import Settings
        from cloud.intel.platform import Platform, PlatformConfig
        from cloud.intel.store.memory import MemoryStore
        from cloud.shared.storage import LocalFileStorage
        from cloud.worker.dispatcher import NullDispatcher

        scratch = tempfile.TemporaryDirectory()
        self.addCleanup(scratch.cleanup)
        root = Path(scratch.name)
        self.platform = Platform(MemoryStore(), storage=LocalFileStorage(root / "platform"),
                                 config=PlatformConfig(files_dir=root / "platform"))
        self.email = EmailValidationService(self.platform, local=LocalValidator(resolver=world_resolver))
        self.platform.override("email", self.email)
        self.platform.service("email_jobs").dns_client = WorldDns()
        issuer = DevTokenIssuer("email-hardening-tests-secret-0123456789ab")
        app = create_app(Settings(auth_mode="dev", results_dir=root / "results"),
                         storage=LocalFileStorage(root / "results"), token_verifier=issuer,
                         dispatcher=NullDispatcher(), platform=self.platform)
        self.client = TestClient(app)
        self.client.headers["Authorization"] = f"Bearer {issuer.issue('alice@example.com')['access_token']}"
        self.client.__enter__()
        self.addCleanup(self.client.__exit__, None, None, None)
        self.ws = self.client.post("/api/v1/workspaces", json={"name": "Alice", "seed": False}).json()["id"]
        self.base = f"/api/v1/w/{self.ws}"
        self.elv = ScriptedElv()

    def test_paid_needs_the_exact_confirmed_estimate(self) -> None:
        self.email.paid_factory = lambda c: self.elv
        emails = ["a@acme-corp.com", "b@acme-corp.com", "x@gone-domain.com", "info@acme-corp.com"]
        estimate = self.client.post(self.base + "/email/validate/estimate", json={"emails": emails}).json()
        self.assertEqual((estimate["to_check"], estimate["decided_free"]), (2, 2))
        refused = self.client.post(self.base + "/email/validate", json={"emails": emails, "allow_paid": True})
        self.assertEqual(refused.status_code, 409, refused.text)
        self.assertEqual(self.elv.calls, [])
        wrong = self.client.post(self.base + "/email/validate", json={"emails": emails, "allow_paid": True,
                                                                      "confirm": True, "expected_credits": 1})
        self.assertEqual(wrong.status_code, 409)
        self.assertEqual(self.elv.calls, [])
        too_many = self.client.post(self.base + "/email/validate",
                                    json={"emails": [f"u{i}@acme-corp.com" for i in range(30)], "allow_paid": True})
        self.assertEqual(too_many.status_code, 422)
        free = self.client.post(self.base + "/email/validate", json={"emails": emails})
        self.assertEqual(free.status_code, 200, free.text)
        self.assertEqual({i["final_status"] for i in free.json()["items"]}, {"NOT_VERIFIED", "INVALID"})
        self.assertEqual(self.elv.calls, [])


# --- 5. failures, recovery, restarts ---------------------------------------------------------------

class FailureTests(unittest.TestCase):
    def setUp(self) -> None:
        self.platform, self.ctx, self.email, self.jobs = world_platform()

    def test_one_bad_row_never_fails_the_chunk(self) -> None:
        original = engine.contact_signals

        def explode(email, contact, **kw):
            if email.startswith("boom"):
                raise RuntimeError("bad row")
            return original(email, contact, **kw)

        rows = [{"email": "boom@acme-corp.com", "first_name": "B"}, {"email": "fine@acme-corp.com", "first_name": "F"}]
        with mock.patch.object(engine, "contact_signals", explode):
            job = self.jobs.create_from_rows(self.ctx, name="t", rows=rows, email_field="email", start=True)
            done = run_task_inline(self.platform, self.ctx.workspace_id, job["task_id"])
        self.assertEqual(done["status"], "completed", done.get("error"))
        items = {i["row"]["email"]: i for i in self.jobs.items(self.ctx, job["id"]).rows}
        self.assertIn("evidence_error", items["boom@acme-corp.com"]["checks"])
        self.assertEqual(items["fine@acme-corp.com"]["final_status"], "NOT_VERIFIED")
        self.assertNotIn("evidence_error", items["fine@acme-corp.com"]["checks"])

    def test_a_built_in_check_that_crashes_is_unknown_not_a_failed_batch(self) -> None:
        validator = LocalValidator(resolver=lambda d: (_ for _ in ()).throw(RuntimeError("resolver bug")))
        result = validator.check("a@acme-corp.com")
        self.assertEqual(result.status, "UNKNOWN")

        class Broken(LocalValidator):
            def _check(self, email):
                raise KeyError("boom")

        broken = Broken().check("a@acme-corp.com")
        self.assertEqual((broken.status, final_status(broken.status, broken.provider)), ("UNKNOWN", "NOT_VERIFIED"))
        self.assertIn("built-in check failed", broken.checks["error"])

    def test_a_transient_database_error_is_retried_inside_the_run(self) -> None:
        store = self.platform.store
        original = store.update_many
        failures = []

        def flaky(ctx, entity, rows):
            if entity == "email_validation_items" and not failures:
                failures.append(1)
                raise ConnectionError("server closed the connection unexpectedly")
            return original(ctx, entity, rows)

        with mock.patch.object(store, "update_many", flaky), mock.patch.object(jobs_module, "CHUNK_BACKOFF", 0.0):
            job = self.jobs.create_from_rows(self.ctx, name="t", rows=[{"email": "a@acme-corp.com"}],
                                             email_field="email", start=True)
            done = run_task_inline(self.platform, self.ctx.workspace_id, job["task_id"])
        self.assertEqual((done["status"], done["attempts"], failures), ("completed", 1, [1]))

    def test_an_interrupted_run_resumes_and_never_pays_twice(self) -> None:
        elv = ScriptedElv()
        enable_elv(self.platform, self.ctx, elv)
        emails = [f"r{i}@acme-corp.com" for i in range(250)]
        job, _ = run_rows(self.platform, self.ctx, self.jobs, emails)
        estimate = self.jobs.unknowns_estimate(self.ctx, job["id"])
        queued = self.jobs.verify_unknowns(self.ctx, job["id"], confirm=True, expected_credits=estimate["credits"])
        calls = {"n": 0}
        original = self.jobs._process  # noqa: SLF001

        def crash_after_first_chunk(*a, **k):
            calls["n"] += 1
            if calls["n"] == 2:
                raise SystemExit("worker killed")  # not an Exception: no in-run retry, like a dead worker
            return original(*a, **k)

        with mock.patch.object(self.jobs, "_process", crash_after_first_chunk):
            with self.assertRaises(SystemExit):
                run_task_inline(self.platform, self.ctx.workspace_id, queued["task_id"])
        self.assertEqual(len(elv.calls), jobs_module.WRITE_CHUNK)
        # The lease expires and the task is picked up again (here: reclaimed directly).
        task = self.platform.store.get(self.ctx.as_system(), "platform_tasks", queued["task_id"])
        self.platform.store.update(self.ctx.as_system(), "platform_tasks", task["id"],
                                   {"status": "queued", "worker_id": None, "lease_expires_at": None})
        done = run_task_inline(self.platform, self.ctx.workspace_id, queued["task_id"])
        self.assertEqual(done["status"], "completed", done.get("error"))
        self.assertEqual(sorted(elv.calls), sorted(emails))  # every address exactly once
        self.assertEqual(self.jobs.get(self.ctx, job["id"])["counts"]["final_valid"], 250)

    def test_a_failed_job_can_be_retried_from_where_it_stopped(self) -> None:
        job = self.jobs.create_from_rows(self.ctx, name="t", rows=[{"email": f"f{i}@acme-corp.com"} for i in range(5)],
                                         email_field="email", start=True)
        self.platform.store.update(self.ctx, "email_validation_jobs", job["id"], {"status": "failed",
                                                                                   "error": "attempts exhausted"})
        self.platform.store.update(self.ctx.as_system(), "platform_tasks", job["task_id"], {"status": "failed"})
        retried = self.jobs.retry(self.ctx, job["id"])
        self.assertEqual(retried["status"], "queued")
        self.assertEqual(run_task_inline(self.platform, self.ctx.workspace_id, retried["task_id"])["status"],
                         "completed")
        self.assertEqual(self.jobs.get(self.ctx, job["id"])["counts"]["processed"], 5)
        with self.assertRaises(ConflictError):
            self.jobs.retry(self.ctx, job["id"])  # completed jobs are not retried

    def test_dns_that_failed_transiently_is_asked_again_later_in_a_long_run(self) -> None:
        dns = WorldDns()
        run = jobs_module._EvidenceRun(self.jobs, self.ctx, {})  # noqa: SLF001
        run.dns = dns
        self.assertEqual(run.domain("flaky-dns.com")["dns_error"], "servfail")
        run.domain("flaky-dns.com")
        self.assertEqual(len([c for c in dns.calls if c == ("flaky-dns.com", "MX")]), 1)  # reused briefly
        run.transient["flaky-dns.com"] -= jobs_module.DNS_TRANSIENT_SECONDS + 1
        WORLD["flaky-dns.com"] = None
        try:
            self.assertTrue(run.domain("flaky-dns.com")["mail"])  # recovered
        finally:
            WORLD["flaky-dns.com"] = "servfail"

    def test_the_built_in_domain_cache_expires_and_is_bounded(self) -> None:
        from cloud.intel.email import providers

        now = [1000.0]
        answers = {"a.com": True}
        validator = LocalValidator(resolver=lambda d: answers.get(d, True), clock=lambda: now[0])
        validator.check("x@a.com")
        answers["a.com"] = False
        now[0] += providers.MX_TTL - 1
        self.assertIs(validator.check("x@a.com").checks["mx"], True)  # still the cached answer
        now[0] += 2
        self.assertIs(validator.check("x@a.com").checks["mx"], False)  # re-resolved after the TTL
        with mock.patch.object(providers, "MX_CACHE_MAX", 10):
            for i in range(50):
                validator.seed(f"d{i}.com", True)
            self.assertLessEqual(len(validator._mx_cache), 10)  # noqa: SLF001

    def test_empty_inputs_are_refused_cleanly(self) -> None:
        for call in (lambda: self.jobs.create_from_rows(self.ctx, name="t", rows=[], email_field="email"),
                     lambda: self.jobs.create_upload(self.ctx, "empty.csv", b""),
                     lambda: self.jobs.create_upload(self.ctx, "header-only.csv", b"Email\n")):
            with self.assertRaises(ValidationError):
                call()

    def test_a_malformed_xlsx_upload_is_a_clear_error(self) -> None:
        with self.assertRaises(ValidationError):
            self.jobs.create_upload(self.ctx, "broken.xlsx", b"PK\x03\x04 this is not a workbook")


class BackgroundTests(unittest.TestCase):
    def setUp(self) -> None:
        from cloud.shared.storage import LocalFileStorage

        scratch = tempfile.TemporaryDirectory()
        self.addCleanup(scratch.cleanup)
        self.platform, self.ctx, self.email, self.jobs = world_platform()
        self.platform.storage = LocalFileStorage(Path(scratch.name))
        for patch in (mock.patch.object(jobs_module, "INLINE_INGEST_BYTES", 100),
                      mock.patch.object(jobs_module, "INGEST_CHUNK", 100),
                      mock.patch.object(jobs_module, "INLINE_ROWS", 50)):
            patch.start()
            self.addCleanup(patch.stop)

    def test_a_large_paste_is_read_in_the_background_and_starts_by_itself(self) -> None:
        rows = [{"email": f"p{i}@acme-corp.com"} for i in range(230)]
        job = self.jobs.create_from_rows(self.ctx, name="Pasted", rows=rows, email_field="email",
                                         source_type="manual", start=True, settings={"max_age_days": 30})
        self.assertEqual((job["status"], job["row_count"], job["settings"]["ingest"]["state"]), ("ready", 0, "pending"))
        self.assertEqual(self.platform.store.count(self.ctx, "email_validation_items", {"job_id": job["id"]}), 0)
        done = run_task_inline(self.platform, self.ctx.workspace_id, job["settings"]["ingest"]["task_id"])
        self.assertEqual(done["status"], "completed", done.get("error"))
        job = self.jobs.get(self.ctx, job["id"])
        self.assertEqual((job["status"], job["row_count"]), ("queued", 230))  # started automatically
        self.assertEqual(run_task_inline(self.platform, self.ctx.workspace_id, job["task_id"])["status"], "completed")
        counts = self.jobs.get(self.ctx, job["id"])["counts"]
        self.assertEqual((counts["processed"], counts["final_not_verified"]), (230, 230))

    def test_a_damaged_csv_keeps_the_rows_before_the_damage(self) -> None:
        # Damage after the 200-row sample the upload request reads: found by the background reader.
        good = "".join(f"d{i}@acme-corp.com\n" for i in range(250))
        data = ("Email\n" + good + '"unterminated' + "x" * 200_000 + "\n").encode()
        job = self.jobs.create_upload(self.ctx, "damaged.csv", data)
        done = run_task_inline(self.platform, self.ctx.workspace_id, job["settings"]["ingest"]["task_id"])
        self.assertEqual(done["status"], "completed", done.get("error"))
        job = self.jobs.get(self.ctx, job["id"])
        self.assertEqual(job["row_count"], 250)
        self.assertTrue(any("stopped being readable" in p for p in job["settings"]["problems"]))

    def test_damage_inside_the_sample_is_a_clean_upload_error(self) -> None:
        with self.assertRaises(ValidationError):  # a 422 for the API, never a 500
            self.jobs.create_upload(self.ctx, "garbage.csv", b"Email\n" + b'"' + b"x" * 300_000 + b"\n")

    def test_an_unreadable_file_fails_the_job_visibly_and_permanently(self) -> None:
        job = self.jobs.create_upload(self.ctx, "ok.csv", ("Email\n" + "a@acme-corp.com\n" * 40).encode())
        key = job["settings"]["ingest"]["storage_key"]
        with tempfile.TemporaryDirectory() as scratch:  # the stored copy is damaged before it is read
            bad = Path(scratch) / "bad.csv"
            bad.write_bytes(b"Email\n" + b'"' + b"x" * 300_000 + b"\n")
            self.platform.storage.put_file(key, bad, content_type="text/csv")
        done = run_task_inline(self.platform, self.ctx.workspace_id, job["settings"]["ingest"]["task_id"])
        self.assertEqual((done["status"], done["attempts"]), ("failed", 1))  # not retried pointlessly
        job = self.jobs.get(self.ctx, job["id"])
        self.assertEqual((job["status"], job["settings"]["ingest"]["state"]), ("failed", "failed"))
        self.assertIn("could not be read", job["error"])

    def test_a_dead_reader_task_does_not_leave_the_job_reading_forever(self) -> None:
        job = self.jobs.create_upload(self.ctx, "big.csv", ("Email\n" + "a@acme-corp.com\n" * 40).encode())
        task_id = job["settings"]["ingest"]["task_id"]
        self.platform.store.update(self.ctx.as_system(), "platform_tasks", task_id,
                                   {"status": "failed", "error": "worker stopped heartbeating"})
        job = self.jobs.get(self.ctx, job["id"])
        self.assertEqual((job["status"], job["settings"]["ingest"]["state"]), ("failed", "failed"))
        retried = self.jobs.retry(self.ctx, job["id"])  # the stored file is read again
        self.assertEqual(retried["settings"]["ingest"]["state"], "pending")
        done = run_task_inline(self.platform, self.ctx.workspace_id, retried["settings"]["ingest"]["task_id"])
        self.assertEqual(done["status"], "completed", done.get("error"))
        self.assertEqual(self.jobs.get(self.ctx, job["id"])["row_count"], 40)


# --- 6. tenant isolation -----------------------------------------------------------------------------

class IsolationTests(unittest.TestCase):
    def test_jobs_results_and_caches_stay_in_their_workspace(self) -> None:
        platform, ctx_a, email, jobs = world_platform()
        elv = ScriptedElv()
        enable_elv(platform, ctx_a, elv)
        job, _ = run_rows(platform, ctx_a, jobs, ["shared.person@acme-corp.com"])
        estimate = jobs.unknowns_estimate(ctx_a, job["id"])
        queued = jobs.verify_unknowns(ctx_a, job["id"], confirm=True, expected_credits=estimate["credits"])
        run_task_inline(platform, ctx_a.workspace_id, queued["task_id"])
        owner_b = str(uuid.uuid4())
        ws_b = platform.store.create_workspace(owner_b, "Other", f"other-{uuid.uuid4().hex[:8]}")
        ctx_b = Ctx(ws_b["id"], owner_b, "owner")
        from cloud.intel.core.context import NotFoundError

        with self.assertRaises(NotFoundError):
            jobs.get(ctx_b, job["id"])
        with self.assertRaises(NotFoundError):
            jobs.items(ctx_b, job["id"])
        self.assertEqual(platform.store.count(ctx_b, "email_validations", {}), 0)
        self.assertEqual(platform.store.count(ctx_b, "email_check_cache", {}), 0)
        # B validating the same address gets its own (built-in) answer, never A's paid VALID.
        [result] = email.validate(ctx_b, ["shared.person@acme-corp.com"])
        self.assertEqual((result["cached"], result["final_status"]), (False, "NOT_VERIFIED"))


# --- 7. contacts, exports, parsing ---------------------------------------------------------------------

class ContactAndExportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.platform, self.ctx, self.email, self.jobs = world_platform()

    def test_contact_status_follows_the_resolver(self) -> None:
        store = self.platform.store
        disposable = store.insert(self.ctx, "contacts", {"full_name": "T", "email": "temp@mailinator.com"})
        dead = store.insert(self.ctx, "contacts", {"full_name": "D", "email": "d@gone-domain.com"})
        self.email.validate(self.ctx, ["temp@mailinator.com", "d@gone-domain.com"])
        self.assertEqual(store.get(self.ctx, "contacts", disposable["id"])["validation_status"], "needs_verification")
        self.assertEqual(store.get(self.ctx, "contacts", dead["id"])["validation_status"], "rejected")
        self.assertEqual(store.get(self.ctx, "contacts", disposable["id"])["email_status"], "DISPOSABLE")  # evidence

    def test_streamed_csv_equals_the_whole_file_and_is_filtered_by_final_status(self) -> None:
        job, _ = run_rows(self.platform, self.ctx, self.jobs,
                          ["a@acme-corp.com", "=HYPERLINK(1)@x", "b@gone-domain.com"])
        _, whole, _ = self.jobs.export(self.ctx, job["id"], "csv")
        _, parts, _ = self.jobs.export(self.ctx, job["id"], "csv", stream=True)
        self.assertEqual(b"".join(parts), whole)
        rows = list(csv.reader(io.StringIO(whole.decode("utf-8-sig"))))
        self.assertTrue(any(cell.startswith("'=") for row in rows for cell in row))  # formula neutralised
        _, invalid, _ = self.jobs.export(self.ctx, job["id"], "csv", final="INVALID")
        lines = invalid.decode("utf-8-sig").strip().splitlines()
        self.assertEqual(len(lines), 3)  # header + 2 invalid rows
        with self.assertRaises(ValidationError):
            self.jobs.export(self.ctx, job["id"], "csv", final="MAYBE")

    def test_xlsx_cells_are_never_formulas(self) -> None:
        from openpyxl import load_workbook

        job, _ = run_rows(self.platform, self.ctx, self.jobs, ["=cmd|' /C calc'!A0"])
        _, content, _ = self.jobs.export(self.ctx, job["id"], "xlsx")
        sheet = load_workbook(io.BytesIO(content)).active
        cell = sheet.cell(row=2, column=1)
        self.assertEqual((cell.value, cell.data_type), ("=cmd|' /C calc'!A0", "s"))

    def test_items_can_be_filtered_by_final_status(self) -> None:
        job, _ = run_rows(self.platform, self.ctx, self.jobs, ["a@acme-corp.com", "b@gone-domain.com", "c@gmail.com"])
        for final, n in (("VALID", 0), ("INVALID", 1), ("NOT_VERIFIED", 2), ("not verified", 2)):
            self.assertEqual(self.jobs.items(self.ctx, job["id"], {"final_status": final}).total, n, final)
        with self.assertRaises(ValidationError):
            self.jobs.items(self.ctx, job["id"], {"final_status": "RISKY"})

    def test_a_superseded_placeholder_invalid_in_the_cache_is_checked_again(self) -> None:
        from datetime import timedelta

        from cloud.intel.core.context import utcnow

        now = utcnow()
        self.platform.store.insert(self.ctx, "email_validations", {
            "email": "test@acme-corp.com", "status": "INVALID", "score": 0.0, "provider": "local",
            "checks": {"syntax": True, "placeholder": True}, "validated_at": now, "expires_at": now + timedelta(days=30)})
        [result] = self.email.validate(self.ctx, ["test@acme-corp.com"])
        self.assertEqual((result["cached"], result["final_status"]), (False, "NOT_VERIFIED"))


class StreamingParserTests(unittest.TestCase):
    def test_streamed_rows_match_the_in_memory_reader(self) -> None:
        from openpyxl import Workbook

        from cloud.intel.imports.parse import iter_rows, iter_rows_stream

        text = "Email;Name\r\nä@x.com;Jörg\r\n\r\nb@y.com;\"Smith; Jr\"\r\n"
        for encoding in ("utf-8", "utf-8-sig", "cp1252"):
            data = text.encode(encoding)
            self.assertEqual(list(iter_rows_stream("csv", io.BytesIO(data))), list(iter_rows("csv", data)), encoding)
        workbook = Workbook()
        workbook.active.append(["Email", "N"])
        workbook.active.append(["a@x.com", 1.0])
        workbook.active.append([None, None])
        workbook.active.append(["b@x.com", "two"])
        buffer = io.BytesIO()
        workbook.save(buffer)
        data = buffer.getvalue()
        self.assertEqual(list(iter_rows_stream("xlsx", io.BytesIO(data))), list(iter_rows("xlsx", data)))

    def test_malformed_csv_is_a_parse_error(self) -> None:
        from cloud.intel.imports.parse import ParseError, iter_rows_stream

        with self.assertRaises(ParseError):
            list(iter_rows_stream("csv", io.BytesIO(b'Email\n"' + b"x" * 300_000)))


if __name__ == "__main__":
    unittest.main()
