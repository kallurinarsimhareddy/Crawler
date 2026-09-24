"""Track A: exports — formats, provenance columns, formula neutralisation, async task."""

from __future__ import annotations

import csv
import io
import json
import tempfile
import unittest
import uuid
from pathlib import Path

from openpyxl import load_workbook

from cloud.intel.core.context import Ctx, ValidationError
from cloud.intel.platform import Platform, PlatformConfig
from cloud.intel.store.memory import MemoryStore
from cloud.intel.tasks.worker import run_task_inline
from cloud.shared.storage import LocalFileStorage


class ExportTests(unittest.TestCase):
    def setUp(self) -> None:
        scratch = tempfile.TemporaryDirectory()
        self.addCleanup(scratch.cleanup)
        store = MemoryStore()
        user = str(uuid.uuid4())
        ws = store.create_workspace(user, "Exports", f"exp-{uuid.uuid4().hex[:8]}")
        self.ctx = Ctx(ws["id"], user, "owner")
        self.platform = Platform(store, storage=LocalFileStorage(Path(scratch.name)),
                                 config=PlatformConfig(files_dir=Path(scratch.name)))
        crm = self.platform.service("crm")
        self.company = crm.upsert_company(self.ctx, {"name": "=HYPERLINK(\"http://evil\")", "domain": "evil.com",
                                                     "technologies": ["SAP", "AS400"]},
                                          source_kind="import", source_name="list.csv")["company"]
        crm.upsert_company(self.ctx, {"name": "Plain Co", "domain": "plain.com"}, source_kind="manual",
                           source_name="user")
        self.exports = self.platform.service("exports")

    def _read(self, record) -> bytes:
        with self.platform.storage.open(record["storage_key"]) as handle:
            return handle.read()

    def test_csv_neutralises_formulas_and_carries_provenance(self) -> None:
        record = self.exports.export_entity(self.ctx, "companies", {"domain": "evil.com"}, "csv")
        self.assertEqual((record["status"], record["row_count"]), ("completed", 1))
        rows = list(csv.DictReader(io.StringIO(self._read(record).decode("utf-8-sig"))))
        self.assertEqual(rows[0]["name"], "'=HYPERLINK(\"http://evil\")")
        self.assertEqual(rows[0]["technologies"], "SAP; AS400")
        self.assertEqual(rows[0]["_source_kinds"], "import")
        self.assertEqual(rows[0]["_source_names"], "list.csv")
        self.assertEqual(rows[0]["_workspace_id"], self.ctx.workspace_id)
        self.assertTrue(rows[0]["_exported_at"])
        self.assertIn("first_seen_at", rows[0])

    def test_xlsx_and_json(self) -> None:
        xlsx = self.exports.export_entity(self.ctx, "companies", {}, "xlsx")
        sheet = load_workbook(io.BytesIO(self._read(xlsx)), read_only=True).active
        values = list(sheet.iter_rows(values_only=True))
        self.assertEqual(len(values), 3)
        self.assertTrue(any(str(v).startswith("'=") for row in values for v in row if v))
        js = self.exports.export_entity(self.ctx, "companies", {"domain": "plain.com"}, "json")
        payload = json.loads(self._read(js))
        self.assertEqual(payload["rows"][0]["name"], "Plain Co")
        self.assertEqual(payload["rows"][0]["_source_kinds"], ["manual"])

    def test_list_export_async_task_and_validation(self) -> None:
        crm = self.platform.service("crm")
        target = crm.create_list(self.ctx, "Targets", "companies")
        crm.add_to_list(self.ctx, target["id"], "companies", [self.company["id"]], reason="ERP hiring")
        queued = self.exports.export_entity(self.ctx, "list", {"list_id": target["id"]}, "csv", async_=True)
        self.assertEqual(queued["status"], "queued")
        run_task_inline(self.platform, self.ctx.workspace_id, queued["task_id"])
        done = self.platform.store.get(self.ctx, "exports", queued["id"])
        self.assertEqual((done["status"], done["row_count"]), ("completed", 1))
        rows = list(csv.DictReader(io.StringIO(self._read(done).decode("utf-8-sig"))))
        self.assertEqual(rows[0]["_added_reason"], "ERP hiring")
        with self.assertRaises(ValidationError):
            self.exports.export_entity(self.ctx, "audit_log", {}, "csv")
        with self.assertRaises(ValidationError):
            self.exports.export_entity(self.ctx, "companies", {}, "pdf")
        with self.assertRaises(ValidationError):
            self.exports.export_entity(self.ctx, "companies", {"secret_column": 1}, "csv")


if __name__ == "__main__":
    unittest.main()
