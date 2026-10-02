"""Job relevance (keyword workbook -> 0-100 score, HIGH / REVIEW / REJECT) and the JobSpy source
(gated boards, search pages, stored fields). The workbook here is a small synthetic one in the
same layout as IT_Crawler_Keywords.xlsx; the real workbook is user data and stays out of the repo."""

from __future__ import annotations

import io
import tempfile
import unittest
import uuid
from datetime import date
from pathlib import Path
from typing import Any, Dict, List

from cloud.intel.core.context import Ctx, ValidationError
from cloud.intel.job_monitor.jobspy_source import record_from_jobspy, validate_jobspy_filters
from cloud.intel.job_monitor.relevance import RelevanceEngine, parse_keyword_workbook
from cloud.intel.platform import Platform, PlatformConfig
from cloud.intel.store.memory import MemoryStore
from cloud.intel.tasks.worker import run_task_inline
from cloud.shared.storage import LocalFileStorage

CATEGORIES = {
    "Category Words (Broad Net)": ["ERP", "CRM", "MES", "DevOps", "Cybersecurity", "Agile"],
    "ERP Platforms": ["SAP", "Oracle", "Workday", "NetSuite", "JDEdwards", "Sage", "Epicor"],
    "ERP Modules (SAP Specific)": ["FICO", "ABAP", "MM", "SD", "S4HANA", "Basis"],
    "Cloud Platforms": ["AWS", "Azure"],
    "Data & Analytics": ["Snowflake", "PowerBI", "Spark"],
    "Programming Languages": ["Python", "Swift", "PL/SQL", "Assembler"],
    "Frontend Frameworks": ["NextJS"],
    "Cybersecurity Tools": ["CrowdStrike", "Splunk"],
    "Collaboration & Project Management": ["Monday"],
    "Testing Tools": ["Selenium"],
}


def workbook_bytes(extra_uncategorized: bool = True) -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    flat = book.active
    flat.title = "All Keywords"
    flat.append(["#", "Keyword", "Category"])
    n = 0
    for cat, words in list(CATEGORIES.items())[:-1]:            # Testing Tools only in By Category
        for w in words:
            n += 1
            flat.append([n, w, cat])
    if extra_uncategorized:
        flat.append([n + 1, "Salesforce", None])
    by = book.create_sheet("By Category")
    by.append(["IT Job Crawler Keywords"])
    for cat, words in CATEGORIES.items():
        by.append([f"  {cat}  ({len(words)} keywords)"])
        by.append(words)
    summary = book.create_sheet("Summary")
    summary.append(["IT Keyword Categories Summary"])
    summary.append(["Category", "Keyword Count", "% of Total"])
    for cat in list(CATEGORIES) + ["CRM Platforms"]:
        summary.append([cat, 1, 1.0])
    summary.append(["TOTAL", 1, "100%"])
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


def engine() -> RelevanceEngine:
    return RelevanceEngine(parse_keyword_workbook(workbook_bytes())["keywords"])


class WorkbookTests(unittest.TestCase):
    def test_parse_merges_sheets_and_reports_problems(self) -> None:
        parsed = parse_keyword_workbook(workbook_bytes())
        words = {k["keyword"] for k in parsed["keywords"]}
        self.assertIn("Selenium", words)                       # only in By Category
        self.assertIn({"keyword": "Salesforce", "category": "Uncategorized"}, parsed["keywords"])
        self.assertTrue(any("Salesforce" in p for p in parsed["problems"]))
        self.assertTrue(any("CRM Platforms" in p for p in parsed["problems"]))
        self.assertEqual(len([k for k in parsed["keywords"] if k["keyword"] == "SAP"]), 1)   # no duplicates
        with self.assertRaises(ValueError):
            from openpyxl import Workbook

            empty = io.BytesIO()
            Workbook().save(empty)
            parse_keyword_workbook(empty.getvalue())


class MatchingTests(unittest.TestCase):
    def setUp(self) -> None:
        self.e = engine()

    def kws(self, title: str, description: str = "") -> List[str]:
        return self.e.score(title=title, description=description).matched_keywords

    def test_whole_words_only(self) -> None:
        self.assertEqual(self.kws("Enterprise Account Executive", "interpret enterprise data asap"), [])
        self.assertEqual(self.kws("Superpower Bistro"), [])
        self.assertNotIn("SAP", self.kws("Hiring ASAP"))
        self.assertNotIn("ERP", self.kws("erp lowercase noise"))          # acronyms are case-sensitive
        self.assertIn("ERP", self.kws("ERP Analyst"))
        self.assertIn("SAP", self.kws("SAP Consultant"))

    def test_spelling_variants(self) -> None:
        self.assertIn("S4HANA", self.kws("SAP Lead", "S/4HANA migration"))
        self.assertIn("PowerBI", self.kws("Analyst", "Power BI dashboards"))
        self.assertIn("NextJS", self.kws("Developer", "Next.js and React"))
        self.assertIn("JDEdwards", self.kws("JD Edwards Developer"))
        self.assertIn("PL/SQL", self.kws("Oracle Developer", "PL/SQL packages"))
        self.assertIn("PL/SQL", self.kws("Applications Manager", "PL-SQL, BI Publisher"))
        self.assertIn("Cybersecurity", self.kws("Analyst", "Cyber Security operations"))
        self.assertEqual(self.kws("Cyber Café attendant"), [])          # "cyber" alone is not the keyword
        self.assertEqual(self.e.score(title="Cyber Cafe attendant").classification, "REJECT")
        self.assertEqual(self.e.score(title="Cyber Security Analyst", description="Splunk").classification, "HIGH")

    def test_ambiguous_words_need_context(self) -> None:
        self.assertEqual(self.kws("Line Cook", "Monday to Friday, swift service"), [])
        self.assertIn("Monday", self.kws("Project Manager", "SAP rollout tracked in monday.com"))
        self.assertNotIn("Monday", self.kws("SAP Developer", "Schedule: Monday-Friday"))
        # equal-opportunity boilerplate is not the SAP Basis module
        self.assertEqual(self.kws("3D Print Technician", "no discrimination on the basis of race"), [])
        self.assertNotIn("Basis", self.kws("SAP Developer", "on the basis of experience"))
        self.assertIn("Basis", self.kws("SAP BASIS Specialist"))
        self.assertNotIn("MM", self.kws("Buyer", "MM and SD experience"))      # SAP modules need SAP
        self.assertIn("MM", self.kws("SAP Buyer", "MM and SD experience"))
        self.assertNotIn("Sage", self.kws("Sage advice for gardeners"))
        self.assertEqual(self.kws("Assembler", "Monday-Friday, 2nd shift"), [])     # the job, not the language
        self.assertIn("Assembler", self.kws("Embedded Developer", "x86 Assembler and Python"))

    def test_classification(self) -> None:
        cases = {
            ("SAP FICO Consultant", "S/4HANA, ABAP, MM"): "HIGH",
            ("Azure Cloud Engineer", "AWS, Python, Snowflake"): "HIGH",
            ("Data Engineer", "Snowflake, Spark, PowerBI, Python"): "REVIEW",
            ("CDL Truck Driver", "ASAP hiring, Monday start"): "REJECT",
            ("Registered Nurse", "Epic charting"): "REJECT",
            ("Security Officer", "patrol"): "REJECT",
            ("Enterprise Account Manager", "enterprise customers"): "REJECT",
        }
        for (title, desc), expected in cases.items():
            result = self.e.score(title=title, description=desc)
            self.assertEqual(result.classification, expected, (title, result.reason))
            self.assertTrue(0 <= result.score <= 100)
        noise = self.e.score(title="Warehouse Associate", description="SAP scanner, forklift")
        self.assertTrue(noise.reason.startswith("REJECT"))
        self.assertIn("noise", noise.reason)

    def test_categories_groups_and_search_term(self) -> None:
        r = self.e.score(title="SAP MES Engineer", description="plant floor, CrowdStrike", search_term="SAP")
        self.assertIn("ERP Platforms", r.matched_categories)
        self.assertIn("ERP / Enterprise Applications", r.groups)
        self.assertIn("Manufacturing Systems", r.groups)
        self.assertIn("search term 'SAP' in title", r.reason)
        self.assertEqual(r.as_dict()["relevance_class"], r.classification)


def _jobspy_row(n: int, **over: Any) -> Dict[str, Any]:
    row = {"job_url": f"https://www.indeed.com/viewjob?jk=abc{n:04d}", "title": f"SAP FICO Consultant {n}",
           "company": "Acme Manufacturing", "location": "Austin, TX, US", "date_posted": date(2026, 9, 30),
           "description": "S/4HANA FICO implementation with ABAP and Power BI reporting.", "is_remote": True,
           "job_level": None, "min_amount": 120000.0, "max_amount": 150000.0, "interval": "yearly",
           "currency": "USD"}
    row.update(over)
    return row


class JobSpyTests(unittest.TestCase):
    def setUp(self) -> None:
        scratch = tempfile.TemporaryDirectory()
        self.addCleanup(scratch.cleanup)
        self.store = MemoryStore()
        user = str(uuid.uuid4())
        ws = self.store.create_workspace(user, "W", "w-jobspy")
        self.ctx = Ctx(ws["id"], user, "owner")
        self.calls: List[Dict[str, Any]] = []
        self.rows: Dict[str, List[Dict[str, Any]]] = {}

        def scrape(**kwargs):
            self.calls.append(kwargs)
            return self.rows.get(kwargs["search_term"], [])

        self.extra: Dict[str, Any] = {"jobspy_scrape": scrape, "jobspy_boards_enabled": [],
                                      "job_monitor_retry_delays": [0], "job_monitor_end_delay": 0,
                                      "job_monitor_sleep": lambda _s: None}
        self.platform = Platform(self.store, storage=LocalFileStorage(Path(scratch.name)),
                                 config=PlatformConfig(extra=self.extra))
        self.svc = self.platform.service("job_monitors")
        self.svc.upload_keyword_set(self.ctx, "IT_Crawler_Keywords.xlsx", workbook_bytes())

    def monitor(self, **filters: Any) -> Dict[str, Any]:
        return self.svc.create_monitor(self.ctx, {"strategy": "jobspy", "filters": {
            "boards": ["indeed"], "search_terms": ["SAP", "ERP"], "location": "United States", "hours_old": 24,
            "results_wanted": 3, **filters}})

    def run_monitor(self, monitor: Dict[str, Any], mode: str = "incremental") -> Dict[str, Any]:
        run = self.svc.start_run(self.ctx, monitor["id"], mode=mode)
        run_task_inline(self.platform, self.ctx.workspace_id, run["task_id"])
        return self.store.get(self.ctx, "job_monitor_runs", run["id"])

    def test_filters_are_validated(self) -> None:
        params = validate_jobspy_filters({"search_terms": "SAP, Oracle\nERP"})
        self.assertEqual(params["search_terms"], ["SAP", "Oracle", "ERP"])
        self.assertEqual((params["boards"], params["hours_old"], params["results_wanted"]), (["indeed"], 24, 25))
        for bad in ({"search_terms": []}, {"search_terms": ["x"], "boards": ["monster"]},
                    {"search_terms": ["x"], "hours_old": 0}, {"search_terms": ["x"], "results_wanted": 5000}):
            with self.assertRaises(ValidationError):
                validate_jobspy_filters(bad)

    def test_monitor_shape(self) -> None:
        m = self.monitor()
        self.assertEqual((m["strategy"], m["source_name"], m["source_url"]), ("jobspy", "JobSpy",
                                                                             "https://www.indeed.com/"))
        self.assertEqual(m["name"], "JobSpy Indeed - United States")
        self.assertEqual(m["filters"]["search_terms"], ["SAP", "ERP"])

    def test_disabled_board_is_reported_and_never_called(self) -> None:
        run = self.run_monitor(self.monitor())
        self.assertEqual(run["status"], "partial")
        self.assertIn("disabled", run["stop_reason"])
        self.assertEqual(self.calls, [])
        self.assertEqual(self.store.count(self.ctx, "job_postings", {}), 0)

    def test_enabled_board_stores_classified_jobs(self) -> None:
        self.extra["jobspy_boards_enabled"] = ["indeed"]
        self.rows["SAP"] = [_jobspy_row(1), _jobspy_row(2), _jobspy_row(3), _jobspy_row(4)]
        self.rows["ERP"] = [_jobspy_row(1),                       # same job from another search: deduped
                            _jobspy_row(9, title="Delivery Driver", description="ASAP, enterprise routes",
                                        is_remote=False, min_amount=None, max_amount=None)]
        monitor = self.monitor()
        run = self.run_monitor(monitor)
        self.assertEqual((run["status"], run["pages"], run["new_count"]), ("completed", 2, 4))
        self.assertEqual(len(self.calls), 2)
        self.assertEqual(self.calls[0]["site_name"], ["indeed"])
        self.assertEqual((self.calls[0]["hours_old"], self.calls[0]["results_wanted"],
                          self.calls[0]["location"]), (24, 3, "United States"))
        jobs = {j["url_key"]: j for j in self.store.all(self.ctx, "job_postings", {})}
        sap = jobs["https://indeed.com/viewjob?jk=abc0001"]
        self.assertEqual((sap["source"], sap["source_board"], sap["search_term"]), ("JobSpy", "Indeed", "SAP"))
        self.assertEqual((sap["remote"], sap["salary_budget"]), ("Remote", "$120k–$150k/yearly"))
        self.assertEqual(sap["posted_at"].date(), date(2026, 9, 30))
        self.assertEqual(sap["relevance_class"], "HIGH")
        self.assertGreaterEqual(sap["relevance_score"], 70)
        self.assertIn("FICO", sap["matched_keywords"])
        self.assertIn("ERP Modules (SAP Specific)", sap["matched_categories"])
        self.assertTrue(sap["relevance_reason"].startswith("HIGH"))
        self.assertEqual([sap[f"keyword_{i}"] for i in range(1, 4)], ["SAP", "FICO", "ABAP"])
        driver = jobs["https://indeed.com/viewjob?jk=abc0009"]
        self.assertEqual((driver["relevance_class"], driver["matched_keywords"]), ("REJECT", []))
        self.assertIsNone(driver["salary_budget"])
        self.assertIsNone(driver["keyword_1"])
        # filters over relevance
        q = self.svc.search_jobs
        self.assertEqual(q(self.ctx, {"relevance": "HIGH"})["total"], 3)
        self.assertEqual(q(self.ctx, {"relevance": "REJECT", "source_board": "Indeed"})["total"], 1)
        self.assertEqual(q(self.ctx, {"category": "ERP Platforms"})["total"], 3)
        self.assertEqual(q(self.ctx, {"search_term": "SAP"})["total"], 3)
        self.assertEqual(q(self.ctx, {"conditions": {"all": [{"field": "relevance_score", "op": "gte",
                                                              "value": 70}]}})["total"], 3)
        # JobSpy is a search window: a full sweep never closes anything
        self.rows = {"SAP": [], "ERP": []}
        for _ in range(3):
            self.assertEqual(self.run_monitor(monitor, "full")["closed_count"], 0)
        self.assertEqual(self.store.count(self.ctx, "job_postings", {"status": "closed"}), 0)

    def test_record_mapping_never_invents(self) -> None:
        rec = record_from_jobspy({"job_url": "https://www.indeed.com/viewjob?jk=1", "title": "X",
                                  "company": float("nan"), "min_amount": float("nan"), "is_remote": None},
                                 board="indeed", search_term="ERP")
        self.assertIsNone(rec["company_name"])
        self.assertIsNone(rec["salary_budget"])
        self.assertIsNone(rec["remote"])
        self.assertEqual((rec["source_board"], rec["search_term"]), ("Indeed", "ERP"))

    def test_rescore_and_keyword_set_switch(self) -> None:
        self.extra["jobspy_boards_enabled"] = ["indeed"]
        self.rows["SAP"] = [_jobspy_row(1)]
        self.run_monitor(self.monitor(search_terms=["SAP"]))
        from openpyxl import Workbook

        book = Workbook()
        sheet = book.active
        sheet.title = "All Keywords"
        sheet.append(["#", "Keyword", "Category"])
        sheet.append([1, "Kubernetes", "Containerization & Orchestration"])
        buf = io.BytesIO()
        book.save(buf)
        self.svc.upload_keyword_set(self.ctx, "k8s.xlsx", buf.getvalue())
        self.assertEqual(self.svc.rescore(self.ctx)["rescored"], 1)
        job = self.store.all(self.ctx, "job_postings", {})[0]
        self.assertEqual(job["matched_keywords"], [])
        # Keyword 1-5 had been derived from the old matches: they follow the new keyword set
        self.assertEqual([job[f"keyword_{i}"] for i in range(1, 6)], [None] * 5)

    def test_rescore_keeps_source_tags(self) -> None:
        values = {"job_url": "https://example.com/j/1", "title": "SAP FICO Consultant",
                  "keyword_1": "Go", "keyword_2": "Kotlin", "matched_keywords": ["SAP", "FICO"]}
        patch = self.svc._rescored(self.svc.engine(self.ctx), values)
        self.assertNotIn("keyword_1", patch)                        # the source's own tags stay
        self.assertEqual(patch["relevance_class"], "HIGH")
        self.assertEqual(len(self.store.all(self.ctx, "job_keyword_sets", {"active": True})), 1)

    def test_viewer_cannot_upload_keywords(self) -> None:
        from cloud.intel.core.context import ForbiddenError

        viewer = Ctx(self.ctx.workspace_id, str(uuid.uuid4()), "viewer")
        self.store.add_member(self.ctx, viewer.user_id, "viewer")
        with self.assertRaises(ForbiddenError):
            self.svc.upload_keyword_set(viewer, "k.xlsx", workbook_bytes())
        self.assertEqual(self.svc.score_text(viewer, title="SAP FICO Consultant")["relevance_class"], "HIGH")


if __name__ == "__main__":
    unittest.main()
