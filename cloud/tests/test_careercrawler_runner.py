"""The CareerCrawler runner adapter, driving the *real* crawler engine.

The engine, its seed selection, platform detection, dispatch and the crawler's
own XLSX exporter all run for real. Only the adapters are replaced (through the
engine's registry injection point), and DNS is a fake resolver — so nothing here
touches the network, and the assertions are about isolation and progress rather
than about any particular job board.
"""

from __future__ import annotations

import csv
import json
import os
import sys
import tempfile
import threading
import unittest
import uuid
from pathlib import Path
from typing import List

from cloud.shared.models import JobStatus, JobType, ResultKind, TargetStatus
from cloud.shared.repository import InMemoryJobRepository
from cloud.shared.schemas import parse_job_request
from cloud.shared.service import JobService, RetryPolicy
from cloud.shared.storage import LocalFileStorage
from cloud.worker.careercrawler_runner import (
    CareerCrawlerRunner,
    apply_crawler_settings,
    restore_crawler_settings,
)
from cloud.worker.executor import JobExecutor
from cloud.worker.results import ResultWriter
from cloud.worker.workspace import PROTECTED_DIRECTORIES, REPO_ROOT, JobWorkspace, UnsafeRuntimeRootError, check_runtime_root

OWNER = str(uuid.uuid4())


def public_resolver(host, port, type=None):
    return [(2, 1, 6, "", ("93.184.216.34", port))]


def private_resolver(host, port, type=None):
    return [(2, 1, 6, "", ("10.0.0.7", port))]


def snapshot(directory: Path):
    if not directory.exists():
        return None
    return sorted((str(p.relative_to(directory)), p.stat().st_mtime_ns) for p in directory.rglob("*"))


class RunnerHarness(unittest.TestCase):
    def setUp(self) -> None:
        from config.settings import SETTINGS

        self.settings_before = {name: getattr(SETTINGS, name) for name in vars(SETTINGS)}
        self.protected_before = {str(d): snapshot(d) for d in PROTECTED_DIRECTORIES}

        scratch = tempfile.TemporaryDirectory()
        self.addCleanup(scratch.cleanup)
        self.scratch = Path(scratch.name)
        self.runtime = self.scratch / "runtime"
        self.repo = InMemoryJobRepository()
        self.service = JobService(self.repo)
        self.storage = LocalFileStorage(self.scratch / "results")
        self.adapter_calls: List[str] = []
        self.release = threading.Event()
        self.release.set()

    def tearDown(self) -> None:
        from config.settings import SETTINGS

        restore_crawler_settings(self.settings_before)
        self.assertEqual({name: getattr(SETTINGS, name) for name in vars(SETTINGS)}, self.settings_before)
        # Nothing a cloud job did touched CareerCrawler's production directories.
        self.assertEqual({str(d): snapshot(d) for d in PROTECTED_DIRECTORIES}, self.protected_before)

    # --- fakes -----------------------------------------------------------------

    def fake_greenhouse(self, url, company, session=None):
        from models.job import Job

        self.release.wait(10)
        self.adapter_calls.append(url)
        if "broken" in url:
            raise RuntimeError("board exploded")
        return [
            Job(company_name=company, job_title="Data Engineer", job_url=f"{url}/jobs/1", location="Austin, TX", platform="Greenhouse"),
            Job(company_name=company, job_title="=HYPERLINK(\"http://evil\",\"x\")", job_url=f"{url}/jobs/2", location="Remote", platform="Greenhouse"),
            Job(company_name=company, job_title="Data Engineer", job_url=f"{url}/jobs/1", location="Austin, TX", platform="Greenhouse"),
        ]

    def engine_factory(self):
        from crawler.crawler_engine import CrawlerEngine
        from crawler.platform_detector import Platform

        return CrawlerEngine(registry={Platform.GREENHOUSE: self.fake_greenhouse})

    def runner(self, **kwargs) -> CareerCrawlerRunner:
        options = dict(
            engine_factory=self.engine_factory,
            session_factory=lambda: object(),
            resolver=public_resolver,
            runtime_root=self.runtime,
            company_concurrency=2,
        )
        options.update(kwargs)
        return CareerCrawlerRunner(**options)

    def execute(self, payload, runner=None):
        job = self.service.create_job(parse_job_request(payload), owner_id=OWNER, max_attempts=1)
        executor = JobExecutor(
            self.service,
            runner or self.runner(),
            result_writer=ResultWriter(self.storage),
            runtime_root=self.runtime,
            retry_policy=RetryPolicy(max_attempts=1),
            cancel_check_interval=0,
        )
        return executor.execute(job.job_id)

    def result_path(self, job, kind: ResultKind) -> Path:
        result = next(r for r in self.service.list_results(job.job_id, owner_id=OWNER) if r.kind is kind)
        return self.storage.root / result.storage_key


class TestSingleCompany(RunnerHarness):
    def test_a_company_is_crawled_through_the_real_engine(self) -> None:
        job = self.execute({"type": "single_company", "website": "boards.greenhouse.io/acme", "company_name": "Acme"})

        self.assertIs(job.status, JobStatus.COMPLETED, job.error)
        self.assertEqual(self.adapter_calls, ["https://boards.greenhouse.io/acme"])
        progress = job.progress
        self.assertEqual((progress.total, progress.completed, progress.failed, progress.jobs_found), (1, 1, 0, 3))
        self.assertEqual(progress.current_phase, "completed")
        self.assertIsNone(progress.current_company)

        (target,) = self.service.list_targets(job.job_id, owner_id=OWNER)
        self.assertEqual((target.status, target.platform, target.outcome, target.jobs_found), (TargetStatus.COMPLETED, "Greenhouse", "jobs", 3))
        self.assertIsNotNone(target.started_at)
        self.assertIsNotNone(target.completed_at)

        kinds = {r.kind for r in self.service.list_results(job.job_id, owner_id=OWNER)}
        self.assertTrue({ResultKind.SUMMARY_JSON, ResultKind.JOBS_CSV, ResultKind.JOBS_XLSX} <= kinds)

        summary = json.loads(self.result_path(job, ResultKind.SUMMARY_JSON).read_text(encoding="utf-8"))
        self.assertEqual(summary["summary"]["jobs_found"], 2, "duplicates are removed exactly as the engine's crawl() does")
        self.assertEqual(summary["companies"][0]["platform"], "Greenhouse")
        self.assertEqual({p["job_url"] for p in summary["postings"]}, {"https://boards.greenhouse.io/acme/jobs/1", "https://boards.greenhouse.io/acme/jobs/2"})

    def test_spreadsheet_formulas_from_scraped_titles_are_neutralised(self) -> None:
        from openpyxl import load_workbook

        job = self.execute({"type": "single_company", "website": "boards.greenhouse.io/acme"})
        with self.result_path(job, ResultKind.JOBS_CSV).open(encoding="utf-8-sig", newline="") as handle:
            titles = [row["job_title"] for row in csv.DictReader(handle)]
        self.assertIn("'=HYPERLINK(\"http://evil\",\"x\")", titles)

        workbook = load_workbook(self.result_path(job, ResultKind.JOBS_XLSX))
        cells = [cell for row in workbook.active.iter_rows(min_row=2) for cell in row]
        self.assertTrue(cells)
        self.assertFalse([c.coordinate for c in cells if c.data_type == "f"], "no cell may be a formula")

    def test_the_workspace_is_private_and_removed_afterwards(self) -> None:
        job = self.execute({"type": "single_company", "website": "boards.greenhouse.io/acme"})
        self.assertIs(job.status, JobStatus.COMPLETED)
        self.assertFalse((self.runtime / job.job_id).exists())

    def test_crawler_settings_are_pointed_away_from_production_paths(self) -> None:
        from config.settings import SETTINGS

        apply_crawler_settings(self.runtime, browser_fallback=False)
        self.assertFalse(SETTINGS.diagnostics)
        for path in (SETTINGS.output_dir, SETTINGS.diagnostics_dir):
            self.assertIn(self.runtime.resolve(), Path(path).resolve().parents)
        self.assertEqual(SETTINGS.max_workers, self.settings_before["max_workers"], "the production worker default is not touched")


class TestBulkAndFailures(RunnerHarness):
    def test_bulk_progress_counts_failures_without_failing_the_job(self) -> None:
        job = self.execute(
            {
                "type": "bulk_companies",
                "companies": [
                    {"website": "boards.greenhouse.io/one"},
                    {"website": "boards.greenhouse.io/broken"},
                    {"company_name": "Name Only Inc"},
                    {"website": "acme-marketing-site.com"},
                ],
            }
        )
        self.assertIs(job.status, JobStatus.COMPLETED, job.error)
        progress = job.progress
        self.assertEqual((progress.total, progress.completed), (4, 4))
        self.assertEqual(progress.jobs_found, 3)
        self.assertEqual(progress.failed, 3)

        targets = {t.position: t for t in self.service.list_targets(job.job_id, owner_id=OWNER)}
        self.assertEqual(targets[0].status, TargetStatus.COMPLETED)
        self.assertEqual(targets[1].status, TargetStatus.FAILED)
        self.assertIn("board exploded", targets[1].error)
        self.assertEqual(targets[2].status, TargetStatus.SKIPPED)
        self.assertIn("name-only", targets[2].error)
        self.assertEqual((targets[3].status, targets[3].platform), (TargetStatus.FAILED, "Generic HTML"))

    def test_websites_resolving_to_private_addresses_are_never_crawled(self) -> None:
        job = self.execute(
            {"type": "single_company", "website": "boards.greenhouse.io/acme"},
            runner=self.runner(resolver=private_resolver),
        )
        self.assertIs(job.status, JobStatus.COMPLETED)
        self.assertEqual(self.adapter_calls, [])
        (target,) = self.service.list_targets(job.job_id, owner_id=OWNER)
        self.assertEqual(target.status, TargetStatus.FAILED)
        self.assertIn("non-public", target.error)

    def test_unsupported_job_types_are_refused_without_crawling(self) -> None:
        runner = self.runner()
        for payload in ({"type": "weekly_crawl"}, {"type": "discovery", "website": "acme.com"}):
            with self.subTest(payload["type"]):
                self.assertFalse(runner.supports(JobType(payload["type"])))
                job = self.execute(payload, runner=runner)
                self.assertIs(job.status, JobStatus.QUEUED, "left for a runner that supports it")
        self.assertEqual(self.adapter_calls, [])

    def test_the_time_limit_skips_companies_not_yet_started(self) -> None:
        ticks = iter([0.0, 0.0] + [10_000.0] * 50)
        runner = self.runner(company_concurrency=1, max_runtime_seconds=60, clock=lambda: next(ticks))
        job = self.execute(
            {"type": "bulk_companies", "companies": [{"website": "boards.greenhouse.io/a"}, {"website": "boards.greenhouse.io/b"}]},
            runner=runner,
        )
        self.assertIs(job.status, JobStatus.COMPLETED)
        # The clock passes the deadline after the first company has started: it
        # finishes, and the second is recorded as skipped rather than dropped.
        targets = self.service.list_targets(job.job_id, owner_id=OWNER)
        self.assertEqual([t.status for t in targets], [TargetStatus.COMPLETED, TargetStatus.SKIPPED])
        self.assertIn("time limit", targets[1].error)
        summary = json.loads(self.result_path(job, ResultKind.SUMMARY_JSON).read_text(encoding="utf-8"))
        self.assertTrue(summary["summary"]["timed_out"])


class TestCancellation(RunnerHarness):
    def test_cancelling_a_bulk_job_stops_between_companies(self) -> None:
        self.release.clear()
        companies = [{"website": f"boards.greenhouse.io/c{n}"} for n in range(20)]
        job = self.service.create_job(parse_job_request({"type": "bulk_companies", "companies": companies}), owner_id=OWNER)
        executor = JobExecutor(
            self.service, self.runner(company_concurrency=1), result_writer=ResultWriter(self.storage),
            runtime_root=self.runtime, cancel_check_interval=0,
        )
        outcome = {}
        thread = threading.Thread(target=lambda: outcome.setdefault("job", executor.execute(job.job_id)))
        thread.start()
        for _ in range(500):
            if self.service.get_job(job.job_id).status is JobStatus.RUNNING:
                break
            threading.Event().wait(0.01)
        self.service.request_cancel(job.job_id, owner_id=OWNER)
        self.release.set()
        thread.join(30)

        final = outcome["job"]
        self.assertIs(final.status, JobStatus.CANCELLED)
        self.assertLess(len(self.adapter_calls), 20)
        self.assertEqual(self.service.list_results(job.job_id), [])


class TestIsolationGuards(unittest.TestCase):
    def test_runtime_root_may_not_overlap_crawler_production_directories(self) -> None:
        for bad in (REPO_ROOT / "state", REPO_ROOT / "output", REPO_ROOT / "output" / "cloud", REPO_ROOT / "secrets", REPO_ROOT, REPO_ROOT.parent):
            with self.subTest(bad=str(bad)), self.assertRaises(UnsafeRuntimeRootError):
                check_runtime_root(bad)
        with self.assertRaises(UnsafeRuntimeRootError):
            CareerCrawlerRunner(runtime_root=REPO_ROOT / "state")

    def test_workspace_ids_cannot_escape_the_root(self) -> None:
        with tempfile.TemporaryDirectory() as root:
            for bad in ("../../etc", "job_../../x", "job_ABC", "", "job_" + "0" * 31):
                with self.subTest(bad=bad), self.assertRaises(ValueError):
                    JobWorkspace.create(Path(root), bad, 1)
            workspace = JobWorkspace.create(Path(root), "job_" + "a" * 32, 2)
            self.assertTrue(workspace.state.is_dir() and workspace.output.is_dir() and workspace.logs.is_dir())
            workspace.remove()
            self.assertEqual(list(Path(root).iterdir()), [])

    def test_a_new_attempt_clears_a_dead_attempts_workspace(self) -> None:
        with tempfile.TemporaryDirectory() as root:
            job_id = "job_" + "b" * 32
            dead = JobWorkspace.create(Path(root), job_id, 1)
            (dead.logs / "crawl.log").write_text("half a crawl", encoding="utf-8")
            JobWorkspace.create(Path(root), job_id, 2)
            self.assertEqual(sorted(p.name for p in (Path(root) / job_id).iterdir()), ["attempt-2"])

    def test_abandoned_workspaces_are_swept_and_live_ones_kept(self) -> None:
        import time as _time

        from cloud.worker.workspace import sweep_abandoned_workspaces

        with tempfile.TemporaryDirectory() as root:
            old = JobWorkspace.create(Path(root), "job_" + "c" * 32, 1)
            live = JobWorkspace.create(Path(root), "job_" + "d" * 32, 1)
            (Path(root) / "not-a-job").mkdir()
            past = _time.time() - 10 * 3600
            for item in [*old.path.parent.rglob("*"), old.path.parent]:
                os.utime(item, (past, past))
            removed = sweep_abandoned_workspaces(Path(root), older_than_seconds=3600)
            self.assertEqual(removed, [old.path.parent.name])
            self.assertTrue(live.path.exists())
            self.assertTrue((Path(root) / "not-a-job").exists())

    def test_the_worker_process_never_opens_the_crawler_database(self) -> None:
        """Run a job in a separate process and prove no SQLite file was opened."""
        import subprocess

        script = (
            "import sys, builtins, tempfile, pathlib, uuid\n"
            "opened = []\n"
            "real_open = builtins.open\n"
            "def spy(file, *a, **k):\n"
            "    opened.append(str(file)); return real_open(file, *a, **k)\n"
            "builtins.open = spy\n"
            "from cloud.tests.test_careercrawler_runner import RunnerHarness\n"
            "class T(RunnerHarness):\n"
            "    def runTest(self):\n"
            "        job = self.execute({'type': 'single_company', 'website': 'boards.greenhouse.io/acme'})\n"
            "        assert job.status.value == 'completed', job\n"
            "import unittest\n"
            "result = unittest.TextTestRunner(stream=open(__import__('os').devnull, 'w')).run(T())\n"
            "assert result.wasSuccessful(), result.errors + result.failures\n"
            "assert 'sqlite3' not in sys.modules, 'sqlite3 was imported'\n"
            "bad = [p for p in opened if p.endswith(('.db', '.db-wal', 'checkpoint.json')) or 'secrets' in p]\n"
            "assert not bad, bad\n"
            "print('ok')\n"
        )
        result = subprocess.run(
            [sys.executable, "-c", script], cwd=REPO_ROOT, capture_output=True, text=True, timeout=180,
            env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"},
        )
        self.assertEqual(result.returncode, 0, result.stderr[-3000:])
        self.assertIn("ok", result.stdout)


if __name__ == "__main__":
    unittest.main()
