"""CSV, XLSX, JSON and NDJSON output, and the Companies / Jobs / All-fields views.

Files written per run:

=====================  ===========================================================
``results.csv``        All fields (one row per record)
``jobs.csv``           Jobs view (job requests)
``companies.csv``      Companies view (one row per company, with ``job_count``)
``pages.csv``          Every page visited, with its outcome and browser use
``errors.csv``         Refused/failed pages, input problems, rejected values
``results.xlsx``       Sheets: All Fields, Companies, Jobs, Pages, Errors, Run Summary, Inputs
``results.json``       Everything, with complete provenance (evidence, field status,
                       conflicts, rejected values, pages, errors, summary, schema)
``results.ndjson``     One record per line, with its provenance
=====================  ===========================================================

Every row carries ``source_url`` (plus ``source_urls`` when duplicates were
merged), ``extraction_method``, ``confidence`` and ``extracted_at``. Cells that a
spreadsheet would run as a formula (``=``, ``+``, ``-``, ``@``) are neutralised.
"""

from __future__ import annotations

import csv
import json
import tempfile
from pathlib import Path
from typing import Any, Dict, List, Mapping, Optional, Sequence

from cloud.intel.core.normalize import company_name_key, domain_of

__all__ = ["FILE_NAMES", "META_COLUMNS", "PAGE_COLUMNS", "ERROR_COLUMNS", "INPUT_COLUMNS", "columns_for", "views",
           "write_outputs"]

META_COLUMNS = ["source_url", "extraction_method", "confidence", "extracted_at"]
PAGE_COLUMNS = ["input_row", "input_url", "url", "final_url", "kind", "outcome", "page_no", "http_status", "attempts",
                "records", "browser_used", "browser_reason", "browser_duration_ms", "browser_outcome", "error",
                "fetched_at"]
ERROR_COLUMNS = ["kind", "input_row", "url", "outcome", "field", "value", "error"]
INPUT_COLUMNS = ["row", "batch", "source", "url", "final_url", "outcome", "reason", "records", "pages_fetched",
                 "ai_used", "extraction_method"]
#: file key (as used by ``/files/{fmt}?view=``) -> published file name
FILE_NAMES = {"csv": "results.csv", "jobs.csv": "jobs.csv", "companies.csv": "companies.csv",
              "pages.csv": "pages.csv", "errors.csv": "errors.csv", "xlsx": "results.xlsx", "json": "results.json",
              "ndjson": "results.ndjson"}
_TYPES = {".csv": "text/csv", ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
          ".json": "application/json", ".ndjson": "application/x-ndjson"}


def columns_for(schema: Mapping[str, Any], view: str = "all") -> List[str]:
    fields = list(schema["fields"])
    if view == "jobs":
        names = [f["name"] for f in fields if f.get("level") == "job"]
        names = [n for n in ("company_name", "website") if any(f["name"] == n for f in fields)] + names
    elif view == "companies":
        names = [f["name"] for f in fields if f.get("level") != "job"]
    else:
        names = [f["name"] for f in fields]
    return names + META_COLUMNS


def views(records: Sequence[Mapping[str, Any]], schema: Mapping[str, Any]) -> Dict[str, List[Dict[str, Any]]]:
    """``{"all": [...], "companies": [...], "jobs": [...]}`` (``jobs`` is empty for company requests)."""
    all_rows = [dict(r) for r in records]
    if schema.get("entity") != "job":
        return {"all": all_rows, "companies": all_rows, "jobs": []}
    company_cols = [c for c in columns_for(schema, "companies") if c not in META_COLUMNS]
    companies: Dict[str, Dict[str, Any]] = {}
    for row in records:
        key = (row.get("domain") or (domain_of(row["website"]) if row.get("website") else None)
               or company_name_key(row.get("company_name")) or row.get("input_url") or row.get("source_url") or "")
        if key not in companies:
            companies[key] = {**{c: row.get(c) for c in company_cols}, "source_url": row.get("source_url"),
                              "extraction_method": row.get("extraction_method"), "confidence": row.get("confidence"),
                              "extracted_at": row.get("extracted_at"), "job_count": 0,
                              "_evidence": {c: e for c, e in (row.get("_evidence") or {}).items() if c in company_cols}}
        else:
            for c in company_cols:
                if companies[key].get(c) in (None, "", []) and row.get(c) not in (None, "", []):
                    companies[key][c] = row[c]
        if row.get("job_title") or row.get("job_url"):
            companies[key]["job_count"] += 1
    return {"all": all_rows, "companies": list(companies.values()), "jobs": [r for r in all_rows
                                                                             if r.get("job_title") or r.get("job_url")]}


def _cell(value: Any) -> Any:
    from cloud.worker.results import neutralise_cell

    if isinstance(value, list):
        value = " | ".join(str(v) for v in value)
    elif isinstance(value, dict):
        value = json.dumps(value, default=str)
    return neutralise_cell(value)


def _xlsx_cell(value: Any) -> Any:
    if value is None or isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value
    return _cell(value if isinstance(value, (list, dict)) else str(value))


def _write_csv(path: Path, columns: Sequence[str], rows: Sequence[Mapping[str, Any]]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(columns)
        for row in rows:
            writer.writerow([_cell(row.get(c)) for c in columns])


def write_outputs(storage: Any, base_key: str, run_id: str, schema: Mapping[str, Any],
                  records: Sequence[Mapping[str, Any]], inputs: Sequence[Mapping[str, Any]],
                  summary: Mapping[str, Any], *, pages: Optional[Sequence[Mapping[str, Any]]] = None,
                  errors: Optional[Sequence[Mapping[str, Any]]] = None) -> Dict[str, Dict[str, Any]]:
    """Write every output file to ``storage`` under ``base_key``; returns ``{file key: file info}``."""
    from openpyxl import Workbook

    pages = list(pages or [])
    errors = list(errors or [])
    tables = views(records, schema)
    all_cols = columns_for(schema, "all")
    job_cols = columns_for(schema, "jobs")
    company_cols = columns_for(schema, "companies") + (["job_count"] if schema.get("entity") == "job" else [])
    with tempfile.TemporaryDirectory() as scratch:
        root = Path(scratch)
        _write_csv(root / "results.csv", all_cols + ["source_urls"], tables["all"])
        _write_csv(root / "jobs.csv", job_cols + ["source_urls"], tables["jobs"])
        _write_csv(root / "companies.csv", company_cols, tables["companies"])
        _write_csv(root / "pages.csv", PAGE_COLUMNS, pages)
        _write_csv(root / "errors.csv", ERROR_COLUMNS, errors)

        book = Workbook()
        book.remove(book.active)
        for title, columns, rows in (("All Fields", all_cols + ["source_urls"], tables["all"]),
                                     ("Companies", company_cols, tables["companies"]),
                                     ("Jobs", job_cols + ["source_urls"], tables["jobs"]),
                                     ("Pages", PAGE_COLUMNS, pages),
                                     ("Errors", ERROR_COLUMNS, errors)):
            sheet = book.create_sheet(title)
            sheet.append(columns)
            for row in rows:
                sheet.append([_xlsx_cell(row.get(c)) for c in columns])
        sheet = book.create_sheet("Run Summary")
        sheet.append(["metric", "value"])
        for key, value in summary.items():
            sheet.append([key, _xlsx_cell(value)])
        sheet = book.create_sheet("Inputs")
        sheet.append(INPUT_COLUMNS)
        for row in inputs:
            sheet.append([_xlsx_cell(row.get(c)) for c in INPUT_COLUMNS])
        book.save(root / "results.xlsx")

        (root / "results.json").write_text(json.dumps({
            "run_id": run_id, "instruction": schema.get("instruction"), "schema": schema, "summary": dict(summary),
            "columns": all_cols, "records": list(records), "companies": tables["companies"], "jobs": tables["jobs"],
            "inputs": list(inputs), "pages": pages, "errors": errors}, default=str, indent=1), encoding="utf-8")
        with (root / "results.ndjson").open("w", encoding="utf-8") as handle:
            for record in records:
                handle.write(json.dumps({"run_id": run_id, **record}, default=str) + "\n")

        files: Dict[str, Dict[str, Any]] = {}
        for key, name in FILE_NAMES.items():
            path = root / name
            content_type = _TYPES[path.suffix]
            stored = storage.put_file(f"{base_key}/{name}", path, content_type=content_type)
            stem = name.rsplit(".", 1)[0]
            filename = f"scrape-{run_id}.{path.suffix[1:]}" if stem == "results" else f"scrape-{run_id}-{name}"
            files[key] = {"storage_key": f"{base_key}/{name}", "content_type": content_type, "filename": filename,
                          "size_bytes": getattr(stored, "size_bytes", None), "sha256": getattr(stored, "sha256", None)}
    return files
