"""The SANA GTM AI Scraper: URLs + a plain-language instruction in, validated, de-duplicated,
provenance-rich records out — as a persistent background job.

    service = platform.service("scraper")
    service.plan(ctx, urls, instruction, options=...)       # schema, sources, limits, work/cost estimate
    run = service.start(ctx, urls, instruction, schema=edited, options=..., confirm=True)
    service.pause / resume / cancel / retry / restart (ctx, run_id)
    service.records(ctx, run_id, view="jobs") · pages · errors · evidence
    service.templates.list / create / update / duplicate / delete
    service.match_crm / propose / review / apply_proposals        # PROPOSE → REVIEW → APPLY

Module map (``cloud/intel/scraper``)::

    models.py      data types, outcomes, crawl options, evidence precedence and conflicts
    schemas.py     field vocabulary and types; JSON schemas sent to AI
    planner.py     instruction -> schema (entities, types, filters, criteria, custom fields)
    limits.py      per-domain limiter, retry/backoff policy, run budget
    fetcher.py     SafeFetcher + outcome classification + retries + browser fallback
    pagination.py  next / numbered / cursor / load-more detection
    discovery.py   homepage -> careers page / ATS board candidates
    extractor.py   JSON-LD, official ATS APIs, links/headings, regex; AI with evidence checks
    detail.py      job detail pages and listing+detail merging
    crawler.py     one input URL, deep but bounded
    normalizer.py  URLs, websites, domains, phones, dates, numbers, derived fields
    validator.py   input URLs, field validation and statuses, filters
    dedupe.py      job URL / company / fingerprint keys; sources and evidence kept
    exports.py     CSV / XLSX / JSON / NDJSON and the Companies / Jobs / All views
    runner.py      the background task: concurrency, progress, pause/cancel, recovery
    templates.py   built-in and saved templates
    crm.py         CRM matching and proposals (never applied without approval)
    toolkit.py     the scraper as tools (research agent, AI Control Room)
    service.py     this file

The scraper never logs in, solves a CAPTCHA, gets past a WAF or rotates
proxies; such pages are reported with their outcome. AI goes only through the
provider layer (``platform.service("ai")``), so free-only mode, the $0 budget
and quota exhaustion (rules fallback) apply unchanged.
"""

from __future__ import annotations

import csv
import io
import json
import re
import uuid
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple, Union

from cloud.intel.core.audit import audit
from cloud.intel.core.context import Ctx, ValidationError
from cloud.intel.core.http import UnsafeTargetError, check_url
from cloud.intel.scraper import crm as crm_actions
from cloud.intel.scraper.exports import FILE_NAMES, columns_for, views
from cloud.intel.scraper.models import TRANSIENT, CrawlOptions, ScrapeInput
from cloud.intel.scraper.planner import instruction_to_schema
from cloud.intel.scraper.runner import ACTIVE_STATES, run_scrape_task
from cloud.intel.scraper.schemas import NORMALIZE_RULES, canonical_type
from cloud.intel.scraper.templates import TemplateStore
from cloud.intel.scraper.validator import check_input_url
from cloud.intel.vendor import ats_detect

__all__ = ["HIGH_VOLUME_REQUESTS", "HIGH_VOLUME_URLS", "MAX_URLS", "ScraperService", "clean_schema", "parse_inputs",
           "read_file_rows", "run_scrape_task"]

MAX_URLS = 5000
#: A run above either threshold needs ``confirm=True`` (the UI asks first).
HIGH_VOLUME_URLS = 50
HIGH_VOLUME_REQUESTS = 1500
_URL_HEADER = re.compile(r"url|website|link|site|domain|homepage|careers", re.I)
_LOOKS_LIKE_URL = re.compile(r"^(?:https?://)?[\w-]+(?:\.[\w-]+)+(?:[/:?#]|$)", re.I)
_CAREERS_PATH = re.compile(r"/(?:careers?|jobs?|join(?:-us)?|work-with-us|opportunities|open-positions|vacancies)"
                           r"(?:[/?#.]|$)", re.I)


def read_file_rows(data: bytes, filename: str, column: Optional[str] = None) -> List[Tuple[int, str]]:
    """``(spreadsheet row, cell)`` pairs from the URL column of a CSV, TXT, TSV or XLSX upload.

    The column is the one named, else the first whose header mentions url/website/
    link/domain, else the column whose cells most often look like URLs. A file
    without a header row (a bare list of URLs) is read from row 1.
    """
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
            table = [["" if v is None else str(v).strip() for v in row] for row in sheet.iter_rows(values_only=True)]
        except Exception as error:  # noqa: BLE001 - a corrupt upload is the user's input error
            raise ValidationError(f"could not read the XLSX file: {type(error).__name__}") from None
    elif name.endswith((".csv", ".txt", ".tsv")):
        text = data.decode("utf-8-sig", errors="replace")
        dialect = csv.excel_tab if name.endswith(".tsv") else csv.excel
        table = [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]
    else:
        raise ValidationError("upload a .csv, .txt or .xlsx file")
    if not table:
        return []
    width = max((len(r) for r in table), default=0)
    if width == 0:
        return []
    table = [r + [""] * (width - len(r)) for r in table]
    header = table[0]
    has_header = not any(_LOOKS_LIKE_URL.match(v or "") for v in header)
    index: Optional[int] = None
    if column:
        if not has_header or column not in header:
            raise ValidationError(f"column {column!r} is not in the file (columns: {', '.join(h for h in header if h)})")
        index = header.index(column)
    elif has_header:
        index = next((i for i, h in enumerate(header) if _URL_HEADER.search(h or "")), None)
    if index is None:
        body = table[1:] if has_header else table
        scores = [sum(1 for r in body if _LOOKS_LIKE_URL.match(r[i] or "")) for i in range(width)]
        if max(scores) == 0:
            raise ValidationError("no URL column found; name the column that holds the URLs")
        index = scores.index(max(scores))
    start = 2 if has_header else 1
    return [(start + offset, row[index]) for offset, row in enumerate(table[start - 1:])]


def parse_inputs(urls: Union[None, str, Sequence[Any]] = None, *, file_bytes: Optional[bytes] = None,
                 filename: Optional[str] = None, column: Optional[str] = None
                 ) -> Tuple[List[ScrapeInput], List[Dict[str, Any]], Dict[str, int]]:
    """Validate pasted URLs and/or an uploaded file.

    Returns ``(accepted, rejected, report)``. Every accepted input keeps its row
    number and import batch; rejected ones carry a reason (invalid syntax,
    unsupported protocol, private address, duplicate). Empty rows are counted.
    """
    candidates: List[Tuple[int, str, str, str]] = []   # (row, raw, batch, source)
    if urls:
        batch = "b_" + uuid.uuid4().hex[:12]
        if isinstance(urls, str):
            lines = urls.splitlines()
            if len(lines) == 1 and re.search(r"[,\s]", lines[0].strip()):   # "a.com, b.com" on one line
                lines = re.split(r"[,\s]+", lines[0].strip())
            source = "paste"
        else:
            lines, source = ["" if u is None else str(u) for u in urls], "list"
        candidates += [(i + 1, line, batch, source) for i, line in enumerate(lines)]
    if file_bytes is not None:
        batch = "b_" + uuid.uuid4().hex[:12]
        candidates += [(row, cell, batch, (filename or "upload")[:120])
                       for row, cell in read_file_rows(file_bytes, filename or "", column)]
    accepted: List[ScrapeInput] = []
    rejected: List[Dict[str, Any]] = []
    report = {"received": 0, "accepted": 0, "empty": 0, "invalid": 0, "unsafe": 0, "duplicates": 0}
    seen = set()
    for row, raw, batch, source in candidates:
        if not str(raw or "").strip():
            report["empty"] += 1
            continue
        report["received"] += 1
        url, reason = check_input_url(raw)
        if url is None:
            report["invalid"] += 1
            rejected.append({"row": row, "batch": batch, "source": source, "url": str(raw)[:300], "reason": reason})
            continue
        try:
            check_url(url, resolve=False)   # DNS is re-checked on every hop at fetch time
        except UnsafeTargetError as error:
            report["unsafe"] += 1
            rejected.append({"row": row, "batch": batch, "source": source, "url": url[:300], "reason": str(error)})
            continue
        key = url.rstrip("/").lower()
        if key in seen:
            report["duplicates"] += 1
            rejected.append({"row": row, "batch": batch, "source": source, "url": url[:300], "reason": "duplicate"})
            continue
        seen.add(key)
        accepted.append(ScrapeInput(url[:2048], row, batch, source))
    report["accepted"] = len(accepted)
    return accepted, rejected, report


def clean_schema(schema: Mapping[str, Any]) -> Dict[str, Any]:
    """A schema the client supplied or edited (renamed, retyped, reordered, fields added or
    removed): keep only well-formed fields, in the order given."""
    fields = []
    seen = set()
    for item in schema.get("fields") or []:
        if not isinstance(item, Mapping):
            continue
        name = re.sub(r"[^a-z0-9]+", "_", str(item.get("name", "")).lower()).strip("_")[:60]
        if not name or name in seen:
            continue
        seen.add(name)
        kind = canonical_type(item.get("type"))
        enum = [str(v)[:100] for v in (item.get("enum") or []) if str(v).strip()][:100] if kind == "enum" else None
        if kind == "enum" and not enum:
            raise ValidationError(f"field {name!r}: an enum needs its allowed values")
        pattern = str(item["pattern"])[:300] if item.get("pattern") else None
        if pattern:
            try:
                re.compile(pattern)
            except re.error:
                raise ValidationError(f"field {name!r}: invalid pattern") from None
        max_length = item.get("max_length")
        fields.append({"name": name, "label": str(item.get("label") or name.replace("_", " ").capitalize())[:120],
                       "type": kind, "level": "job" if item.get("level") == "job" else "company",
                       "description": str(item.get("description") or name)[:300],
                       "required": bool(item.get("required")), "source": str(item.get("source") or "user")[:20],
                       "normalize": item.get("normalize") if item.get("normalize") in NORMALIZE_RULES else None,
                       "enum": enum, "pattern": pattern,
                       "max_length": int(max_length) if isinstance(max_length, (int, float)) and max_length > 0 else None,
                       "hint": str(item["hint"])[:300] if item.get("hint") else None})
    if not fields:
        raise ValidationError("the schema has no fields")
    entity = "job" if schema.get("entity") == "job" or any(f["level"] == "job" for f in fields) else "company"
    names = {f["name"] for f in fields}
    filters = []
    for f in schema.get("filters") or []:
        if not isinstance(f, Mapping) or f.get("field") not in names:
            continue
        if f.get("op") == "within_days" and isinstance(f.get("value"), int):
            filters.append({"field": f["field"], "op": "within_days", "value": f["value"],
                            "mode": "soft" if f.get("mode") == "soft" else "hard"})
        elif f.get("op") == "contains_any" and isinstance(f.get("value"), list):
            filters.append({"field": f["field"], "op": "contains_any", "value": [str(v)[:100] for v in f["value"]][:20],
                            "mode": "soft" if f.get("mode") == "soft" else "hard"})
    return {"version": 3, "entity": entity,
            "entities": ["company", "job"] if entity == "job" and any(f["level"] == "company" for f in fields)
            else [entity],
            "fields": fields, "filters": filters, "criteria": dict(schema.get("criteria") or {}),
            "custom": list(schema.get("custom") or [])[:20], "parser": str(schema.get("parser") or "user")[:40],
            "instruction": str(schema.get("instruction") or "")[:4000]}


def detect_source(url: str) -> Dict[str, Any]:
    """What an input URL looks like, from the URL alone (nothing is fetched)."""
    found = ats_detect.detect(url)
    if found:
        return {"kind": "ats_board", "label": f"{found.get('platform')} job board", "ats": found.get("platform"),
                "official_api": bool(ats_detect.api_endpoints(found))}
    path = url.split("://", 1)[-1].partition("/")[2]
    if _CAREERS_PATH.search("/" + path):
        return {"kind": "careers_page", "label": "Careers page"}
    if not path.strip("/"):
        return {"kind": "homepage", "label": "Homepage (careers page discovered automatically)"}
    return {"kind": "page", "label": "Web page"}


class ScraperService:
    def __init__(self, platform: Any) -> None:
        self.platform = platform
        self.templates = TemplateStore(platform)

    def _ai(self, ctx: Ctx):
        return self.platform.service("ai").for_ctx(ctx, "extraction")

    # --- schema & planning -------------------------------------------------------------------

    def instruction_to_schema(self, ctx: Ctx, instruction: str) -> Dict[str, Any]:
        """The extraction schema for ``instruction`` (AI is asked only about unrecognised phrases)."""
        if not (instruction or "").strip():
            raise ValidationError("describe what to extract")
        ai = self._ai(ctx)
        return instruction_to_schema(instruction, ai=ai if ai.external else None)

    preview = instruction_to_schema
    preview_schema = instruction_to_schema

    def _template_defaults(self, ctx: Ctx, template_id: Optional[str], instruction: str,
                           schema: Optional[Mapping[str, Any]], options: Optional[Mapping[str, Any]]):
        if not template_id:
            return instruction, schema, options
        template = self.templates.get(ctx, template_id)
        return (instruction or template["instruction"], schema or (template.get("schema") or None),
                {**(template.get("options") or {}), **dict(options or {})})

    def estimate(self, ctx: Ctx, inputs: int, options: CrawlOptions, schema: Mapping[str, Any],
                 sources: Sequence[Mapping[str, Any]] = ()) -> Dict[str, Any]:
        """Upper-bound work for a run, and its AI cost (always $0 in free-only mode)."""
        pages_per_input = options.max_pages if options.pagination or options.follow_details else 4
        pages_max = inputs * pages_per_input
        api_boards = sum(1 for s in sources if s.get("official_api"))
        requests_max = pages_max * (1 + options.max_retries) + api_boards * 25
        requests_typical = inputs * (2 + (3 if options.pagination else 0)
                                     + (min(options.max_detail_pages, 20) if options.follow_details else 0))
        registry = self.platform.service("ai")
        provider = registry.for_ctx(ctx, "extraction") if options.use_ai else None
        ai_external = bool(provider is not None and provider.external)
        status = {"provider": getattr(provider, "name", None),
                  "free_only": bool(registry.workspace_config(ctx).get("free_only"))}
        per_input_ai = 2 if schema.get("entity") == "job" else 1
        ai_max = 0
        if ai_external:
            limit = options.max_ai_calls if options.max_ai_calls is not None else int(
                self.platform.config.extra.get("scraper_max_ai_calls") or 25)
            ai_max = min(limit, inputs * per_input_ai)
        free_only = bool(status.get("free_only"))
        if not ai_external:
            cost, note = 0.0, "no AI calls (rules only)"
        elif free_only:
            cost, note = 0.0, "free-only mode: Gemini free tier, $0 budget, no paid fallback"
        else:
            cost, note = None, "depends on the configured provider's pricing; see AI settings"
        return {"inputs": inputs, "pages_max": pages_max, "requests_max": requests_max,
                "requests_typical": min(requests_typical, requests_max), "ai_calls_max": ai_max,
                "ai_provider": status.get("provider") if ai_external else None, "ai_free_only": free_only,
                "estimated_cost_usd": cost, "cost_note": note, "browser_pages_max":
                min(options.max_browser_pages, inputs * 2) if options.browser else 0}

    def plan(self, ctx: Ctx, urls: Union[None, str, Sequence[Any]] = None, instruction: str = "", *,
             file_bytes: Optional[bytes] = None, filename: Optional[str] = None, column: Optional[str] = None,
             schema: Optional[Mapping[str, Any]] = None, options: Optional[Mapping[str, Any]] = None,
             template_id: Optional[str] = None) -> Dict[str, Any]:
        """Everything the UI shows before a run: inputs, detected sources, schema, filters,
        limits, browser status, work and cost estimate, and whether confirmation is needed."""
        instruction, schema, options = self._template_defaults(ctx, template_id, instruction, schema, options)
        accepted, rejected, report = parse_inputs(urls, file_bytes=file_bytes, filename=filename, column=column)
        crawl = CrawlOptions.from_mapping(options)
        built = clean_schema(schema) if schema else self.instruction_to_schema(ctx, instruction)
        sources = [{"url": i.url, "row": i.row, **detect_source(i.url)} for i in accepted]
        estimate = self.estimate(ctx, len(accepted), crawl, built, sources)
        reasons = []
        if len(accepted) > HIGH_VOLUME_URLS:
            reasons.append(f"{len(accepted)} URLs (more than {HIGH_VOLUME_URLS})")
        if estimate["requests_max"] > HIGH_VOLUME_REQUESTS:
            reasons.append(f"up to {estimate['requests_max']} requests (more than {HIGH_VOLUME_REQUESTS})")
        browser_available = bool(self.platform.config.extra.get("scraper_browser_enabled")
                                 or self.platform.config.extra.get("renderer_factory"))
        return {"inputs": {"accepted": len(accepted), "rejected": rejected[:200], "report": report,
                           "sample": [i.as_dict() for i in accepted[:50]]},
                "sources": sources[:200], "schema": built, "filters": built.get("filters", []),
                "criteria": built.get("criteria", {}), "options": crawl.as_dict(),
                "limits": {k: getattr(crawl, k) for k in ("max_pages", "max_records", "max_runtime_s",
                                                         "max_requests_per_domain", "concurrency",
                                                         "max_detail_pages", "max_browser_pages")},
                "browser": {"requested": crawl.browser, "available": browser_available,
                            "note": None if not crawl.browser or browser_available else
                            "browser rendering is disabled on this server; the run will use HTTP only"},
                "estimate": estimate, "requires_confirmation": bool(reasons), "confirmation_reasons": reasons}

    # --- runs --------------------------------------------------------------------------------

    def start(self, ctx: Ctx, urls: Union[None, str, Sequence[Any]] = None, instruction: str = "", *,
              file_bytes: Optional[bytes] = None, filename: Optional[str] = None, column: Optional[str] = None,
              schema: Optional[Mapping[str, Any]] = None, use_ai: Optional[bool] = None,
              max_ai_calls: Optional[int] = None, options: Optional[Mapping[str, Any]] = None,
              template_id: Optional[str] = None, confirm: bool = False, enqueue: bool = True) -> Dict[str, Any]:
        ctx.require_write()
        instruction, schema, options = self._template_defaults(ctx, template_id, instruction, schema, options)
        accepted, rejected, report = parse_inputs(urls, file_bytes=file_bytes, filename=filename, column=column)
        if not accepted:
            reasons = sorted({str(r["reason"]) for r in rejected})[:3]
            raise ValidationError("no fetchable public http(s) URLs were given"
                                  + (f" ({'; '.join(reasons)})" if reasons else ""))
        if len(accepted) > MAX_URLS:
            raise ValidationError(f"at most {MAX_URLS} URLs per run (got {len(accepted)})")
        raw_options = dict(options or {})
        if use_ai is not None:
            raw_options["use_ai"] = bool(use_ai)
        if max_ai_calls is not None:
            raw_options["max_ai_calls"] = max_ai_calls
        crawl = CrawlOptions.from_mapping(raw_options)
        schema = clean_schema(schema) if schema else self.instruction_to_schema(ctx, instruction)
        sources = [detect_source(i.url) for i in accepted]
        estimate = self.estimate(ctx, len(accepted), crawl, schema, sources)
        if not confirm and (len(accepted) > HIGH_VOLUME_URLS or estimate["requests_max"] > HIGH_VOLUME_REQUESTS):
            raise ValidationError(f"high-volume run ({len(accepted)} URLs, up to {estimate['requests_max']} "
                                  "requests): review the plan and confirm to start")
        stored_options = crawl.as_dict()
        stored_options.pop("LIMITS", None)
        run = self.platform.store.insert(ctx, "scrape_runs", {
            "instruction": (instruction or schema.get("instruction") or "(schema supplied)")[:4000],
            "schema": schema, "urls": [], "status": "queued",
            "stats": {"inputs": [i.as_dict() for i in accepted], "url_count": len(accepted),
                      "rejected": rejected[:500], "input_report": report, "options": stored_options,
                      "estimate": estimate, "template_id": template_id,
                      "progress": {"total": len(accepted), "processed": 0, "completed": 0, "failed": 0, "blocked": 0,
                                   "pages": 0, "records": 0, "stage": "queued", "current_url": None}}})
        if enqueue:
            task = self.platform.tasks.submit(ctx, "scraper", {"run_id": run["id"]}, entity_type="scrape_runs",
                                              entity_id=run["id"])
            run = self.platform.store.update(ctx, "scrape_runs", run["id"], {"task_id": task["id"]})
        audit(self.platform.store, ctx, "scraper.start", entity_type="scrape_runs", entity_id=run["id"],
              summary=f"{len(accepted)} URLs: {(instruction or '')[:200]}")
        return run

    create_run = start

    def enqueue(self, ctx: Ctx, run_id: str) -> Dict[str, Any]:
        """Queue a run that was created without starting (``start(..., enqueue=False)``)."""
        ctx.require_write()
        run = self.get(ctx, run_id)
        if run.get("task_id"):
            raise ValidationError("the run was already started")
        task = self.platform.tasks.submit(ctx, "scraper", {"run_id": run_id}, entity_type="scrape_runs",
                                          entity_id=run_id)
        return self.platform.store.update(ctx, "scrape_runs", run_id, {"task_id": task["id"]})

    def get(self, ctx: Ctx, run_id: str) -> Dict[str, Any]:
        return self.platform.store.get(ctx, "scrape_runs", run_id)

    def _set_stage(self, ctx: Ctx, run: Mapping[str, Any], status: str, **extra: Any) -> Dict[str, Any]:
        progress = {**(run["stats"].get("progress") or {}), "stage": status}
        return self.platform.store.update(ctx, "scrape_runs", run["id"], {
            "status": status, **extra, "stats": {**run["stats"], "progress": progress}})

    def pause(self, ctx: Ctx, run_id: str) -> Dict[str, Any]:
        """Pause a run. A queued run pauses now; a working one after its in-flight URLs stop."""
        ctx.require_write()
        run = self.get(ctx, run_id)
        if run["status"] not in ACTIVE_STATES:
            raise ValidationError(f"a {run['status']} run cannot be paused")
        task = self.platform.tasks.pause(ctx, run["task_id"]) if run.get("task_id") else None
        audit(self.platform.store, ctx, "scraper.pause", entity_type="scrape_runs", entity_id=run_id)
        if task is None or task.get("status") == "paused":
            return self._set_stage(ctx, run, "paused")
        return run

    def resume(self, ctx: Ctx, run_id: str) -> Dict[str, Any]:
        ctx.require_write()
        run = self.get(ctx, run_id)
        if run["status"] != "paused":
            raise ValidationError(f"a {run['status']} run cannot be resumed")
        task = self.platform.tasks.get(ctx, run["task_id"]) if run.get("task_id") else None
        if task is not None and task["status"] == "paused":
            self.platform.tasks.resume(ctx, task["id"])
        else:   # the task is gone or finished: queue a fresh one that continues from the checkpoint
            task = self.platform.tasks.submit(ctx, "scraper", {"run_id": run_id}, entity_type="scrape_runs",
                                              entity_id=run_id)
            run = self.platform.store.update(ctx, "scrape_runs", run_id, {"task_id": task["id"]})
        audit(self.platform.store, ctx, "scraper.resume", entity_type="scrape_runs", entity_id=run_id)
        return self._set_stage(ctx, run, "queued")

    def cancel(self, ctx: Ctx, run_id: str) -> Dict[str, Any]:
        """Stop a run. A queued/paused run stops now; a working one at its in-flight URLs'
        next page, keeping (and exporting) what it has scraped so far."""
        ctx.require_write()
        run = self.get(ctx, run_id)
        if run["status"] not in ACTIVE_STATES and run["status"] != "paused":
            raise ValidationError(f"a {run['status']} run cannot be cancelled")
        task = None
        if run.get("task_id"):
            try:
                task = self.platform.tasks.cancel(ctx, run["task_id"])
            except Exception:  # noqa: BLE001 - the task already finished: just mark the run
                task = None
        audit(self.platform.store, ctx, "scraper.cancel", entity_type="scrape_runs", entity_id=run_id)
        if task is not None and task.get("status") == "running":
            return run
        return self._set_stage(ctx, run, "cancelled")

    def retry(self, ctx: Ctx, run_id: str) -> Dict[str, Any]:
        """Run again: URLs without a result, plus those that failed transiently (timeout,
        network error, rate limit). Pages that refused us (CAPTCHA, WAF, login) are not retried."""
        ctx.require_write()
        run = self.get(ctx, run_id)
        if run["status"] in ACTIVE_STATES or run["status"] == "paused":
            raise ValidationError("the run is still in progress")
        results = self.platform.store.all(ctx, "scrape_results", {"run_id": run_id})
        final = {int(r["data"].get("index", -1)) for r in results if r["data"].get("outcome") not in TRANSIENT}
        todo = len(run["stats"].get("inputs") or []) - len(final)
        if todo <= 0:
            raise ValidationError("nothing to retry: every URL has a final result")
        task = self.platform.tasks.submit(ctx, "scraper", {"run_id": run_id, "retry": True},
                                          entity_type="scrape_runs", entity_id=run_id)
        run = self.platform.store.update(ctx, "scrape_runs", run_id, {"task_id": task["id"], "error": None})
        audit(self.platform.store, ctx, "scraper.retry", entity_type="scrape_runs", entity_id=run_id,
              summary=f"{todo} URLs")
        return self._set_stage(ctx, run, "queued")

    def restart(self, ctx: Ctx, run_id: str) -> Dict[str, Any]:
        """Start over as a new run with the same inputs, schema and options (the old run is kept)."""
        run = self.get(ctx, run_id)
        if run["status"] in ACTIVE_STATES or run["status"] == "paused":
            raise ValidationError("cancel the run before restarting it")
        inputs = run["stats"].get("inputs") or []
        new = self.start(ctx, [i["url"] for i in inputs], run["instruction"], schema=run["schema"],
                         options=run["stats"].get("options"), confirm=True)
        stats = {**new["stats"], "restarted_from": run_id,
                 "inputs": [{**i, "row": old.get("row"), "batch": old.get("batch"), "source": old.get("source")}
                            for i, old in zip(new["stats"]["inputs"], inputs)]}
        return self.platform.store.update(ctx, "scrape_runs", new["id"], {"stats": stats})

    def recover(self, ctx: Ctx) -> int:
        """Crash recovery (called by worker maintenance): a run whose task ended without the
        run itself finishing is set to the task's state — failed runs keep every saved result
        and can be retried; paused tasks leave the run paused."""
        fixed = 0
        for run in self.platform.store.all(ctx, "scrape_runs", {"status__in": sorted(ACTIVE_STATES)},
                                           cap=1000):
            if not run.get("task_id"):
                continue
            task = self.platform.store.find(ctx, "platform_tasks", run["task_id"])
            if task is None:
                self._set_stage(ctx, run, "failed", error="the run's background task disappeared; retry to resume")
                fixed += 1
            elif task["status"] == "failed":
                self._set_stage(ctx, run, "failed", error=f"worker: {task.get('error') or 'failed'}; saved results "
                                                          "are kept — retry to resume")
                fixed += 1
            elif task["status"] == "cancelled":
                self._set_stage(ctx, run, "cancelled")
                fixed += 1
            elif task["status"] == "paused" and run["status"] != "queued":
                self._set_stage(ctx, run, "paused")
                fixed += 1
        return fixed

    # --- results -----------------------------------------------------------------------------

    def _payload(self, ctx: Ctx, run: Mapping[str, Any]) -> Optional[Dict[str, Any]]:
        info = (run["stats"].get("files") or {}).get("json")
        if not info:
            return None
        with self.platform.storage.open(info["storage_key"]) as handle:
            return json.loads(handle.read())

    def records(self, ctx: Ctx, run_id: str, view: str = "all", *, limit: int = 500, offset: int = 0
                ) -> Dict[str, Any]:
        """Final (normalised, de-duplicated) rows of a finished run, by view."""
        if view not in ("all", "companies", "jobs"):
            raise ValidationError("view must be all, companies or jobs")
        run = self.get(ctx, run_id)
        payload = self._payload(ctx, run)
        if payload is None:
            return {"view": view, "columns": [], "items": [], "total": 0, "final": False}
        rows = payload.get(view if view != "all" else "records")
        if rows is None:
            rows = views(payload.get("records") or [], run["schema"])[view]
        columns = columns_for(run["schema"], view)
        if view == "companies" and run["schema"].get("entity") == "job":
            columns = columns + ["job_count"]
        limit = max(1, min(int(limit), 10000))
        return {"view": view, "columns": columns, "items": rows[offset: offset + limit], "total": len(rows),
                "final": True, "summary": payload.get("summary") or {}}

    def pages(self, ctx: Ctx, run_id: str, *, limit: int = 500, offset: int = 0,
              outcome: Optional[str] = None) -> Dict[str, Any]:
        self.get(ctx, run_id)
        filters: Dict[str, Any] = {"run_id": run_id}
        if outcome:
            filters["outcome"] = outcome
        page = self.platform.store.list(ctx, "scrape_pages", filters, order="created_at", limit=limit, offset=offset)
        return {"items": page.rows, "total": page.total}

    def errors(self, ctx: Ctx, run_id: str, *, limit: int = 1000) -> Dict[str, Any]:
        payload = self._payload(ctx, self.get(ctx, run_id)) or {}
        items = payload.get("errors") or []
        return {"items": items[:limit], "total": len(items)}

    def evidence(self, ctx: Ctx, run_id: str, *, limit: int = 500, offset: int = 0) -> Dict[str, Any]:
        """Per record: every field's value, method, confidence, evidence, source, status, conflicts
        and rejected values."""
        payload = self._payload(ctx, self.get(ctx, run_id)) or {}
        rows = payload.get("records") or []
        out = []
        for record in rows[offset: offset + limit]:
            out.append({"source_url": record.get("source_url"), "input_row": record.get("input_row"),
                        "label": record.get("job_title") or record.get("company_name") or record.get("source_url"),
                        "fields": {name: {**info, "value": record.get(name)}
                                   for name, info in (record.get("_evidence") or {}).items()},
                        "status": record.get("_field_status") or {}, "conflicts": record.get("_conflicts") or {},
                        "rejected": record.get("_rejected") or {}, "source_urls": record.get("source_urls") or []})
        return {"items": out, "total": len(rows)}

    def file(self, ctx: Ctx, run_id: str, fmt: str, view: str = "all") -> Optional[Dict[str, Any]]:
        files = self.get(ctx, run_id)["stats"].get("files") or {}
        key = f"{view}.{fmt}" if fmt == "csv" and view in ("companies", "jobs", "pages", "errors") else fmt
        return files.get(key) if key in FILE_NAMES else None   # unknown format: 404, as in V1

    # --- CRM ---------------------------------------------------------------------------------

    def match_crm(self, ctx: Ctx, run_id: str) -> Dict[str, Any]:
        matches = crm_actions.match_run(self.platform, ctx, run_id)
        summary: Dict[str, int] = {}
        for m in matches:
            summary[m["match"]] = summary.get(m["match"], 0) + 1
        return {"items": matches, "summary": summary}

    def propose(self, ctx: Ctx, run_id: str, actions: Sequence[str] = ("company", "job"), **kw: Any) -> Dict[str, Any]:
        return crm_actions.propose(self.platform, ctx, run_id, actions, **kw)

    def proposals(self, ctx: Ctx, run_id: str, *, status: Optional[str] = None) -> Dict[str, Any]:
        filters: Dict[str, Any] = {"run_id": run_id}
        if status:
            filters["status"] = status
        rows = self.platform.store.all(ctx, "scrape_proposals", filters, order="created_at", cap=10000)
        return {"items": rows, "total": len(rows)}

    def review(self, ctx: Ctx, proposal_ids: Iterable[str], decision: str) -> List[Dict[str, Any]]:
        return crm_actions.review(self.platform, ctx, proposal_ids, decision)

    def apply_proposals(self, ctx: Ctx, proposal_ids: Iterable[str]) -> Dict[str, Any]:
        return crm_actions.apply(self.platform, ctx, proposal_ids)

