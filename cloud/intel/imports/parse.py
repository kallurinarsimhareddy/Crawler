"""Reading uploaded CSV, XLSX and JSON files into header + rows, exactly as written.

Cell values are returned as strings (numbers without a spurious ``.0``, dates in
ISO form) and are never reinterpreted here: normalisation happens later, and the
original value is always kept beside the normalised one.
"""

from __future__ import annotations

import csv
import io
import json
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, BinaryIO, Dict, Iterator, List, Optional, Tuple

__all__ = ["ParseError", "ParsedFile", "detect_format", "parse_file", "iter_rows", "iter_rows_stream"]

DEFAULT_MAX_ROWS = 200_000


class ParseError(ValueError):
    pass


@dataclass
class ParsedFile:
    format: str
    columns: List[str]
    row_count: int
    sheet: Optional[str] = None
    problems: List[str] = field(default_factory=list)


def detect_format(filename: str) -> str:
    lowered = filename.lower()
    if lowered.endswith(".csv") or lowered.endswith(".tsv") or lowered.endswith(".txt"):
        return "csv"
    if lowered.endswith(".xlsx") or lowered.endswith(".xlsm"):
        return "xlsx"
    if lowered.endswith(".json"):
        return "json"
    raise ParseError(f"{filename}: only CSV, XLSX and JSON files can be imported")


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value).strip()


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _csv_rows(data: bytes) -> Iterator[List[str]]:
    text = _decode(data)
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text, newline=""), dialect)
    while True:
        try:
            row = next(reader)
        except StopIteration:
            return
        except csv.Error as error:  # an oversized field, a NUL byte ...: a file problem, not a crash
            raise ParseError(f"malformed CSV near line {reader.line_num}: {error}") from None
        yield [cell.strip() for cell in row]


def _xlsx_rows(data: bytes, sheet: Optional[str]) -> Tuple[Optional[str], Iterator[List[str]]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as error:  # noqa: BLE001 - any openpyxl failure means "not a readable workbook"
        raise ParseError(f"not a readable XLSX workbook ({type(error).__name__})") from error
    if sheet is not None:
        if sheet not in workbook.sheetnames:
            raise ParseError(f"the workbook has no sheet named {sheet!r}")
        names = [sheet]
    else:
        names = list(workbook.sheetnames)
    for name in names:
        rows = ([_cell(v) for v in row] for row in workbook[name].iter_rows(values_only=True))
        first = next((r for r in rows if any(r)), None)
        if first is None:
            continue

        def chain(first=first, rows=rows):
            yield first
            yield from rows

        return name, chain()
    return names[0] if names else None, iter(())


def _json_rows(data: bytes) -> Iterator[List[str]]:
    try:
        payload = json.loads(_decode(data))
    except json.JSONDecodeError as error:
        raise ParseError(f"not valid JSON ({error.msg} at line {error.lineno})") from error
    if isinstance(payload, dict):
        for key in ("rows", "data", "items", "records"):
            if isinstance(payload.get(key), list):
                payload = payload[key]
                break
    if not isinstance(payload, list) or not all(isinstance(item, dict) for item in payload):
        raise ParseError("JSON must be a list of objects (or {\"rows\": [...]})")
    columns: List[str] = []
    for item in payload:
        for key in item:
            if key not in columns:
                columns.append(str(key))
    yield columns
    for item in payload:
        yield [_cell(item.get(column)) for column in columns]


def _raw_rows(fmt: str, data: bytes, sheet: Optional[str]) -> Tuple[Optional[str], Iterator[List[str]]]:
    if fmt == "csv":
        return None, _csv_rows(data)
    if fmt == "xlsx":
        return _xlsx_rows(data, sheet)
    if fmt == "json":
        return None, _json_rows(data)
    raise ParseError(f"unsupported format {fmt}")


def iter_rows(fmt: str, data: bytes, *, sheet: Optional[str] = None, max_rows: int = DEFAULT_MAX_ROWS
              ) -> Iterator[Tuple[int, Dict[str, str]]]:
    """``(row_number, {header: value})`` for every non-empty data row.

    ``row_number`` is 1-based and counts data rows, as a person counting rows
    under the header would. Duplicate headers keep the first column.
    """
    _sheet, rows = _raw_rows(fmt, data, sheet)
    header: Optional[List[str]] = None
    number = 0
    for raw in rows:
        if header is None:
            if not any(raw):
                continue
            header = raw
            continue
        if not any(raw):
            continue
        number += 1
        if number > max_rows:
            return
        record: Dict[str, str] = {}
        for index, name in enumerate(header):
            if name and name not in record:
                record[name] = raw[index] if index < len(raw) else ""
        yield number, record


def _stream_text(handle: BinaryIO) -> io.TextIOWrapper:
    """A text view of a binary file without reading it all: the encoding is chosen from the
    first 64 KB (UTF-8 with or without BOM, else cp1252, as :func:`_decode` does); a later
    undecodable UTF-8 byte becomes U+FFFD instead of failing the whole file."""
    head = handle.read(65536)
    handle.seek(0)
    encoding = "cp1252"
    try:
        head.decode("utf-8-sig")
        encoding = "utf-8-sig"
    except UnicodeDecodeError as error:
        if error.start >= len(head) - 4:  # only a multi-byte character cut at the sample edge
            encoding = "utf-8-sig"
    return io.TextIOWrapper(handle, encoding=encoding, errors="replace", newline="")


def iter_rows_stream(fmt: str, handle: BinaryIO, *, max_rows: int = DEFAULT_MAX_ROWS
                     ) -> Iterator[Tuple[int, Dict[str, str]]]:
    """:func:`iter_rows` over a seekable binary file handle with memory bounded by the row,
    not the file: CSV is decoded and parsed incrementally, XLSX is read in openpyxl's
    read-only (streaming) mode. Malformed CSV raises :class:`ParseError` (never ``csv.Error``)."""
    if fmt == "csv":
        text = _stream_text(handle)
        sample = text.read(20000)
        text.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel

        def rows() -> Iterator[List[str]]:
            reader = csv.reader(text, dialect)
            while True:
                try:
                    row = next(reader)
                except StopIteration:
                    return
                except csv.Error as error:
                    raise ParseError(f"malformed CSV near line {reader.line_num}: {error}") from None
                yield [cell.strip() for cell in row]

        raw_rows: Iterator[List[str]] = rows()
    elif fmt == "xlsx":
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(handle, read_only=True, data_only=True)
        except Exception as error:  # noqa: BLE001 - any openpyxl failure means "not a readable workbook"
            raise ParseError(f"not a readable XLSX workbook ({type(error).__name__})") from error

        def xlsx() -> Iterator[List[str]]:
            for name in workbook.sheetnames:  # the first sheet with data, as _xlsx_rows does
                sheet_rows = ([_cell(v) for v in row] for row in workbook[name].iter_rows(values_only=True))
                first = next((r for r in sheet_rows if any(r)), None)
                if first is None:
                    continue
                yield first
                yield from sheet_rows
                return

        raw_rows = xlsx()
    else:
        raise ParseError(f"unsupported format {fmt}")
    header: Optional[List[str]] = None
    number = 0
    for raw in raw_rows:
        if header is None:
            if any(raw):
                header = raw
            continue
        if not any(raw):
            continue
        number += 1
        if number > max_rows:
            return
        record: Dict[str, str] = {}
        for index, name in enumerate(header):
            if name and name not in record:
                record[name] = raw[index] if index < len(raw) else ""
        yield number, record


def parse_file(filename: str, data: bytes, *, sheet: Optional[str] = None,
               max_rows: int = DEFAULT_MAX_ROWS) -> ParsedFile:
    """Header and row count, plus every structural problem found."""
    fmt = detect_format(filename)
    if not data:
        return ParsedFile(fmt, [], 0, problems=["the file is empty"])
    chosen, rows = _raw_rows(fmt, data, sheet)
    header: Optional[List[str]] = None
    count = 0
    problems: List[str] = []
    for raw in rows:
        if header is None:
            if any(raw):
                header = raw
            continue
        if any(raw):
            count += 1
            if count > max_rows:
                problems.append(f"more than {max_rows:,} rows; split the file")
                break
    if header is None:
        return ParsedFile(fmt, [], 0, sheet=chosen, problems=["the file has no header row"])
    blanks = [i + 1 for i, name in enumerate(header) if not name]
    if blanks:
        problems.append("unnamed column(s) at position " + ", ".join(map(str, blanks[:10])))
    seen, duplicates = set(), []
    for name in header:
        if name and name.lower() in seen and name not in duplicates:
            duplicates.append(name)
        seen.add(name.lower())
    if duplicates:
        problems.append("duplicate column header(s): " + ", ".join(duplicates[:10]))
    if count == 0:
        problems.append("the file has a header but no data rows")
    return ParsedFile(fmt, [name for name in header if name], min(count, max_rows), sheet=chosen, problems=problems)
