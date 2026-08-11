"""Unit tests for :mod:`exporters.excel_exporter`."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from exporters.excel_exporter import DEFAULT_OUTPUT_PATH, HEADERS, export_jobs
from models.job import Job


def _job(index: int = 0) -> Job:
    """Build one exportable posting."""
    return Job(
        company_name=f"Company {index}",
        job_title=f"Engineer {index}",
        job_url=f"https://acme.com/jobs/{index}",
        location="Austin, TX",
        country="United States",
        career_page_url="https://acme.com/careers",
        platform="Workday",
    )


class TestExportScales(unittest.TestCase):
    """Writing a full run's worth of postings must stay linear.

    ``Worksheet.max_row`` is ``max()`` over every cell in the sheet. Reading it
    once per hyperlink made the export quadratic, which was invisible at the
    fifteen hundred rows an early run produced and cost three minutes at forty
    thousand. Counting accesses catches a reintroduction exactly, where a
    wall-clock assertion would only catch it flakily.
    """

    def setUp(self) -> None:
        self._directory = tempfile.TemporaryDirectory()
        self.path = Path(self._directory.name) / "jobs.xlsx"

    def tearDown(self) -> None:
        self._directory.cleanup()

    def test_max_row_is_not_consulted_per_row(self) -> None:
        real = Worksheet.max_row
        reads = []

        def counted(self):  # type: ignore[no-untyped-def]
            reads.append(1)
            return real.fget(self)

        with mock.patch.object(Worksheet, "max_row", property(counted)):
            export_jobs([_job(index) for index in range(400)], self.path)

        self.assertLess(len(reads), 10, f"max_row was read {len(reads)} times for 400 rows")

    def test_links_land_on_the_right_rows_at_scale(self) -> None:
        export_jobs([_job(index) for index in range(300)], self.path)

        sheet = load_workbook(self.path)["Jobs"]

        self.assertEqual(sheet.max_row, 301)
        for row in (2, 150, 301):
            with self.subTest(row=row):
                cell = sheet.cell(row=row, column=5)
                self.assertEqual(cell.hyperlink.target, cell.value)


class TestExportJobs(unittest.TestCase):
    """The workbook is written with the agreed columns and formatting."""

    def setUp(self) -> None:
        self._directory = tempfile.TemporaryDirectory()
        self.path = Path(self._directory.name) / "jobs.xlsx"

    def tearDown(self) -> None:
        self._directory.cleanup()

    def _export(self, jobs) -> object:
        export_jobs(jobs, self.path)
        return load_workbook(self.path)["Jobs"]

    def test_headers_in_the_agreed_order(self) -> None:
        sheet = self._export([_job()])
        self.assertEqual([cell.value for cell in sheet[1]], list(HEADERS))
        self.assertEqual(
            list(HEADERS),
            [
                "Company Name",
                "Job Title",
                "Location",
                "Country",
                "Job URL",
                "Career Page URL",
                "Platform",
            ],
        )

    def test_header_is_bold(self) -> None:
        sheet = self._export([_job()])
        self.assertTrue(all(cell.font.bold for cell in sheet[1]))

    def test_header_is_frozen(self) -> None:
        sheet = self._export([_job()])
        self.assertEqual(sheet.freeze_panes, "A2")

    def test_one_row_per_job(self) -> None:
        sheet = self._export([_job(i) for i in range(7)])
        self.assertEqual(sheet.max_row, 8)  # 7 jobs + header

    def test_values_land_in_the_right_columns(self) -> None:
        sheet = self._export([_job(3)])
        self.assertEqual(
            [cell.value for cell in sheet[2]],
            [
                "Company 3",
                "Engineer 3",
                "Austin, TX",
                "United States",
                "https://acme.com/jobs/3",
                "https://acme.com/careers",
                "Workday",
            ],
        )

    def test_columns_are_widened_to_fit(self) -> None:
        long_title = "Principal Distributed Systems Engineer, Platform Infrastructure Group"
        job = Job(company_name="Acme", job_title=long_title, job_url="https://acme.com/1")
        export_jobs([job], self.path)
        sheet = load_workbook(self.path)["Jobs"]

        self.assertGreater(sheet.column_dimensions["B"].width, len("Job Title"))
        self.assertLessEqual(sheet.column_dimensions["B"].width, 70)

    def test_job_urls_are_clickable(self) -> None:
        sheet = self._export([_job()])
        self.assertEqual(sheet.cell(row=2, column=5).hyperlink.target, "https://acme.com/jobs/0")

    def test_empty_export_still_writes_a_header(self) -> None:
        sheet = self._export([])
        self.assertEqual(sheet.max_row, 1)
        self.assertEqual([cell.value for cell in sheet[1]], list(HEADERS))

    def test_creates_missing_directories(self) -> None:
        nested = Path(self._directory.name) / "deep" / "output" / "jobs.xlsx"
        export_jobs([_job()], nested)
        self.assertTrue(nested.is_file())

    def test_overwrites_a_previous_export(self) -> None:
        export_jobs([_job(i) for i in range(5)], self.path)
        export_jobs([_job()], self.path)
        self.assertEqual(load_workbook(self.path)["Jobs"].max_row, 2)

    def test_no_partial_file_is_left_behind(self) -> None:
        export_jobs([_job()], self.path)
        self.assertEqual(list(self.path.parent.glob("*.partial.xlsx")), [])

    def test_default_path_is_the_agreed_output(self) -> None:
        self.assertEqual(DEFAULT_OUTPUT_PATH.as_posix(), "output/jobs.xlsx")


if __name__ == "__main__":
    unittest.main()
