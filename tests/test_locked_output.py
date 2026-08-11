"""Tests for writing when the destination file is locked.

A crawl costs tens of minutes. If the destination happens to be open in Excel —
which takes an exclusive lock on Windows — the results must land beside it, not
be thrown away. ``os.replace`` is patched to raise, which is exactly what
Windows does against a locked target.

Every exporter routes its write through :mod:`exporters._atomic`, so patching
``os.replace`` there covers the workbook, the failure CSV and the version 2
reports at once.
"""

from __future__ import annotations

import csv
import os
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from crawler.crawler_engine import CrawlResult
from crawler.platform_detector import Platform
from exporters.excel_exporter import export_jobs
from exporters.failure_exporter import export_failures
from models.job import Job


def _job(index: int = 0) -> Job:
    """Build one posting."""
    return Job(
        company_name=f"Company {index}",
        job_title=f"Engineer {index}",
        job_url=f"https://acme.com/jobs/{index}",
    )


_ACCESS_DENIED = OSError(5, "Access is denied")


class TestLockedWorkbook(unittest.TestCase):
    """The workbook survives a locked destination."""

    def setUp(self) -> None:
        self._directory = tempfile.TemporaryDirectory()
        self.path = Path(self._directory.name) / "jobs.xlsx"

    def tearDown(self) -> None:
        self._directory.cleanup()

    def test_falls_back_to_a_sibling_file(self) -> None:
        real_replace = os.replace

        def deny_target(src, dst):
            if str(dst) == str(self.path):
                raise _ACCESS_DENIED
            return real_replace(src, dst)

        with mock.patch("exporters._atomic.os.replace", side_effect=deny_target):
            written = export_jobs([_job(i) for i in range(3)], self.path)

        self.assertNotEqual(written, self.path)
        self.assertTrue(written.is_file())
        self.assertTrue(written.name.startswith("jobs-"))
        self.assertEqual(written.suffix, ".xlsx")

    def test_fallback_holds_the_complete_data(self) -> None:
        real_replace = os.replace

        def deny_target(src, dst):
            if str(dst) == str(self.path):
                raise _ACCESS_DENIED
            return real_replace(src, dst)

        with mock.patch("exporters._atomic.os.replace", side_effect=deny_target):
            written = export_jobs([_job(i) for i in range(25)], self.path)

        sheet = load_workbook(written)["Jobs"]
        self.assertEqual(sheet.max_row, 26)  # 25 jobs + header
        self.assertEqual(sheet.cell(row=2, column=2).value, "Engineer 0")

    def test_an_existing_locked_file_is_left_untouched(self) -> None:
        export_jobs([_job()], self.path)
        original = self.path.read_bytes()

        real_replace = os.replace

        def deny_target(src, dst):
            if str(dst) == str(self.path):
                raise _ACCESS_DENIED
            return real_replace(src, dst)

        with mock.patch("exporters._atomic.os.replace", side_effect=deny_target):
            export_jobs([_job(i) for i in range(9)], self.path)

        self.assertEqual(self.path.read_bytes(), original)

    def test_no_partial_file_survives_the_fallback(self) -> None:
        real_replace = os.replace

        def deny_target(src, dst):
            if str(dst) == str(self.path):
                raise _ACCESS_DENIED
            return real_replace(src, dst)

        with mock.patch("exporters._atomic.os.replace", side_effect=deny_target):
            export_jobs([_job()], self.path)

        self.assertEqual(list(self.path.parent.glob("*.partial.xlsx")), [])

    def test_fallback_can_be_refused(self) -> None:
        with mock.patch("exporters._atomic.os.replace", side_effect=_ACCESS_DENIED):
            with self.assertRaises(OSError):
                export_jobs([_job()], self.path, fallback_when_locked=False)


class TestLockedFailureCsv(unittest.TestCase):
    """The failure CSV behaves the same way."""

    def setUp(self) -> None:
        self._directory = tempfile.TemporaryDirectory()
        self.path = Path(self._directory.name) / "failed_companies.csv"

    def tearDown(self) -> None:
        self._directory.cleanup()

    def test_falls_back_and_keeps_the_rows(self) -> None:
        real_replace = os.replace

        def deny_target(src, dst):
            if str(dst) == str(self.path):
                raise _ACCESS_DENIED
            return real_replace(src, dst)

        pairs = [({"company": "Acme"}, CrawlResult("Acme", Platform.ICIMS, error="blocked"))]

        with mock.patch("exporters._atomic.os.replace", side_effect=deny_target):
            written = export_failures(pairs, self.path)

        self.assertNotEqual(written, self.path)
        with written.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle))
        self.assertEqual(rows[1][0], "Acme")
        self.assertEqual(rows[1][5], "blocked")


if __name__ == "__main__":
    unittest.main()
