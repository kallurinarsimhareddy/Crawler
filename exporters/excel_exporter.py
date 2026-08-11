"""Write crawl results to ``output/jobs.xlsx``.

The workbook is the deliverable of the whole pipeline. It has one row per job
posting and exactly these columns, in this order::

    Company Name | Job Title | Location | Country | Job URL | Career Page URL | Platform

    >>> from exporters.excel_exporter import export_jobs
    >>> export_jobs(jobs)
    WindowsPath('output/jobs.xlsx')

The sheet is written for a human reader: a bold, frozen header row, column
widths fitted to their contents, and job URLs as clickable links. Writing goes
to a temporary file that replaces the target only on success, so an interrupted
run cannot leave a half-written workbook behind.
"""

from __future__ import annotations

from pathlib import Path
from typing import Final, Iterable, List, Sequence

from loguru import logger
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from exporters._atomic import write_atomically
from models.job import EXPORT_COLUMNS, Job

__all__ = ["DEFAULT_OUTPUT_PATH", "SHEET_NAME", "export_jobs"]

#: Default destination, relative to the project root.
DEFAULT_OUTPUT_PATH: Final[Path] = Path("output") / "jobs.xlsx"

#: Name of the single worksheet.
SHEET_NAME: Final[str] = "Jobs"

#: Column headers, in export order.
HEADERS: Final[tuple] = tuple(header for _, header in EXPORT_COLUMNS)

#: Bounds on a fitted column, in characters.
_MIN_WIDTH: Final[int] = 12
_MAX_WIDTH: Final[int] = 70

#: Excel refuses a hyperlink longer than this, so over-long URLs stay plain text.
_MAX_LINK_LENGTH: Final[int] = 255

#: Columns rendered as clickable links.
_LINK_HEADERS: Final[frozenset] = frozenset({"Job URL", "Career Page URL"})


def _write_header(sheet: Worksheet) -> None:
    """Write and style the header row.

    Args:
        sheet: The worksheet to write into.
    """
    sheet.append(list(HEADERS))

    bold = Font(bold=True)
    for cell in sheet[1]:
        cell.font = bold
        cell.alignment = Alignment(vertical="center")

    # Everything below row 1 scrolls under the header.
    sheet.freeze_panes = "A2"


def _write_rows(sheet: Worksheet, jobs: Sequence[Job]) -> None:
    """Write one row per job.

    Args:
        sheet: The worksheet to write into.
        jobs: The jobs to write, in order.
    """
    # The row number is counted here rather than read back from
    # ``sheet.max_row``. That property is ``max()`` over every cell in the
    # sheet, so consulting it once per link made the whole export quadratic —
    # forty thousand postings took three minutes, almost all of it in max().
    # One row is appended per job and the header is row 1, so the number is
    # already known.
    for row_number, job in enumerate(jobs, start=2):
        row = job.to_row()
        sheet.append([row[header] for header in HEADERS])

        for index, header in enumerate(HEADERS, start=1):
            if header not in _LINK_HEADERS:
                continue

            value = row[header]
            if value and len(value) <= _MAX_LINK_LENGTH:
                cell = sheet.cell(row=row_number, column=index)
                cell.hyperlink = value
                cell.style = "Hyperlink"


def _fit_columns(sheet: Worksheet, jobs: Sequence[Job]) -> None:
    """Set each column's width to fit its contents.

    Args:
        sheet: The worksheet to adjust.
        jobs: The rows that were written, measured without re-reading cells.
    """
    for index, (attribute, header) in enumerate(EXPORT_COLUMNS, start=1):
        longest = len(header)
        for job in jobs:
            longest = max(longest, len(getattr(job, attribute)))

        width = min(max(longest + 2, _MIN_WIDTH), _MAX_WIDTH)
        sheet.column_dimensions[get_column_letter(index)].width = width


def export_jobs(
    jobs: Iterable[Job],
    output_path: Path | str = DEFAULT_OUTPUT_PATH,
    fallback_when_locked: bool = True,
) -> Path:
    """Write jobs to an Excel workbook.

    If the destination is locked — the usual cause being that it is open in
    Excel, which takes an exclusive lock on Windows — the workbook is written
    beside it under a timestamped name rather than discarded. A crawl costs
    tens of minutes; an open spreadsheet must not throw that away.

    Args:
        jobs: The postings to export, in the order they should appear.
        output_path: Destination workbook. Parent directories are created.
        fallback_when_locked: Whether to write beside a locked destination.
            Set ``False`` to require the exact path or fail.

    Returns:
        The path actually written, which is the fallback if the destination was
        locked.

    Raises:
        OSError: If the workbook cannot be written at all.
    """
    rows: List[Job] = list(jobs)
    path = Path(output_path)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME

    _write_header(sheet)
    _write_rows(sheet, rows)
    _fit_columns(sheet, rows)
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(sheet.max_row, 1)}"

    try:
        written = write_atomically(path, workbook.save, fallback_when_locked)
    finally:
        workbook.close()

    logger.success("Wrote {} job(s) to {}", len(rows), written)
    return written
